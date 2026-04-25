from __future__ import annotations

import base64
import io
import json
import hashlib
import random
import re
import secrets
from datetime import date, datetime, timedelta
from functools import lru_cache
from pathlib import Path
from typing import Any

import altair as alt
import pandas as pd
import plotly.graph_objects as go
import requests
import streamlit as st
import streamlit.components.v1 as components
from dotenv import load_dotenv

load_dotenv()


st.set_page_config(
    page_title="ITI Monthly Feedback Dashboard",
    page_icon="📘",
    layout="wide",
)


DEFAULT_API_BASE = "http://127.0.0.1:8000"
ROLE_LABEL_TO_KEY = {
    "ITI Trainee": "iti_trainee",
    "Student": "student",
    "Trainer": "trainer",
    "Principal": "principal",
    "OJT Institute Officer (ATO/DTO/TO)": "ojt_institute_officer",
    "OJT Supervisor": "ojt_supervisor",
    "Admin": "admin",
}
ROLE_ALLOWED_FORM_IDS = {
    "iti_trainee": {"iti_trainee_feedback_form"},
    "student": {"iti_trainee_feedback_form", "student_feedback_form_in_plant_training"},
    "trainee": {"iti_trainee_feedback_form", "student_feedback_form_in_plant_training"},
    "trainer": {"iti_training_officer_feedback_form"},
    "principal": {"principal_feedback_form"},
    "ojt_institute_officer": {"institute_feedback_form_in_plant_training_ato_dto_to"},
    "ojt_trainer": {"institute_feedback_form_in_plant_training_ato_dto_to"},
    "ojt_supervisor": {"supervisor_feedback_form_iti_in_plant_training_evaluation"},
    "supervisor": {"supervisor_feedback_form_iti_in_plant_training_evaluation"},
    "admin": set(),
}
TRAINER_TECH_ROLES = {"trainer"}
STUDENT_TECH_ROLES = {"student", "ojt_student", "iti_trainee", "trainee"}
ANALYTICS_ALLOWED_ROLES = {"trainer", "principal"}
LOGIN_TO_CATEGORY_ROLE = {
    "iti_trainee": "iti_trainee",
    "student": "student",
    "ojt_student": "student",
    "trainee": "student",
    "trainer": "trainer",
    "principal": "principal",
    "ojt_institute_officer": "ojt_institute_officer",
    "ojt_trainer": "ojt_institute_officer",
    "ojt_supervisor": "ojt_supervisor",
    "supervisor": "ojt_supervisor",
    "admin": "admin",
}
BASE_DIR = Path(__file__).resolve().parent
THEORY_BANK_PATH = BASE_DIR / "data" / "question_bank_theory.csv"
PRACTICAL_BANK_PATH = BASE_DIR / "data" / "question_bank_practical.csv"
QUESTION_IMAGE_DIR = BASE_DIR / "data" / "question_images"
AUTH_SESSION_STORE_PATH = BASE_DIR / "data" / "auth_sessions.json"
AUTH_SESSION_TTL_HOURS = 12
DEFAULT_TRADE_OPTIONS = [
    "Electrician",
    "Fitter",
    "Welder",
    "Instrument Mechanic",
    "Fashion Design Technology",
    "Mechanic Diesel",
    "Mechanic Motor Vehicle",
    "COPA",
    "Turner",
    "Machinist",
    "Draughtsman Civil",
    "Draughtsman Mechanical",
]


def api_get(base_url: str, path: str, params: dict[str, Any] | None = None, headers: dict[str, str] | None = None) -> tuple[bool, Any]:
    try:
        resp = requests.get(
            f"{base_url.rstrip('/')}{path}",
            params=params,
            headers=headers,
            timeout=20,
        )
        try:
            data = resp.json()
        except Exception:
            data = resp.text
        return resp.ok, data
    except Exception as exc:
        return False, str(exc)


def api_post(base_url: str, path: str, payload: dict[str, Any], headers: dict[str, str] | None = None) -> tuple[bool, Any]:
    try:
        resp = requests.post(
            f"{base_url.rstrip('/')}{path}",
            json=payload,
            headers=headers,
            timeout=120,
        )
        try:
            data = resp.json()
        except Exception:
            data = resp.text
        return resp.ok, data
    except Exception as exc:
        return False, str(exc)


def check_api(base_url: str) -> tuple[bool, str]:
    ok, data = api_get(base_url, "/openapi.json")
    if ok:
        return True, "Connected"
    return False, f"Not connected: {data}"


def show_json_block(title: str, data: Any) -> None:
    st.markdown(f"**{title}**")
    st.code(json.dumps(data, indent=2, ensure_ascii=False), language="json")


def show_warning_once(key: str, message: str) -> None:
    warning_state = st.session_state.setdefault("_shown_warnings", set())
    if key in warning_state:
        return
    warning_state.add(key)
    st.warning(message)


def value_or_empty(value: Any) -> str:
    return "" if value is None else str(value)


def auth_api_headers() -> dict[str, str]:
    auth_user = st.session_state.get("auth_user") or {}
    access_token = value_or_empty(auth_user.get("access_token")).strip()
    if not access_token:
        return {}
    return {"Authorization": f"Bearer {access_token}"}


def clear_registration_form_state() -> None:
    registration_keys = [
        "sidebar_reg_full_name",
        "sidebar_reg_username",
        "sidebar_reg_email",
        "sidebar_reg_password",
        "sidebar_reg_confirm_password",
        "sidebar_reg_role",
        "sidebar_reg_district",
        "sidebar_reg_department",
        "sidebar_reg_trade_choice",
        "sidebar_reg_assigned_trade_other",
        "sidebar_reg_assigned_year",
        "sidebar_reg_semester",
        "sidebar_reg_status",
    ]
    for key in registration_keys:
        if key in st.session_state:
            del st.session_state[key]


QUESTION_BANK_COLUMNS = [
    "source",
    "year_level",
    "trade",
    "month",
    "topic",
    "question_text",
    "question_image",
    "option_a",
    "option_b",
    "option_c",
    "option_d",
    "correct_option",
    "answer_text",
]


IMAGE_UPLOAD_EXTENSIONS = {"jpg", "jpeg", "png", "webp", "bmp", "tif", "tiff"}
QUESTION_UPLOAD_TYPES = ["csv", "json", "xlsx", "xls", "pdf", "docx", "jpg", "jpeg", "png", "webp", "bmp", "tif", "tiff"]


def empty_question_bank_df() -> pd.DataFrame:
    return pd.DataFrame(columns=QUESTION_BANK_COLUMNS)


def normalize_question_source(value: Any) -> str:
    text = _norm_text(str(value or ""))
    if not text:
        return ""
    if "practical" in text:
        return "practical"
    if "theory" in text:
        return "theory"
    if text in {"p", "prac"}:
        return "practical"
    if text in {"t", "theo"}:
        return "theory"
    return text


def infer_source_from_text(value: Any) -> str:
    text = _norm_text(str(value or ""))
    compact = re.sub(r"[^a-z0-9]+", "", text)
    if "practical" in text:
        if "theoryquestionbank" in compact or "tradetheory" in compact:
            return "theory"
        return "practical"
    if "theory" in text or "theoretical" in text:
        return "theory"
    return ""


def infer_sheet_source(sheet_name: str, sheet_text: str) -> str:
    name_source = infer_source_from_text(sheet_name)
    if name_source:
        return name_source
    compact = re.sub(r"[^a-z0-9]+", "", f"{sheet_name} {sheet_text}".lower())
    if "theoryquestionbank" in compact or "tradetheory" in compact:
        return "theory"
    if "practicalquestionbank" in compact or "tradepractical" in compact:
        return "practical"
    text = _norm_text(f"{sheet_name} {sheet_text}")
    theory_hits = len(re.findall(r"\b(theory|theoretical)\b", text))
    practical_hits = len(re.findall(r"\bpractical\b", text))
    if theory_hits > practical_hits:
        return "theory"
    if practical_hits > theory_hits:
        return "practical"
    return ""


def normalize_question_bank_df(df: pd.DataFrame, source_hint: str = "") -> pd.DataFrame:
    if df.empty:
        return empty_question_bank_df()

    aliases = {
        "source": {
            "source",
            "type",
            "question type",
            "question_type",
            "category",
            "section",
            "paper",
            "subject",
            "subject name",
            "subject_name",
        },
        "year_level": {"year", "year level", "year_level", "level", "class"},
        "trade": {"trade", "trade name", "trade_name", "course", "branch"},
        "month": {"month", "month name", "month_name"},
        "topic": {"topic", "topic name", "topic_name", "chapter", "unit", "lesson", "lesson name", "lesson_name"},
        "question_text": {
            "question",
            "questions",
            "question text",
            "question_text",
            "mcq question",
            "question description",
            "description",
        },
        "question_image": {"image", "image url", "image_url", "question image", "question_image", "question_image_url", "question_image_path"},
        "option_a": {"a", "option a", "option_a", "choice a", "choice_a"},
        "option_b": {"b", "option b", "option_b", "choice b", "choice_b"},
        "option_c": {"c", "option c", "option_c", "choice c", "choice_c"},
        "option_d": {"d", "option d", "option_d", "choice d", "choice_d"},
        "correct_option": {"correct", "correct option", "correct_option", "correct answer", "correct_answer", "answer", "answer key", "answer_key", "key"},
        "answer_text": {"answer text", "answer_text", "solution", "explanation", "descriptive answer", "descriptive_answer"},
    }

    normalized_to_original = {
        re.sub(r"[^a-z0-9]+", " ", str(col).strip().lower()).strip(): col
        for col in df.columns
    }
    rename_map: dict[Any, str] = {}
    for canonical, names in aliases.items():
        if canonical in df.columns:
            continue
        for name in names:
            original = normalized_to_original.get(re.sub(r"[^a-z0-9]+", " ", name.lower()).strip())
            if original is not None:
                rename_map[original] = canonical
                break
    out = df.rename(columns=rename_map).copy()
    for col in QUESTION_BANK_COLUMNS:
        if col not in out.columns:
            out[col] = ""

    out["source"] = out["source"].apply(normalize_question_source)
    hint_source = infer_source_from_text(source_hint)
    if hint_source:
        out.loc[out["source"].astype(str).str.strip() == "", "source"] = hint_source

    # Some files keep theory/practical in topic/subject-like text instead of a source column.
    for infer_col in ["topic", "question_text"]:
        inferred = out[infer_col].apply(infer_source_from_text)
        missing = out["source"].astype(str).str.strip() == ""
        out.loc[missing & (inferred != ""), "source"] = inferred

    for col in QUESTION_BANK_COLUMNS:
        out[col] = out[col].apply(clean_cell_text)

    return out[QUESTION_BANK_COLUMNS].copy()


def canonical_header_name(value: Any) -> str:
    normalized = re.sub(r"[^a-z0-9]+", " ", str(value or "").strip().lower()).strip()
    compact = normalized.replace(" ", "")
    header_lookup = {
        "source": "source",
        "type": "source",
        "question type": "source",
        "category": "source",
        "section": "source",
        "subject": "source",
        "subject name": "source",
        "year": "year_level",
        "year level": "year_level",
        "level": "year_level",
        "trade": "trade",
        "trade name": "trade",
        "course": "trade",
        "branch": "trade",
        "month": "month",
        "topic": "topic",
        "topic name": "topic",
        "lesson": "topic",
        "lesson name": "topic",
        "chapter": "topic",
        "unit": "topic",
        "lesson": "topic",
        "question": "question_text",
        "questions": "question_text",
        "question text": "question_text",
        "mcq question": "question_text",
        "description": "question_text",
        "image": "question_image",
        "image url": "question_image",
        "question image": "question_image",
        "a": "option_a",
        "option a": "option_a",
        "choice a": "option_a",
        "b": "option_b",
        "option b": "option_b",
        "choice b": "option_b",
        "c": "option_c",
        "option c": "option_c",
        "choice c": "option_c",
        "d": "option_d",
        "option d": "option_d",
        "choice d": "option_d",
        "correct": "correct_option",
        "correct option": "correct_option",
        "correct answer": "correct_option",
        "answer": "correct_option",
        "answer key": "correct_option",
        "solution": "answer_text",
        "explanation": "answer_text",
    }
    if "lesson" in compact and "name" in compact:
        return "topic"
    if "topic" in compact:
        return "topic"
    if "question" in compact and "no" not in compact:
        return "question_text"
    if "option" in compact and compact.endswith("a"):
        return "option_a"
    if "option" in compact and compact.endswith("b"):
        return "option_b"
    if "option" in compact and compact.endswith("c"):
        return "option_c"
    if "option" in compact and compact.endswith("d"):
        return "option_d"
    if "correct" in compact and "answer" in compact:
        return "correct_option"
    if "trade" in compact and "name" in compact:
        return "trade"
    if "year" in compact:
        return "year_level"
    return header_lookup.get(normalized, "")


def normalize_excel_sheet(raw_df: pd.DataFrame, sheet_name: str, sheet_images: list[dict[str, Any]] | None = None) -> pd.DataFrame:
    if raw_df.empty:
        return empty_question_bank_df()

    raw_df = raw_df.dropna(how="all").copy()
    if raw_df.empty:
        return empty_question_bank_df()
    sheet_text = " ".join(
        str(v)
        for v in raw_df.head(20).to_numpy().flatten().tolist()
        if str(v).strip() and str(v).lower() != "nan"
    )
    sheet_source = infer_sheet_source(sheet_name, sheet_text)

    best_idx = None
    best_score = 0
    max_scan = min(20, len(raw_df))
    for idx in range(max_scan):
        row_values = raw_df.iloc[idx].tolist()
        canonical = [canonical_header_name(v) for v in row_values]
        score = len({c for c in canonical if c})
        if score > best_score:
            best_idx = idx
            best_score = score

    if best_idx is not None and best_score >= 2:
        headers = []
        seen: dict[str, int] = {}
        topic_positions: list[int] = []
        for pos, value in enumerate(raw_df.iloc[best_idx].tolist()):
            canonical = canonical_header_name(value)
            if canonical == "topic":
                topic_positions.append(pos)
            header = canonical or f"extra_{pos}"
            seen[header] = seen.get(header, 0) + 1
            if seen[header] > 1:
                header = f"{header}_{seen[header]}"
            headers.append(header)
        parsed = raw_df.iloc[best_idx + 1 :].copy()
        parsed.columns = headers
        parsed = parsed.dropna(how="all")
        topic_cols = [headers[pos] for pos in topic_positions if pos < len(headers)]
        if topic_cols:
            primary_topic_col = topic_cols[0]
            parsed[primary_topic_col] = parsed[primary_topic_col].apply(clean_cell_text).replace("", pd.NA).ffill().fillna("")
    else:
        parsed = raw_df.copy()
        parsed.columns = [f"col_{i}" for i in range(len(parsed.columns))]
        first_nonempty_cols = [c for c in parsed.columns if parsed[c].notna().any()]
        if first_nonempty_cols:
            parsed = parsed.rename(columns={first_nonempty_cols[0]: "question_text"})

    if sheet_images:
        if "question_image" not in parsed.columns:
            parsed["question_image"] = ""
        parsed_index_values = list(parsed.index)
        for image_info in sheet_images:
            image_row_idx = int(image_info.get("row", 1)) - 1
            image_path = clean_cell_text(image_info.get("path", ""))
            if not image_path or not parsed_index_values:
                continue
            target_idx = min(
                parsed_index_values,
                key=lambda idx: (abs(int(idx) - image_row_idx), 0 if int(idx) >= image_row_idx else 1),
            )
            if not clean_cell_text(parsed.at[target_idx, "question_image"]):
                parsed.at[target_idx, "question_image"] = image_path

    out = normalize_question_bank_df(parsed, source_hint=f"{sheet_name} {sheet_source}")

    current_source = sheet_source or infer_source_from_text(sheet_name)
    current_topic = ""
    rows: list[dict[str, Any]] = []
    for _, row in out.iterrows():
        row_dict = row.to_dict()
        row_text = " ".join(str(v) for v in row_dict.values() if str(v).strip() and str(v).lower() != "nan")
        row_source = normalize_question_source(row_dict.get("source", ""))
        question_text = str(row_dict.get("question_text", "") or "").strip()
        topic_text = str(row_dict.get("topic", "") or "").strip()

        if not question_text or question_text.lower() in {"nan", "none"}:
            inferred_from_section = infer_source_from_text(row_text)
            if inferred_from_section:
                current_source = inferred_from_section
            if topic_text:
                current_topic = topic_text
            continue

        if topic_text and not question_text:
            current_topic = topic_text
            continue

        if not row_source:
            row_dict["source"] = current_source
        else:
            current_source = row_source
        if not str(row_dict.get("topic", "") or "").strip() and current_topic:
            row_dict["topic"] = current_topic
        elif str(row_dict.get("topic", "") or "").strip():
            current_topic = str(row_dict.get("topic", "") or "").strip()
        if not str(row_dict.get("topic", "") or "").strip():
            source_for_topic = normalize_question_source(row_dict.get("source", "")) or current_source
            row_dict["topic"] = "All Practical Questions" if source_for_topic == "practical" else "All Theory Questions"
        rows.append(row_dict)

    if not rows:
        return empty_question_bank_df()
    return pd.DataFrame(rows)[QUESTION_BANK_COLUMNS].copy()


def extract_xlsx_embedded_images(file_bytes: bytes, filename: str) -> dict[str, list[dict[str, Any]]]:
    try:
        from openpyxl import load_workbook
    except Exception:
        st.warning("Install `openpyxl` to extract embedded images from XLSX files.")
        return {}

    try:
        workbook = load_workbook(io.BytesIO(file_bytes), data_only=True)
    except Exception as exc:
        st.warning(f"Excel image extraction failed: {exc}")
        return {}

    images_by_sheet: dict[str, list[dict[str, Any]]] = {}
    for sheet in workbook.worksheets:
        sheet_images: list[dict[str, Any]] = []
        for idx, image in enumerate(getattr(sheet, "_images", []), start=1):
            try:
                anchor = getattr(image, "anchor", None)
                marker = getattr(anchor, "_from", None)
                row = int(getattr(marker, "row", 0)) + 1
                col = int(getattr(marker, "col", 0)) + 1
                image_bytes = image._data()
                ext = clean_cell_text(getattr(image, "format", "")) or "png"
                if not ext.startswith("."):
                    ext = f".{ext}"
                stored_path = store_question_image_bytes(
                    image_bytes,
                    f"{Path(filename).stem}-{sheet.title}-{idx}{ext}",
                )
                sheet_images.append({"row": row, "col": col, "path": stored_path})
            except Exception:
                continue
        if sheet_images:
            images_by_sheet[sheet.title] = sheet_images
    return images_by_sheet


def safe_float(value: Any, default: float = 0.0) -> float:
    try:
        return float(value)
    except Exception:
        return default


def score_band(score: float) -> str:
    if score >= 4.5:
        return "Excellent"
    if score >= 3.5:
        return "Strong"
    if score >= 2.5:
        return "Needs Improvement"
    return "Critical"


def pct(part: int, whole: int) -> float:
    if whole <= 0:
        return 0.0
    return round((part / whole) * 100.0, 1)


def _norm_text(text: str) -> str:
    return " ".join((text or "").strip().lower().split())


def clean_cell_text(value: Any) -> str:
    text = str(value if value is not None else "").strip()
    if text.lower() in {"nan", "none", "nat"}:
        return ""
    return text


@lru_cache(maxsize=1)
def load_topic_bank() -> pd.DataFrame:
    frames: list[pd.DataFrame] = []
    for path, default_source in (
        (THEORY_BANK_PATH, "theory"),
        (PRACTICAL_BANK_PATH, "practical"),
    ):
        if path.exists():
            df = pd.read_csv(path)
            if "source" not in df.columns:
                df["source"] = default_source
            frames.append(df)
    if not frames:
        return empty_question_bank_df()
    combined = pd.concat(frames, ignore_index=True)
    for col in QUESTION_BANK_COLUMNS:
        if col not in combined.columns:
            combined[col] = ""
        combined[col] = combined[col].apply(clean_cell_text)
    return combined


def get_available_trade_names() -> list[str]:
    trades = {t.strip() for t in DEFAULT_TRADE_OPTIONS if t.strip()}
    try:
        bank_df = load_topic_bank()
        if not bank_df.empty and "trade" in bank_df.columns:
            for trade in bank_df["trade"].fillna("").astype(str).tolist():
                cleaned = clean_cell_text(trade)
                if cleaned:
                    trades.add(cleaned)
    except Exception:
        pass
    return sorted(trades)


def _subject_to_source(subject: str) -> str:
    s = _norm_text(subject)
    if "practical" in s:
        return "practical"
    if "theory" in s:
        return "theory"
    return ""


def _year_matches(series: pd.Series, year_value: int) -> pd.Series:
    cleaned = series.fillna("").astype(str).str.upper().str.strip()
    if year_value == 1:
        return cleaned.str.contains(r"\b(I|1|FIRST)\b", regex=True)
    return cleaned.str.contains(r"\b(II|2|SECOND)\b", regex=True)


def collect_topic_options(
    df: pd.DataFrame,
    trade_name: str,
    year_value: int,
    subject_name: str,
) -> list[str]:
    if df.empty:
        return []
    wdf = df.copy()
    trade = _norm_text(trade_name)
    if trade:
        wdf = wdf[wdf["trade"].fillna("").astype(str).str.lower().str.contains(trade, na=False)]
    if "year_level" in wdf.columns:
        ym = _year_matches(wdf["year_level"], year_value)
        if ym.any():
            wdf = wdf[ym]
    source = _subject_to_source(subject_name)
    if source:
        wdf = wdf[wdf["source"].fillna("").astype(str).str.lower() == source]
    topics = sorted({clean_cell_text(t) for t in wdf["topic"].tolist() if clean_cell_text(t)})
    return topics


def collect_bank_month_options(
    df: pd.DataFrame,
    trade_name: str,
    year_value: int,
) -> list[str]:
    if df.empty or "month" not in df.columns:
        return []
    wdf = df.copy()
    trade = _norm_text(trade_name)
    if trade:
        wdf = wdf[wdf["trade"].fillna("").astype(str).str.lower().str.contains(trade, na=False)]
    if "year_level" in wdf.columns:
        ym = _year_matches(wdf["year_level"], year_value)
        if ym.any():
            wdf = wdf[ym]
    month_values = sorted({clean_cell_text(m) for m in wdf["month"].tolist() if clean_cell_text(m)})
    return month_values


def collect_question_options(
    df: pd.DataFrame,
    trade_name: str,
    year_value: int,
    subject_name: str,
    topic_name: str,
) -> list[str]:
    if df.empty:
        return []
    wdf = df.copy()
    trade = _norm_text(trade_name)
    if trade:
        wdf = wdf[wdf["trade"].fillna("").astype(str).str.lower().str.contains(trade, na=False)]
    if "year_level" in wdf.columns:
        ym = _year_matches(wdf["year_level"], year_value)
        if ym.any():
            wdf = wdf[ym]
    source = _subject_to_source(subject_name)
    if source:
        wdf = wdf[wdf["source"].fillna("").astype(str).str.lower() == source]
    topic = _norm_text(topic_name)
    if topic:
        wdf = wdf[wdf["topic"].fillna("").astype(str).str.lower().str.contains(topic, na=False)]
    questions = sorted({clean_cell_text(q) for q in wdf["question_text"].tolist() if clean_cell_text(q)})
    return questions


def load_uploaded_topic_question_df(uploaded_file) -> pd.DataFrame:
    if uploaded_file is None:
        return empty_question_bank_df()
    name = (uploaded_file.name or "").lower()
    if name.endswith(".csv"):
        udf = pd.read_csv(uploaded_file)
    elif name.endswith(".xlsx") or name.endswith(".xls"):
        try:
            file_bytes = uploaded_file.getvalue()
            embedded_images = extract_xlsx_embedded_images(file_bytes, name)
            sheets = pd.read_excel(io.BytesIO(file_bytes), sheet_name=None, header=None)
            frames = []
            for sheet_name, sheet_df in sheets.items():
                normalized_sheet = normalize_excel_sheet(
                    sheet_df,
                    sheet_name=str(sheet_name),
                    sheet_images=embedded_images.get(str(sheet_name), []),
                )
                frames.append(normalized_sheet)
            udf = pd.concat(frames, ignore_index=True) if frames else empty_question_bank_df()
        except Exception:
            st.warning("Install `openpyxl` to parse XLSX uploads.")
            return empty_question_bank_df()
    elif name.endswith(".json"):
        raw = json.load(uploaded_file)
        if isinstance(raw, dict):
            raw = raw.get("rows", [raw])
        udf = pd.DataFrame(raw if isinstance(raw, list) else [])
    elif name.endswith(".pdf"):
        sections = extract_pdf_question_sections(uploaded_file, name)
        if sections:
            udf = parse_question_rows_from_sections(sections)
        else:
            text, image_refs = extract_content_from_document(uploaded_file, name)
            udf = parse_question_rows_from_text(text, image_refs=image_refs)
    elif name.endswith(".docx"):
        text, image_refs = extract_content_from_document(uploaded_file, name)
        udf = parse_question_rows_from_text(text, image_refs=image_refs)
    elif any(name.endswith(f".{ext}") for ext in IMAGE_UPLOAD_EXTENSIONS):
        image_bytes = uploaded_file.getvalue()
        image_ref = store_question_image_bytes(image_bytes, name)
        text = extract_text_from_image_bytes(image_bytes)
        udf = parse_question_rows_from_text(text, image_refs=[image_ref] if image_ref else [])
        if udf.empty and image_ref:
            udf = pd.DataFrame(
                [
                    {
                        "topic": Path(name).stem.replace("_", " ").replace("-", " ").title(),
                        "question_text": "Refer to the attached image and answer the question.",
                        "question_image": image_ref,
                    }
                ]
            )
    else:
        return empty_question_bank_df()

    return normalize_question_bank_df(udf, source_hint=name)


def parse_uploaded_question_files(
    uploaded_files: list[Any],
    default_source: str,
    default_trade: str,
    default_year_level: str,
    default_month: str,
) -> pd.DataFrame:
    frames: list[pd.DataFrame] = []
    for uploaded_file in uploaded_files or []:
        try:
            uploaded_file.seek(0)
        except Exception:
            pass
        df = load_uploaded_topic_question_df(uploaded_file)
        if df.empty:
            continue
        for col in QUESTION_BANK_COLUMNS:
            if col not in df.columns:
                df[col] = ""
        df["source"] = df["source"].fillna("").astype(str)
        if default_source:
            df.loc[df["source"].str.strip() == "", "source"] = default_source
        else:
            filename_lower = value_or_empty(getattr(uploaded_file, "name", "")).lower()
            inferred_source = ""
            has_theory_name = "theory" in filename_lower
            has_practical_name = "practical" in filename_lower
            if has_practical_name and not has_theory_name:
                inferred_source = "practical"
            elif has_theory_name and not has_practical_name:
                inferred_source = "theory"
            if inferred_source:
                df.loc[df["source"].str.strip() == "", "source"] = inferred_source
        df["trade"] = df["trade"].fillna("").astype(str)
        df.loc[df["trade"].str.strip() == "", "trade"] = default_trade
        df["year_level"] = df["year_level"].fillna("").astype(str)
        df.loc[df["year_level"].str.strip() == "", "year_level"] = default_year_level
        df["month"] = df.get("month", "").fillna("").astype(str) if "month" in df.columns else ""
        df.loc[df["month"].str.strip() == "", "month"] = default_month
        frames.append(df[QUESTION_BANK_COLUMNS].copy())
    if not frames:
        return empty_question_bank_df()
    return pd.concat(frames, ignore_index=True)


def image_bytes_to_data_url(image_bytes: bytes, filename: str) -> str:
    suffix = Path(filename or "").suffix.lower().lstrip(".")
    mime_map = {
        "jpg": "image/jpeg",
        "jpeg": "image/jpeg",
        "png": "image/png",
        "webp": "image/webp",
        "bmp": "image/bmp",
        "tif": "image/tiff",
        "tiff": "image/tiff",
    }
    mime = mime_map.get(suffix, "image/png")
    encoded = base64.b64encode(image_bytes).decode("ascii")
    return f"data:{mime};base64,{encoded}"


def store_question_image_bytes(image_bytes: bytes, filename: str) -> str:
    QUESTION_IMAGE_DIR.mkdir(parents=True, exist_ok=True)
    suffix = Path(filename or "").suffix.lower()
    if suffix not in {".jpg", ".jpeg", ".png", ".webp", ".bmp", ".tif", ".tiff"}:
        suffix = ".png"
    digest = hashlib.sha256(image_bytes).hexdigest()[:24]
    image_path = QUESTION_IMAGE_DIR / f"{digest}{suffix}"
    if not image_path.exists():
        image_path.write_bytes(image_bytes)
    return str(image_path.relative_to(BASE_DIR))


def resolve_question_image_source(image_ref: str) -> str:
    ref = value_or_empty(image_ref).strip()
    if not ref:
        return ""
    if ref.startswith(("http://", "https://", "data:")):
        return ref
    path = Path(ref)
    if not path.is_absolute():
        path = BASE_DIR / path
    return str(path)


def extract_text_from_image_bytes(image_bytes: bytes) -> str:
    try:
        from PIL import Image
        import pytesseract
    except Exception:
        show_warning_once(
            "image_ocr_deps_missing",
            "Install `pillow` and `pytesseract` to extract text from image uploads.",
        )
        return ""
    try:
        image = Image.open(io.BytesIO(image_bytes))
        return pytesseract.image_to_string(image) or ""
    except Exception as exc:
        exc_text = str(exc)
        if "tesseract is not installed" in exc_text.lower() or "tesseractnotfounderror" in exc.__class__.__name__.lower():
            show_warning_once(
                "image_ocr_tesseract_missing",
                "Image OCR is unavailable because Tesseract is not installed or not in PATH. Upload will continue without OCR text extraction.",
            )
        else:
            show_warning_once("image_ocr_runtime_failed", f"Image OCR failed: {exc}")
        return ""


def extract_content_from_document(uploaded_file, filename: str) -> tuple[str, list[str]]:
    filename = filename.lower()
    if filename.endswith(".pdf"):
        try:
            from pypdf import PdfReader
        except Exception:
            st.warning("Install `pypdf` to parse PDF uploads.")
            return "", []
        file_bytes = uploaded_file.getvalue()
        reader = PdfReader(io.BytesIO(file_bytes))
        text = "\n".join((p.extract_text() or "") for p in reader.pages)
        try:
            from pdf2image import convert_from_bytes
        except Exception:
            if not text.strip():
                st.warning("This PDF has no selectable text. Install `pdf2image` plus OCR dependencies for scanned PDFs.")
            return text, []
        page_texts: list[str] = []
        image_refs: list[str] = []
        try:
            for idx, page_image in enumerate(convert_from_bytes(file_bytes), start=1):
                buf = io.BytesIO()
                page_image.save(buf, format="PNG")
                page_bytes = buf.getvalue()
                image_refs.append(store_question_image_bytes(page_bytes, f"page-{idx}.png"))
                if not text.strip():
                    page_texts.append(extract_text_from_image_bytes(page_bytes))
        except Exception as exc:
            st.warning(f"Scanned PDF OCR failed: {exc}")
        return text if text.strip() else "\n".join(page_texts), image_refs
    if filename.endswith(".docx"):
        try:
            from docx import Document
        except Exception:
            st.warning("Install `python-docx` to parse Word uploads.")
            return "", []
        doc = Document(uploaded_file)
        text_parts = [p.text for p in doc.paragraphs if p.text]
        image_refs: list[str] = []
        for rel in doc.part.rels.values():
            try:
                if not str(rel.reltype).endswith("/image"):
                    continue
                blob = rel.target_part.blob
                partname = str(rel.target_part.partname)
                image_refs.append(store_question_image_bytes(blob, partname))
                ocr_text = extract_text_from_image_bytes(blob)
                if ocr_text.strip():
                    text_parts.append(ocr_text)
            except Exception:
                continue
        return "\n".join(text_parts), image_refs
    return "", []


def extract_pdf_question_sections(uploaded_file, filename: str) -> list[tuple[str, str]]:
    filename = filename.lower()
    if not filename.endswith(".pdf"):
        return []
    try:
        from pypdf import PdfReader
    except Exception:
        st.warning("Install `pypdf` to parse PDF uploads.")
        return []

    file_bytes = uploaded_file.getvalue()
    try:
        reader = PdfReader(io.BytesIO(file_bytes))
    except Exception as exc:
        st.warning(f"PDF parse failed: {exc}")
        return []

    page_images: list[str] = []
    try:
        from pdf2image import convert_from_bytes

        for idx, page_image in enumerate(convert_from_bytes(file_bytes), start=1):
            buf = io.BytesIO()
            page_image.save(buf, format="PNG")
            page_images.append(store_question_image_bytes(buf.getvalue(), f"{Path(filename).stem}-page-{idx}.png"))
    except Exception:
        page_images = []

    sections: list[tuple[str, str]] = []
    for idx, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        image_ref = page_images[idx - 1] if idx - 1 < len(page_images) else ""
        if text or image_ref:
            sections.append((text, image_ref))
    return sections


def parse_question_rows_from_text(text: str, image_refs: list[str] | None = None) -> pd.DataFrame:
    lines = [ln.strip() for ln in (text or "").splitlines() if ln.strip()]
    if not lines:
        return empty_question_bank_df()

    rows: list[dict[str, str]] = []
    current_topic = "Uploaded Topic"
    image_refs = image_refs or [""]
    image_idx = 0
    i = 0
    while i < len(lines):
        ln = lines[i]
        if len(ln) < 100 and ln == ln.upper() and "?" not in ln and not re.match(r"^\d+[\).]", ln):
            current_topic = ln
            i += 1
            continue

        is_question = ("?" in ln) or bool(re.match(r"^\d+[\).]\s*", ln))
        if not is_question:
            i += 1
            continue
        q_text = re.sub(r"^\d+[\).]\s*", "", ln).strip()
        opts = {"A": "", "B": "", "C": "", "D": ""}
        correct_option = ""
        answer_text = ""
        j = i + 1
        while j < len(lines):
            nxt = lines[j]
            m = re.match(r"^([A-Da-d])[\).:\-]\s*(.+)$", nxt)
            if m:
                opts[m.group(1).upper()] = m.group(2).strip()
                j += 1
                continue
            ans = re.match(r"^(answer|correct)\s*[:\-]\s*([A-Da-d])", nxt, flags=re.I)
            if ans:
                correct_option = ans.group(2).upper()
                answer_text = opts.get(correct_option, "")
                j += 1
                continue
            if "?" in nxt or re.match(r"^\d+[\).]\s*", nxt):
                break
            j += 1
        rows.append(
            {
                "source": "",
                "year_level": "",
                "trade": "",
                "topic": current_topic,
                "question_text": q_text,
                "question_image": image_refs[min(image_idx, len(image_refs) - 1)],
                "option_a": opts["A"],
                "option_b": opts["B"],
                "option_c": opts["C"],
                "option_d": opts["D"],
                "correct_option": correct_option,
                "answer_text": answer_text,
            }
        )
        image_idx += 1
        i = j
    return pd.DataFrame(rows)


def parse_question_rows_from_sections(sections: list[tuple[str, str]]) -> pd.DataFrame:
    frames: list[pd.DataFrame] = []
    for text, image_ref in sections:
        frame = parse_question_rows_from_text(text, image_refs=[image_ref] if image_ref else None)
        if not frame.empty:
            frames.append(frame)
    if not frames:
        return empty_question_bank_df()
    return pd.concat(frames, ignore_index=True)


def build_sample_feedback_payloads() -> list[dict[str, Any]]:
    months = ["2026-01", "2026-02", "2026-03", "2026-04", "2026-05", "2026-06"]
    districts = ["Vijayawada", "Guntur", "Visakhapatnam"]
    trades = ["Electrician", "Fitter", "Welder"]
    topic_sets = [
        (
            "Trade Theory",
            "Ohm's Law",
            "How well did you understand the concept of Ohm's Law in Trade Theory?",
            "Trade Practical",
            "Meters and Measurement",
            "How well were you able to perform the practical activity for Meters and Measurement in Trade Practical?",
            "Workshop Calculation",
            "Series and Parallel Circuits",
            "How well were you able to solve the calculation part of Series and Parallel Circuits in Workshop Calculation?",
        ),
        (
            "Trade Theory",
            "Arc Welding Basics",
            "How well did you understand Arc Welding Basics in Trade Theory?",
            "Trade Practical",
            "Electrode Handling",
            "How confident are you in performing Electrode Handling during practical sessions?",
            "Workshop Calculation",
            "Material Estimation",
            "How well were you able to perform Material Estimation calculations?",
        ),
    ]
    comments = [
        "Training was clear and practical sessions were helpful.",
        "Sessions were useful but I need more hands-on practice.",
        "Good teaching, but tools were sometimes limited.",
        "I understood most topics and want more revision classes.",
    ]

    payloads: list[dict[str, Any]] = []
    for month in months:
        for district in districts:
            trade = random.choice(trades)
            t = random.choice(topic_sets)
            topic_1_score = round(random.uniform(2.0, 5.0), 1)
            topic_2_score = round(random.uniform(2.0, 5.0), 1)
            topic_3_score = round(random.uniform(2.0, 5.0), 1)
            payloads.append(
                {
                    "month": month,
                    "district": district,
                    "trade_name": trade,
                    "year": random.choice([1, 2]),
                    "semester": random.choice([1, 2, 3, 4]),
                    "attendance_pct": round(random.uniform(65.0, 96.0), 1),
                    "subject_1": t[0],
                    "topic_1": t[1],
                    "question_1": t[2],
                    "topic_1_score": topic_1_score,
                    "subject_2": t[3],
                    "topic_2": t[4],
                    "question_2": t[5],
                    "topic_2_score": topic_2_score,
                    "subject_3": t[6],
                    "topic_3": t[7],
                    "question_3": t[8],
                    "topic_3_score": topic_3_score,
                    "teaching_score": round(random.uniform(2.0, 5.0), 1),
                    "practical_score": round(random.uniform(2.0, 5.0), 1),
                    "learning_score": round(random.uniform(2.0, 5.0), 1),
                    "support_score": round(random.uniform(2.0, 5.0), 1),
                    "safety_score": round(random.uniform(2.0, 5.0), 1),
                    "comment_text": random.choice(comments),
                }
            )
    return payloads


@st.cache_data(show_spinner=False, ttl=300)
def build_demo_preview_mirrored_data() -> dict[str, pd.DataFrame]:
    rng = random.Random(17)
    districts = ["Vijayawada", "Guntur", "Visakhapatnam", "Kurnool"]
    trades = ["Electrician", "Fitter", "Welder", "Mechanic Diesel"]
    role_forms = {
        "Trainer": "iti_training_officer_feedback_form",
        "OJT Trainer": "institute_feedback_form_in_plant_training_ato_dto_to",
        "Supervisor": "supervisor_feedback_form_iti_in_plant_training_evaluation",
        "Trainee": "iti_trainee_feedback_form",
    }
    role_titles = {
        "Trainer": "ITI Training Officer Feedback Form",
        "OJT Trainer": "Institute Feedback Form In-Plant Training ATO/DTO/TO",
        "Supervisor": "Supervisor Feedback Form ITI In-Plant Training Evaluation",
        "Trainee": "ITI Trainee Feedback Form",
    }
    role_people = {
        "Trainer": ["Rakesh Kumar", "Priya Nair", "Sandeep Rao", "Lavanya Devi"],
        "OJT Trainer": ["A. Kumar", "Deepa Sharma", "Rohit Varma", "Kiran Patel"],
        "Supervisor": ["Mohan Das", "Ritu Singh", "Farooq Ali", "Sneha Joseph"],
        "Trainee": ["Arjun", "Keerthi", "Bhanu", "Naveen", "Sowmya", "Goutham"],
    }
    monthly_rows: list[dict[str, Any]] = []
    category_rows: list[dict[str, Any]] = []
    now = datetime.now()
    for day_offset in range(48):
        trade = trades[day_offset % len(trades)]
        district = districts[day_offset % len(districts)]
        ts = now - timedelta(days=day_offset, hours=rng.randint(0, 11), minutes=rng.randint(0, 55))
        month = ts.strftime("%Y-%m")
        year_level = 1 if day_offset % 2 == 0 else 2
        semester = (day_offset % 4) + 1
        monthly_rows.append(
            {
                "month": month,
                "district": district,
                "trade_name": trade,
                "year": year_level,
                "semester": semester,
                "attendance_pct": round(rng.uniform(62, 95), 1),
                "teaching_score": round(rng.uniform(2.1, 4.8), 1),
                "practical_score": round(rng.uniform(1.8, 4.7), 1),
                "learning_score": round(rng.uniform(2.0, 4.9), 1),
                "support_score": round(rng.uniform(2.2, 4.8), 1),
                "safety_score": round(rng.uniform(2.4, 4.9), 1),
                "sentiment_label": rng.choices(["positive", "neutral", "negative"], weights=[0.5, 0.3, 0.2])[0],
                "submitted_at": pd.NaT,
            }
        )
        for role_label in role_forms:
            if role_label == "Trainee" and day_offset % 3 == 0:
                continue
            avg_rating = round(rng.uniform(1.7, 3.9), 2)
            person_name = rng.choice(role_people[role_label])
            comment = rng.choice(
                [
                    "Strong delivery and timely response.",
                    "Needs better follow-up on practical execution.",
                    "Good coordination with the department.",
                    "Learners asked for more revision support.",
                ]
            )
            category_rows.append(
                {
                    "submitted_at": pd.Timestamp(ts, tz="UTC"),
                    "source": "category",
                    "form_id": role_forms[role_label],
                    "form_title": role_titles[role_label],
                    "basic_details_json": json.dumps(
                        {
                            "Name": person_name,
                            "District": district,
                            "Trade": trade,
                            "Year": str(year_level),
                            "Semester": str(semester),
                        }
                    ),
                    "parameter_scores_json": "[]",
                    "comment_text": comment,
                    "excellent_count": int(max(avg_rating - 2.5, 0) * 2),
                    "good_count": int(max(avg_rating - 1.8, 0) * 2),
                    "average_count": 1 if 2.0 <= avg_rating < 3.0 else 0,
                    "poor_count": 1 if avg_rating < 2.2 else 0,
                    "avg_rating_score": avg_rating,
                    "basic_details": {
                        "Name": person_name,
                        "District": district,
                        "Trade": trade,
                        "Year": str(year_level),
                        "Semester": str(semester),
                    },
                    "trade_name": trade,
                    "district": district,
                    "year": year_level,
                    "semester": semester,
                    "month": month,
                    "role_group": role_label,
                }
            )
    monthly_df = pd.DataFrame(monthly_rows)
    category_df = pd.DataFrame(category_rows)
    return {"monthly": monthly_df, "category": category_df}


def render_interactive_trend_chart(
    chart_df: pd.DataFrame,
    group_col: str,
    metric_col: str,
    metric_label: str,
    chart_mode: str,
) -> None:
    plot_df = chart_df.reset_index().copy()
    plot_df[group_col] = plot_df[group_col].astype(str)

    # Area charts need at least 2 x-points to produce visible fill.
    # If only one month is available, synthesize a previous month with same value.
    if chart_mode == "Area" and group_col == "month" and len(plot_df) == 1:
        try:
            curr = pd.to_datetime(plot_df.iloc[0][group_col], format="%Y-%m", errors="coerce")
            if pd.notna(curr):
                prev = curr - pd.DateOffset(months=1)
                synthetic = plot_df.iloc[0].copy()
                synthetic[group_col] = prev.strftime("%Y-%m")
                plot_df = pd.concat([pd.DataFrame([synthetic]), plot_df], ignore_index=True)
        except Exception:
            pass

    if group_col == "month":
        plot_df["x_sort"] = pd.to_datetime(plot_df[group_col], format="%Y-%m", errors="coerce")
    else:
        plot_df["x_sort"] = range(len(plot_df))

    base = (
        alt.Chart(plot_df)
        .encode(
            x=alt.X(
                f"{group_col}:N",
                sort=alt.SortField(field="x_sort", order="ascending"),
                axis=alt.Axis(title=group_col.replace("_", " ").title(), labelAngle=-45),
            ),
            y=alt.Y(f"{metric_col}:Q", title=metric_label),
            tooltip=[
                alt.Tooltip(f"{group_col}:N", title=group_col.replace("_", " ").title()),
                alt.Tooltip(f"{metric_col}:Q", title=metric_label, format=".2f"),
            ],
        )
    )

    if chart_mode == "Bar":
        chart = base.mark_bar(cornerRadiusTopLeft=6, cornerRadiusTopRight=6).encode(
            color=alt.Color(f"{group_col}:N", legend=None)
        )
    elif chart_mode == "Area":
        area = base.mark_area(opacity=0.30, color="#264b63")
        area_line = base.mark_line(strokeWidth=3, color="#66c2ff")
        area_points = base.mark_circle(size=85, color="#ff7b3a", stroke="#ffffff", strokeWidth=1.2)
        chart = area + area_line + area_points
    else:
        line = base.mark_line(strokeWidth=3, color="#66c2ff")
        points = base.mark_circle(size=90, color="#ff7b3a", stroke="#ffffff", strokeWidth=1.2)
        if hasattr(alt, "selection_point"):
            hover_sel = alt.selection_point(on="mouseover", nearest=True, fields=[group_col], empty=False)
            highlight = base.mark_circle(size=180, color="#ffd166").encode(
                opacity=alt.condition(hover_sel, alt.value(1), alt.value(0))
            )
            chart = line + points + highlight
            chart = chart.add_params(hover_sel)
        else:
            hover_sel = alt.selection_single(on="mouseover", nearest=True, fields=[group_col], empty="none")
            highlight = base.mark_circle(size=180, color="#ffd166").encode(
                opacity=alt.condition(hover_sel, alt.value(1), alt.value(0))
            )
            chart = (line + points + highlight).add_selection(hover_sel)

    st.altair_chart(chart.properties(height=360), use_container_width=True)


def render_quality_chart(teaching: float, practical: float, learning: float, support: float, safety: float) -> None:
    qdf = pd.DataFrame(
        [
            {"Dimension": "Teaching", "Score": teaching},
            {"Dimension": "Practical", "Score": practical},
            {"Dimension": "Learning", "Score": learning},
            {"Dimension": "Support", "Score": support},
            {"Dimension": "Safety", "Score": safety},
        ]
    )
    chart = (
        alt.Chart(qdf)
        .mark_bar(cornerRadiusTopLeft=6, cornerRadiusTopRight=6)
        .encode(
            x=alt.X("Dimension:N", sort=["Teaching", "Practical", "Learning", "Support", "Safety"]),
            y=alt.Y("Score:Q", scale=alt.Scale(domain=[0, 5])),
            color=alt.Color(
                "Dimension:N",
                scale=alt.Scale(
                    domain=["Teaching", "Practical", "Learning", "Support", "Safety"],
                    range=["#52b3ff", "#ff7b3a", "#82e0aa", "#f7dc6f", "#c39bd3"],
                ),
                legend=None,
            ),
            tooltip=[
                alt.Tooltip("Dimension:N", title="Dimension"),
                alt.Tooltip("Score:Q", title="Score", format=".2f"),
            ],
        )
        .properties(height=260)
    )
    st.altair_chart(chart, use_container_width=True)


def render_sentiment_chart(positive_count: int, neutral_count: int, negative_count: int) -> None:
    sdf = pd.DataFrame(
        [
            {"Sentiment": "Positive", "Count": positive_count},
            {"Sentiment": "Neutral", "Count": neutral_count},
            {"Sentiment": "Negative", "Count": negative_count},
        ]
    )
    chart = (
        alt.Chart(sdf)
        .mark_bar(cornerRadiusTopLeft=8, cornerRadiusTopRight=8)
        .encode(
            x=alt.X("Sentiment:N", sort=["Positive", "Neutral", "Negative"]),
            y=alt.Y("Count:Q"),
            color=alt.Color(
                "Sentiment:N",
                scale=alt.Scale(
                    domain=["Positive", "Neutral", "Negative"],
                    range=["#2ecc71", "#f1c40f", "#e74c3c"],
                ),
                legend=None,
            ),
            tooltip=[
                alt.Tooltip("Sentiment:N", title="Sentiment"),
                alt.Tooltip("Count:Q", title="Count"),
            ],
        )
        .properties(height=240)
    )
    st.altair_chart(chart, use_container_width=True)


def render_category_rating_chart(excellent: int, good: int, average: int, poor: int) -> None:
    rdf = pd.DataFrame(
        [
            {"Rating": "Excellent", "Count": excellent},
            {"Rating": "Good", "Count": good},
            {"Rating": "Average", "Count": average},
            {"Rating": "Poor", "Count": poor},
        ]
    )
    chart = (
        alt.Chart(rdf)
        .mark_bar(cornerRadiusTopLeft=8, cornerRadiusTopRight=8)
        .encode(
            x=alt.X("Rating:N", sort=["Excellent", "Good", "Average", "Poor"]),
            y=alt.Y("Count:Q"),
            color=alt.Color(
                "Rating:N",
                scale=alt.Scale(
                    domain=["Excellent", "Good", "Average", "Poor"],
                    range=["#2ecc71", "#5dade2", "#f1c40f", "#e74c3c"],
                ),
                legend=None,
            ),
            tooltip=[alt.Tooltip("Rating:N"), alt.Tooltip("Count:Q")],
        )
        .properties(height=260)
    )
    st.altair_chart(chart, use_container_width=True)


def render_donut_chart(df: pd.DataFrame, category_col: str, value_col: str, colors: list[str], title: str) -> None:
    if df.empty or category_col not in df.columns or value_col not in df.columns:
        return
    chart = (
        alt.Chart(df)
        .mark_arc(innerRadius=70, outerRadius=120)
        .encode(
            theta=alt.Theta(f"{value_col}:Q"),
            color=alt.Color(
                f"{category_col}:N",
                scale=alt.Scale(range=colors),
                legend=alt.Legend(title=title),
            ),
            tooltip=[alt.Tooltip(f"{category_col}:N"), alt.Tooltip(f"{value_col}:Q")],
        )
        .properties(height=280)
    )
    st.altair_chart(chart, use_container_width=True)


def render_horizontal_accuracy_chart(df: pd.DataFrame, category_col: str = "topic", value_col: str = "accuracy_pct") -> None:
    if df.empty or category_col not in df.columns or value_col not in df.columns:
        return
    sdf = df.copy()
    chart = (
        alt.Chart(sdf)
        .mark_bar(cornerRadius=6)
        .encode(
            y=alt.Y(f"{category_col}:N", sort="-x", title=category_col.replace("_", " ").title()),
            x=alt.X(f"{value_col}:Q", title="Accuracy %", scale=alt.Scale(domain=[0, 100])),
            color=alt.Color(f"{value_col}:Q", scale=alt.Scale(scheme="redyellowgreen"), legend=None),
            tooltip=[alt.Tooltip(f"{category_col}:N"), alt.Tooltip(f"{value_col}:Q", format=".2f")],
        )
        .properties(height=320)
    )
    st.altair_chart(chart, use_container_width=True)


def render_bubble_trend(df: pd.DataFrame, group_col: str) -> None:
    if df.empty:
        return
    needed = {"technical_accuracy_pct", "total_questions_answered", "technical_submissions", group_col}
    if not needed.issubset(set(df.columns)):
        return
    plot_df = df.copy()
    plot_df[group_col] = plot_df[group_col].astype(str)

    # Bubble charts look poor for very small datasets; switch to cleaner combo view.
    if len(plot_df) <= 3:
        base = alt.Chart(plot_df).encode(
            x=alt.X(f"{group_col}:N", title=group_col.replace("_", " ").title(), sort=None),
            tooltip=[
                alt.Tooltip(f"{group_col}:N"),
                alt.Tooltip("total_questions_answered:Q", title="Questions"),
                alt.Tooltip("technical_submissions:Q", title="Submissions"),
                alt.Tooltip("technical_accuracy_pct:Q", title="Accuracy %", format=".2f"),
            ],
        )
        bars = base.mark_bar(opacity=0.45, color="#4dabf7").encode(
            y=alt.Y("total_questions_answered:Q", title="Questions Answered")
        )
        line = base.mark_line(color="#ffb74d", point=alt.OverlayMarkDef(size=80, filled=True)).encode(
            y=alt.Y("technical_accuracy_pct:Q", title="Technical Accuracy %")
        )
        chart = alt.layer(bars, line).resolve_scale(y="independent").properties(height=340)
        st.altair_chart(chart, use_container_width=True)
        return

    chart = (
        alt.Chart(plot_df)
        .mark_circle(opacity=0.8, stroke="white", strokeWidth=1)
        .encode(
            x=alt.X("total_questions_answered:Q", title="Questions Answered"),
            y=alt.Y("technical_accuracy_pct:Q", title="Technical Accuracy %", scale=alt.Scale(domain=[0, 100])),
            size=alt.Size(
                "technical_submissions:Q",
                title="Technical Submissions",
                scale=alt.Scale(range=[120, 1200]),
            ),
            color=alt.Color(f"{group_col}:N", title=group_col.replace("_", " ").title()),
            tooltip=[
                alt.Tooltip(f"{group_col}:N"),
                alt.Tooltip("technical_submissions:Q"),
                alt.Tooltip("total_questions_answered:Q"),
                alt.Tooltip("technical_accuracy_pct:Q", format=".2f"),
            ],
        )
        .properties(height=340)
    )
    st.altair_chart(chart, use_container_width=True)


ROLE_GROUP_FORM_IDS = {
    "Trainer": ["iti_training_officer_feedback_form"],
    "Trainee": ["iti_trainee_feedback_form", "student_feedback_form_in_plant_training"],
    "Principal": ["principal_feedback_form"],
    "OJT Trainer": ["institute_feedback_form_in_plant_training_ato_dto_to"],
    "Supervisor": ["supervisor_feedback_form_iti_in_plant_training_evaluation"],
}

DASHBOARD_VIEW_ROLE_LABELS = {
    "principal": ["Trainer", "OJT Trainer", "Supervisor"],
    "trainer": ["Trainer", "Trainee"],
    "default": ["Trainer", "Trainee", "Supervisor"],
}

ROLE_ACCENT_COLORS = {
    "Trainer": "#38bdf8",
    "Trainee": "#8b5cf6",
    "Principal": "#f97316",
    "OJT Trainer": "#f59e0b",
    "Supervisor": "#22c55e",
}


def get_dashboard_role_labels(viewer_role: str) -> list[str]:
    role_key = value_or_empty(viewer_role).strip().lower()
    return DASHBOARD_VIEW_ROLE_LABELS.get(role_key, DASHBOARD_VIEW_ROLE_LABELS["default"]).copy()


def get_dashboard_role_form_ids(viewer_role: str) -> dict[str, list[str]]:
    return {label: ROLE_GROUP_FORM_IDS.get(label, []).copy() for label in get_dashboard_role_labels(viewer_role)}


def normalize_sentiment_label(value: Any) -> str:
    label = value_or_empty(value).strip().lower()
    mapping = {
        "good": "positive",
        "bad": "negative",
        "average": "neutral",
        "mixed": "neutral",
    }
    return mapping.get(label, label or "neutral")


def sentiment_from_rating(value: Any) -> str:
    score = safe_float(value)
    if score >= 3.0:
        return "positive"
    if score >= 2.0:
        return "neutral"
    return "negative"


def status_from_rating(value: Any) -> str:
    score = safe_float(value)
    if score >= 3.2:
        return "Healthy"
    if score >= 2.5:
        return "Watch"
    return "Alert"


def status_from_performance(avg_rating: Any, submissions: Any) -> str:
    submission_count = int(safe_float(submissions))
    score = safe_float(avg_rating)
    if submission_count <= 0:
        return "No Data"
    if score >= 3.5:
        return "Excellent"
    if score >= 2.5:
        return "Good"
    return "Needs Attention"


def safe_count(df: pd.DataFrame | None) -> int:
    return int(len(df)) if isinstance(df, pd.DataFrame) else 0


def safe_avg(values: Any, default: float = 0.0) -> float:
    if isinstance(values, pd.DataFrame):
        if values.empty:
            return default
        series = pd.to_numeric(values.stack(), errors="coerce")
    elif isinstance(values, pd.Series):
        series = pd.to_numeric(values, errors="coerce")
    elif isinstance(values, list | tuple | set):
        series = pd.to_numeric(pd.Series(list(values), dtype="object"), errors="coerce")
    else:
        series = pd.to_numeric(pd.Series([values], dtype="object"), errors="coerce")
    series = series.dropna()
    if series.empty:
        return default
    return round(float(series.mean()), 2)


def safe_delta(series: list[float] | None) -> tuple[str, str]:
    cleaned = [safe_float(v) for v in (series or []) if v is not None]
    if len(cleaned) < 2:
        return ("No prior data", "flat")
    previous = safe_float(cleaned[-2])
    current = safe_float(cleaned[-1])
    delta = current - previous
    if abs(delta) < 0.005:
        return ("No change", "flat")
    arrow = "UP" if delta > 0 else "DOWN"
    return (f"{arrow} {abs(delta):.2f}", ("up" if delta > 0 else "down"))


def monthly_score_columns() -> list[str]:
    return ["teaching_score", "practical_score", "learning_score", "support_score", "safety_score"]


def parse_json_dict(value: Any) -> dict[str, Any]:
    try:
        parsed = json.loads(value_or_empty(value))
        if isinstance(parsed, dict):
            return parsed
    except Exception:
        pass
    return {}


def parse_iso_dt(value: Any) -> datetime | None:
    text = value_or_empty(value).strip()
    if not text:
        return None
    try:
        return datetime.fromisoformat(text.replace("Z", "+00:00"))
    except Exception:
        try:
            return pd.to_datetime(text, errors="coerce").to_pydatetime()
        except Exception:
            return None


def infer_role_from_form_id(form_id: str) -> str:
    form = value_or_empty(form_id).strip().lower()
    for role_label, form_ids in ROLE_GROUP_FORM_IDS.items():
        if form in {f.lower() for f in form_ids}:
            return role_label
    if "supervisor" in form:
        return "Supervisor"
    if "institute" in form or "officer" in form:
        return "OJT Trainer"
    if "principal" in form:
        return "Principal"
    if "trainer" in form:
        return "Trainer"
    return "Trainee"


def extract_basic_detail(details: dict[str, Any], *needles: str) -> str:
    lowered = {str(k).strip().lower(): value_or_empty(v).strip() for k, v in details.items()}
    for needle in needles:
        for key, value in lowered.items():
            if needle in key and value:
                return value
    return ""


def first_non_empty(*values: Any, default: str = "") -> str:
    for value in values:
        text = value_or_empty(value).strip()
        if text:
            return text
    return default


def infer_feedback_name(details: dict[str, Any], role_label: str, row_index: int) -> str:
    return first_non_empty(
        extract_basic_detail(details, "name"),
        extract_basic_detail(details, "trainee"),
        extract_basic_detail(details, "student"),
        extract_basic_detail(details, "trainer"),
        extract_basic_detail(details, "supervisor"),
        default=f"{role_label} #{row_index}",
    )


def coerce_datetime(value: Any) -> datetime | None:
    if value is None:
        return None
    try:
        if pd.isna(value):
            return None
    except Exception:
        pass
    if isinstance(value, pd.Timestamp):
        if pd.isna(value):
            return None
        return value.to_pydatetime()
    if isinstance(value, datetime):
        return value
    try:
        parsed = pd.to_datetime(value, errors="coerce")
        if pd.isna(parsed):
            return None
        if isinstance(parsed, pd.Timestamp):
            return parsed.to_pydatetime()
    except Exception:
        return None
    return None


def format_relative_time(dt: Any) -> str:
    dt = coerce_datetime(dt)
    if dt is None:
        return "Unavailable"
    now = datetime.now(dt.tzinfo) if dt.tzinfo else datetime.now()
    delta = now - dt
    total_seconds_float = delta.total_seconds()
    if pd.isna(total_seconds_float):
        return "Unavailable"
    total_seconds = int(max(total_seconds_float, 0))
    if total_seconds < 60:
        return f"{total_seconds}s ago"
    if total_seconds < 3600:
        return f"{total_seconds // 60}m ago"
    if total_seconds < 86400:
        return f"{total_seconds // 3600}h ago"
    return f"{total_seconds // 86400}d ago"


def build_sparkline_svg(values: list[float], color: str) -> str:
    points = [safe_float(v) for v in values if v is not None]
    if not points:
        points = [0.0, 0.0]
    if len(points) == 1:
        points = [points[0], points[0]]
    minimum = min(points)
    maximum = max(points)
    spread = maximum - minimum or 1.0
    coords: list[str] = []
    for idx, value in enumerate(points):
        x = 4 + (idx * 92 / max(len(points) - 1, 1))
        y = 30 - (((value - minimum) / spread) * 22)
        coords.append(f"{x:.1f},{y:.1f}")
    return (
        "<svg viewBox='0 0 100 34' preserveAspectRatio='none'>"
        f"<polyline fill='none' stroke='{color}' stroke-width='2.5' stroke-linecap='round' stroke-linejoin='round' points='{' '.join(coords)}'/>"
        "</svg>"
    )


def metric_delta(current: float, previous: float) -> tuple[str, str]:
    delta = current - previous
    arrow = "UP" if delta >= 0 else "DOWN"
    return f"{arrow} {abs(delta):.2f}", ("up" if delta >= 0 else "down")


def series_with_fallback(values: list[float] | None, fallback: float = 0.0) -> list[float]:
    cleaned = [safe_float(v) for v in (values or [])]
    return cleaned if cleaned else [safe_float(fallback)]


def set_auth_session(login_data: dict[str, Any], password: str) -> None:
    st.session_state.auth_user = {
        "user_id": value_or_empty(login_data.get("user_id")),
        "username": value_or_empty(login_data.get("username")),
        "role": value_or_empty(login_data.get("role")).lower(),
        "full_name": value_or_empty(login_data.get("full_name")),
        "email": value_or_empty(login_data.get("email")),
        "assigned_trade": value_or_empty(login_data.get("assigned_trade")),
        "assigned_year": login_data.get("assigned_year"),
        "semester": value_or_empty(login_data.get("semester")),
        "district": value_or_empty(login_data.get("district")),
        "department": value_or_empty(login_data.get("department")),
        "status": value_or_empty(login_data.get("status")),
        "access_token": value_or_empty(login_data.get("access_token")),
        "refresh_token": value_or_empty(login_data.get("refresh_token")),
    }
    st.session_state.auth_password = password
    st.session_state.auth_session = {
        "is_authenticated": True,
        "logged_in_at": datetime.now().isoformat(),
        "last_seen_at": datetime.now().isoformat(),
        "username": value_or_empty(login_data.get("username")),
        "role": value_or_empty(login_data.get("role")).lower(),
        "persistent_session_id": value_or_empty(st.session_state.get("auth_session", {}).get("persistent_session_id")),
        "expires_at": value_or_empty(st.session_state.get("auth_session", {}).get("expires_at")),
    }


def clear_auth_session() -> None:
    st.session_state.auth_user = None
    st.session_state.auth_password = ""
    st.session_state.auth_session = {
        "is_authenticated": False,
        "logged_in_at": None,
        "last_seen_at": None,
        "username": "",
        "role": "",
        "persistent_session_id": "",
        "expires_at": None,
    }
    st.session_state.db_student_qset = None


def load_auth_session_store() -> dict[str, Any]:
    if not AUTH_SESSION_STORE_PATH.exists():
        return {}
    try:
        data = json.loads(AUTH_SESSION_STORE_PATH.read_text(encoding="utf-8"))
        return data if isinstance(data, dict) else {}
    except Exception:
        return {}


def save_auth_session_store(store: dict[str, Any]) -> None:
    AUTH_SESSION_STORE_PATH.parent.mkdir(parents=True, exist_ok=True)
    AUTH_SESSION_STORE_PATH.write_text(json.dumps(store, indent=2), encoding="utf-8")


def cleanup_auth_session_store(store: dict[str, Any] | None = None) -> dict[str, Any]:
    active_store = dict(store or load_auth_session_store())
    now = datetime.now()
    cleaned: dict[str, Any] = {}
    for session_id, payload in active_store.items():
        expires_at = coerce_datetime((payload or {}).get("expires_at"))
        if expires_at and expires_at > now:
            cleaned[session_id] = payload
    if cleaned != active_store:
        save_auth_session_store(cleaned)
    return cleaned


def get_auth_query_session_id() -> str:
    try:
        return value_or_empty(st.query_params.get("auth_sid", "")).strip()
    except Exception:
        return ""


def set_auth_query_session_id(session_id: str) -> None:
    try:
        if session_id:
            st.query_params["auth_sid"] = session_id
        elif "auth_sid" in st.query_params:
            del st.query_params["auth_sid"]
    except Exception:
        pass


def create_persistent_auth_session(login_data: dict[str, Any], password: str) -> str:
    session_id = secrets.token_urlsafe(24)
    store = cleanup_auth_session_store()
    now = datetime.now()
    expires_at = now + timedelta(hours=AUTH_SESSION_TTL_HOURS)
    store[session_id] = {
        "user_id": value_or_empty(login_data.get("user_id")),
        "username": value_or_empty(login_data.get("username")),
        "role": value_or_empty(login_data.get("role")).lower(),
        "full_name": value_or_empty(login_data.get("full_name")),
        "email": value_or_empty(login_data.get("email")),
        "assigned_trade": value_or_empty(login_data.get("assigned_trade")),
        "assigned_year": login_data.get("assigned_year"),
        "semester": value_or_empty(login_data.get("semester")),
        "district": value_or_empty(login_data.get("district")),
        "department": value_or_empty(login_data.get("department")),
        "status": value_or_empty(login_data.get("status")),
        "access_token": value_or_empty(login_data.get("access_token")),
        "refresh_token": value_or_empty(login_data.get("refresh_token")),
        "password": password,
        "created_at": now.isoformat(),
        "expires_at": expires_at.isoformat(),
    }
    save_auth_session_store(store)
    return session_id


def restore_persistent_auth_session(session_id: str) -> bool:
    if not session_id:
        return False
    store = cleanup_auth_session_store()
    payload = store.get(session_id)
    if not isinstance(payload, dict):
        return False
    set_auth_session(payload, value_or_empty(payload.get("password")))
    st.session_state.auth_session["persistent_session_id"] = session_id
    st.session_state.auth_session["expires_at"] = value_or_empty(payload.get("expires_at"))
    return True


def remove_persistent_auth_session(session_id: str) -> None:
    if not session_id:
        return
    store = load_auth_session_store()
    if session_id in store:
        del store[session_id]
        save_auth_session_store(store)


def inject_dashboard_styles() -> None:
    st.markdown(
        """
        <style>
        @import url('https://fonts.googleapis.com/css2?family=Manrope:wght@400;500;600;700;800&display=swap');

        :root {
            --bg-0: #07111f;
            --bg-1: #0d1728;
            --bg-2: rgba(18, 30, 48, 0.88);
            --stroke: rgba(148, 163, 184, 0.18);
            --text-0: #f8fafc;
            --text-1: #cbd5e1;
            --text-2: #7dd3fc;
            --good: #22c55e;
            --warn: #f59e0b;
            --bad: #ef4444;
            --accent: #38bdf8;
            --accent-2: #8b5cf6;
        }

        html, body, [class*="css"]  {
            font-family: 'Manrope', sans-serif;
        }

        .stApp {
            background:
                radial-gradient(circle at top left, rgba(56, 189, 248, 0.18), transparent 24%),
                radial-gradient(circle at top right, rgba(139, 92, 246, 0.15), transparent 22%),
                linear-gradient(180deg, #06101c 0%, #0b1424 52%, #08111d 100%);
            color: var(--text-0);
        }

        .main .block-container {
            padding-top: 2.25rem;
            padding-bottom: 2rem;
        }

        section[data-testid="stSidebar"] {
            background: linear-gradient(180deg, rgba(6, 16, 28, 0.96), rgba(10, 22, 36, 0.96));
            border-right: 1px solid var(--stroke);
        }

        div[data-testid="stMetric"] {
            background: transparent;
            border: none;
            padding: 0;
        }

        .dashboard-hero, .dashboard-card, .kpi-card, .status-card, .feed-card {
            background: linear-gradient(180deg, rgba(15, 23, 42, 0.82), rgba(15, 23, 42, 0.62));
            border: 1px solid var(--stroke);
            box-shadow: 0 18px 50px rgba(2, 6, 23, 0.38);
            backdrop-filter: blur(18px);
            -webkit-backdrop-filter: blur(18px);
        }

        .dashboard-hero {
            padding: 24px 28px;
            border-radius: 24px;
            margin-bottom: 18px;
        }

        .dashboard-hero h1 {
            margin: 0;
            font-size: 2rem;
            font-weight: 800;
            letter-spacing: -0.03em;
            color: var(--text-0);
        }

        .dashboard-hero p {
            margin: 8px 0 0;
            color: var(--text-1);
            font-size: 0.98rem;
        }

        .hero-meta {
            display: flex;
            gap: 10px;
            flex-wrap: wrap;
            align-items: center;
            margin-top: 16px;
        }

        .meta-pill, .badge-pill {
            display: inline-flex;
            align-items: center;
            gap: 8px;
            padding: 9px 14px;
            border-radius: 999px;
            background: rgba(15, 23, 42, 0.72);
            border: 1px solid rgba(148, 163, 184, 0.16);
            color: var(--text-1);
            font-size: 0.84rem;
            font-weight: 700;
        }

        .pulse-dot {
            width: 9px;
            height: 9px;
            border-radius: 999px;
            background: #22c55e;
            box-shadow: 0 0 0 rgba(34, 197, 94, 0.8);
            animation: pulse-live 1.4s infinite;
        }

        @keyframes pulse-live {
            0% { box-shadow: 0 0 0 0 rgba(34, 197, 94, 0.7); }
            70% { box-shadow: 0 0 0 12px rgba(34, 197, 94, 0); }
            100% { box-shadow: 0 0 0 0 rgba(34, 197, 94, 0); }
        }

        .kpi-card {
            border-radius: 22px;
            padding: 18px 18px 12px;
            min-height: 168px;
            transition: transform 160ms ease, border-color 160ms ease, box-shadow 160ms ease;
        }

        .kpi-card:hover {
            transform: translateY(-3px);
            border-color: rgba(56, 189, 248, 0.32);
            box-shadow: 0 22px 56px rgba(15, 23, 42, 0.42);
        }

        .kpi-label {
            color: var(--text-1);
            font-size: 0.78rem;
            text-transform: uppercase;
            letter-spacing: 0.08em;
            font-weight: 800;
        }

        .kpi-value {
            font-size: 1.8rem;
            font-weight: 800;
            margin: 10px 0 4px;
            color: var(--text-0);
        }

        .kpi-sub, .kpi-delta {
            font-size: 0.82rem;
            font-weight: 700;
        }

        .kpi-sub {
            color: var(--text-1);
        }

        .kpi-delta.up {
            color: #4ade80;
        }

        .kpi-delta.down {
            color: #f87171;
        }

        .kpi-delta.flat {
            color: #94a3b8;
        }

        .sparkline-wrap {
            height: 38px;
            margin-top: 10px;
            opacity: 0.92;
        }

        .status-indicator {
            width: 10px;
            height: 10px;
            border-radius: 999px;
            display: inline-block;
        }

        .section-label {
            margin: 16px 0 10px;
            color: #e2e8f0;
            font-size: 0.94rem;
            font-weight: 800;
            letter-spacing: 0.02em;
        }

        .feed-card, .status-card, .dashboard-card {
            border-radius: 22px;
            padding: 18px 18px 10px;
            margin-bottom: 10px;
        }

        .panel-title {
            font-size: 0.95rem;
            color: #f8fafc;
            font-weight: 800;
            margin-bottom: 6px;
        }

        .panel-subtitle {
            color: var(--text-1);
            font-size: 0.83rem;
            margin-bottom: 10px;
        }

        .activity-item, .alert-item {
            display: flex;
            align-items: flex-start;
            gap: 12px;
            padding: 12px 0;
            border-bottom: 1px solid rgba(148, 163, 184, 0.12);
        }

        .activity-item:last-child, .alert-item:last-child {
            border-bottom: none;
        }

        .activity-time {
            color: #7dd3fc;
            font-size: 0.78rem;
            font-weight: 800;
            min-width: 74px;
        }

        .activity-text, .alert-text {
            color: #e2e8f0;
            font-size: 0.88rem;
            line-height: 1.45;
        }

        .activity-meta {
            display: flex;
            flex-wrap: wrap;
            gap: 8px;
            margin-top: 8px;
        }

        .activity-badge {
            display: inline-flex;
            align-items: center;
            padding: 4px 10px;
            border-radius: 999px;
            font-size: 0.76rem;
            font-weight: 800;
            border: 1px solid rgba(148, 163, 184, 0.14);
        }

        .activity-badge.role { background: rgba(59, 130, 246, 0.14); color: #93c5fd; }
        .activity-badge.rating { background: rgba(245, 158, 11, 0.14); color: #fcd34d; }
        .activity-badge.positive { background: rgba(34, 197, 94, 0.14); color: #86efac; }
        .activity-badge.neutral { background: rgba(245, 158, 11, 0.14); color: #fcd34d; }
        .activity-badge.negative { background: rgba(239, 68, 68, 0.14); color: #fca5a5; }
        .activity-badge.status-good { background: rgba(34, 197, 94, 0.14); color: #86efac; }
        .activity-badge.status-warn { background: rgba(245, 158, 11, 0.14); color: #fcd34d; }
        .activity-badge.status-bad { background: rgba(239, 68, 68, 0.14); color: #fca5a5; }

        .alert-item.high .alert-dot { background: var(--bad); }
        .alert-item.medium .alert-dot { background: var(--warn); }
        .alert-item.low .alert-dot { background: var(--good); }

        .alert-dot {
            width: 10px;
            height: 10px;
            border-radius: 999px;
            margin-top: 4px;
        }

        .sidebar-section {
            padding: 14px 16px;
            margin: 0 0 12px;
            border-radius: 18px;
            background: rgba(15, 23, 42, 0.52);
            border: 1px solid rgba(148, 163, 184, 0.14);
        }

        .sidebar-title {
            font-size: 0.78rem;
            color: #7dd3fc;
            text-transform: uppercase;
            letter-spacing: 0.1em;
            font-weight: 800;
            margin-bottom: 6px;
        }

        .sidebar-copy {
            color: var(--text-1);
            font-size: 0.84rem;
            line-height: 1.5;
        }

        .skeleton {
            height: 120px;
            border-radius: 20px;
            background: linear-gradient(90deg, rgba(15, 23, 42, 0.45) 20%, rgba(51, 65, 85, 0.45) 38%, rgba(15, 23, 42, 0.45) 60%);
            background-size: 220% 100%;
            animation: shimmer 1.6s infinite linear;
            border: 1px solid rgba(148, 163, 184, 0.14);
        }

        .empty-state {
            border-radius: 24px;
            padding: 28px 24px;
            border: 1px dashed rgba(148, 163, 184, 0.24);
            background: linear-gradient(180deg, rgba(15, 23, 42, 0.68), rgba(15, 23, 42, 0.38));
            margin: 8px 0;
        }

        .empty-icon {
            font-size: 1.35rem;
            font-weight: 800;
            color: var(--text-2);
            margin-bottom: 10px;
        }

        .empty-title {
            font-size: 1rem;
            font-weight: 800;
            color: var(--text-0);
            margin-bottom: 6px;
        }

        .empty-copy {
            color: var(--text-1);
            font-size: 0.92rem;
            line-height: 1.55;
        }

        @keyframes shimmer {
            0% { background-position: 200% 0; }
            100% { background-position: -200% 0; }
        }
        </style>
        """,
        unsafe_allow_html=True,
    )


@st.cache_data(show_spinner=False, ttl=15)
def load_mirrored_feedback_data(base_url: str | None = None) -> dict[str, pd.DataFrame]:
    monthly_path = BASE_DIR / "data" / "monthly_feedback_records.csv"
    category_path = BASE_DIR / "data" / "category_feedback_records.csv"

    monthly_df = pd.DataFrame()
    if monthly_path.exists():
        try:
            monthly_df = pd.read_csv(monthly_path)
        except Exception:
            monthly_df = pd.DataFrame()
    if not monthly_df.empty:
        for col in [
            "teaching_score",
            "practical_score",
            "learning_score",
            "support_score",
            "safety_score",
            "attendance_pct",
            "sentiment_score",
            "year",
            "semester",
        ]:
            monthly_df[col] = pd.to_numeric(monthly_df.get(col), errors="coerce")
        monthly_df["sentiment_label"] = monthly_df.get("sentiment_label", "").apply(normalize_sentiment_label)
        monthly_df["month"] = monthly_df.get("month", "").fillna("").astype(str)
        monthly_df["submitted_at"] = pd.NaT

    category_df = pd.DataFrame()
    if base_url:
        ok_rows, rows_data = api_get(base_url, "/category-feedback/rows")
        if ok_rows and isinstance(rows_data, list) and rows_data:
            category_df = pd.DataFrame(rows_data)
    if category_df.empty and category_path.exists():
        try:
            category_df = pd.read_csv(category_path)
        except Exception:
            category_df = pd.DataFrame()
    if not category_df.empty:
        category_df["submitted_at"] = pd.to_datetime(category_df.get("submitted_at"), errors="coerce", utc=True)
        category_df["avg_rating_score"] = pd.to_numeric(category_df.get("avg_rating_score"), errors="coerce").fillna(0.0)
        if "basic_details" not in category_df.columns:
            category_df["basic_details"] = category_df.get("basic_details_json", "").apply(parse_json_dict)
        else:
            category_df["basic_details"] = category_df["basic_details"].apply(lambda value: value if isinstance(value, dict) else parse_json_dict(value))
        category_df["trade_name"] = category_df["basic_details"].apply(lambda d: extract_basic_detail(d, "trade"))
        category_df["district"] = category_df["basic_details"].apply(lambda d: extract_basic_detail(d, "district"))
        category_df["year"] = pd.to_numeric(
            category_df["basic_details"].apply(lambda d: extract_basic_detail(d, "year")),
            errors="coerce",
        )
        category_df["semester"] = pd.to_numeric(
            category_df["basic_details"].apply(lambda d: extract_basic_detail(d, "semester")),
            errors="coerce",
        )
        category_df["month"] = category_df["submitted_at"].dt.strftime("%Y-%m").fillna("")
        category_df["role_group"] = category_df.get("form_id", "").apply(infer_role_from_form_id)
    return {"monthly": monthly_df, "category": category_df}


def build_filter_options(mirrored: dict[str, pd.DataFrame]) -> dict[str, list[Any]]:
    monthly_df = mirrored.get("monthly", pd.DataFrame())
    category_df = mirrored.get("category", pd.DataFrame())

    month_values = sorted(
        {
            value_or_empty(v).strip()
            for v in pd.concat(
                [monthly_df.get("month", pd.Series(dtype=str)), category_df.get("month", pd.Series(dtype=str))],
                ignore_index=True,
            ).tolist()
            if value_or_empty(v).strip()
        }
    )
    district_values = sorted(
        {
            value_or_empty(v).strip()
            for v in pd.concat(
                [monthly_df.get("district", pd.Series(dtype=str)), category_df.get("district", pd.Series(dtype=str))],
                ignore_index=True,
            ).tolist()
            if value_or_empty(v).strip()
        }
    )
    trade_values = sorted(
        {
            value_or_empty(v).strip()
            for v in pd.concat(
                [monthly_df.get("trade_name", pd.Series(dtype=str)), category_df.get("trade_name", pd.Series(dtype=str))],
                ignore_index=True,
            ).tolist()
            if value_or_empty(v).strip()
        }
    )
    year_values = sorted(
        {
            int(v)
            for v in pd.concat(
                [monthly_df.get("year", pd.Series(dtype=float)), category_df.get("year", pd.Series(dtype=float))],
                ignore_index=True,
            ).dropna().tolist()
            if safe_float(v) > 0
        }
    )
    semester_values = sorted(
        {
            int(v)
            for v in pd.concat(
                [monthly_df.get("semester", pd.Series(dtype=float)), category_df.get("semester", pd.Series(dtype=float))],
                ignore_index=True,
            ).dropna().tolist()
            if safe_float(v) > 0
        }
    )
    role_values = ["All"] + list(ROLE_GROUP_FORM_IDS.keys())
    return {
        "months": ["All"] + month_values,
        "districts": ["All"] + district_values,
        "trades": ["All"] + trade_values,
        "years": ["All"] + [str(v) for v in year_values],
        "semesters": ["All"] + [str(v) for v in semester_values],
        "roles": role_values,
    }


def apply_dashboard_filters(df: pd.DataFrame, filters: dict[str, Any]) -> pd.DataFrame:
    if df.empty:
        return df.copy()
    out = df.copy()
    month = value_or_empty(filters.get("month"))
    district = value_or_empty(filters.get("district"))
    trade_name = value_or_empty(filters.get("trade_name"))
    role_group = value_or_empty(filters.get("role_group"))
    year = value_or_empty(filters.get("year"))
    semester = value_or_empty(filters.get("semester"))

    if month and month != "All" and "month" in out.columns:
        out = out[out["month"].astype(str) == month]
    if district and district != "All" and "district" in out.columns:
        out = out[out["district"].astype(str).str.lower() == district.lower()]
    if trade_name and trade_name != "All" and "trade_name" in out.columns:
        out = out[out["trade_name"].astype(str).str.lower() == trade_name.lower()]
    if role_group and role_group != "All" and "role_group" in out.columns:
        out = out[out["role_group"].astype(str).str.lower() == role_group.lower()]
    if year and year != "All" and "year" in out.columns:
        out = out[pd.to_numeric(out["year"], errors="coerce").fillna(0).astype(int) == int(year)]
    if semester and semester != "All" and "semester" in out.columns:
        out = out[pd.to_numeric(out["semester"], errors="coerce").fillna(0).astype(int) == int(semester)]
    return out


@st.cache_data(show_spinner=False, ttl=30)
def load_technical_submission_data(base_url: str, filters: dict[str, Any]) -> pd.DataFrame:
    params = {
        "month": None if filters.get("month") in {None, "", "All"} else filters.get("month"),
        "district": None if filters.get("district") in {None, "", "All"} else filters.get("district"),
        "trade_name": None if filters.get("trade_name") in {None, "", "All"} else filters.get("trade_name"),
        "year": None if filters.get("year") in {None, "", "All"} else int(filters.get("year")),
        "semester": None if filters.get("semester") in {None, "", "All"} else int(filters.get("semester")),
    }
    ok, data = api_get(base_url, "/authority/technical-rows", params=params)
    if not ok or not isinstance(data, list) or not data:
        return pd.DataFrame()
    tech_df = pd.DataFrame(data)
    if tech_df.empty:
        return tech_df
    tech_df["submitted_at"] = pd.to_datetime(tech_df.get("submitted_at"), errors="coerce")
    tech_df["technical_rating_4"] = pd.to_numeric(tech_df.get("technical_rating_4"), errors="coerce").fillna(0.0)
    tech_df["technical_score_pct"] = pd.to_numeric(tech_df.get("technical_score_pct"), errors="coerce").fillna(0.0)
    tech_df["year"] = pd.to_numeric(tech_df.get("year"), errors="coerce").fillna(0).astype(int)
    tech_df["semester"] = pd.to_numeric(tech_df.get("semester"), errors="coerce").fillna(0).astype(int)
    return tech_df


def build_filtered_dashboard_data(
    base_url: str,
    mirrored: dict[str, pd.DataFrame],
    filters: dict[str, Any],
    viewer_role: str,
    auth_user: dict[str, Any] | None = None,
    use_demo_preview: bool = False,
) -> dict[str, Any]:
    effective_filters = dict(filters)
    if value_or_empty(viewer_role).strip().lower() == "trainer":
        assigned_trade = value_or_empty((auth_user or {}).get("assigned_trade")).strip()
        if assigned_trade:
            effective_filters["trade_name"] = assigned_trade
    monthly_filtered = apply_dashboard_filters(mirrored.get("monthly", pd.DataFrame()), effective_filters)
    category_filtered = apply_dashboard_filters(mirrored.get("category", pd.DataFrame()), effective_filters)
    technical_filtered = load_technical_submission_data(base_url, effective_filters)
    using_demo_preview = False
    if use_demo_preview and monthly_filtered.empty and category_filtered.empty and technical_filtered.empty:
        demo_mirrored = build_demo_preview_mirrored_data()
        monthly_filtered = apply_dashboard_filters(demo_mirrored.get("monthly", pd.DataFrame()), effective_filters)
        category_filtered = apply_dashboard_filters(demo_mirrored.get("category", pd.DataFrame()), effective_filters)
        using_demo_preview = not monthly_filtered.empty or not category_filtered.empty
    dashboard_role_labels = get_dashboard_role_labels(viewer_role)
    return {
        "monthly_filtered": monthly_filtered,
        "category_filtered": category_filtered,
        "technical_filtered": technical_filtered,
        "dashboard_role_labels": dashboard_role_labels,
        "effective_filters": effective_filters,
        "using_demo_preview": using_demo_preview,
    }


def build_principal_scope_data(category_filtered: pd.DataFrame, viewer_role: str, auth_user: dict[str, Any] | None = None) -> dict[str, Any]:
    dashboard_role_labels = get_dashboard_role_labels(viewer_role)
    scoped_category_filtered = filter_category_roles(category_filtered, dashboard_role_labels)
    if value_or_empty(viewer_role).strip().lower() == "trainer" and not scoped_category_filtered.empty:
        username = value_or_empty((auth_user or {}).get("username")).strip().lower()
        assigned_trade = value_or_empty((auth_user or {}).get("assigned_trade")).strip().lower()
        if assigned_trade and "trade_name" in scoped_category_filtered.columns:
            scoped_category_filtered = scoped_category_filtered[
                scoped_category_filtered["trade_name"].astype(str).str.lower() == assigned_trade
            ].copy()
        if username:
            working = scoped_category_filtered.copy()
            working["person_name"] = [
                infer_feedback_name(
                    row.get("basic_details", {}) if isinstance(row.get("basic_details", {}), dict) else {},
                    value_or_empty(row.get("role_group", "Feedback")),
                    idx + 1,
                )
                for idx, (_, row) in enumerate(working.iterrows())
            ]
            trainer_rows = working["role_group"].astype(str).str.lower() == "trainer"
            trainer_name_match = (
                working["person_name"].astype(str).str.lower().str.contains(username, na=False)
                | pd.Series([username in value_or_empty(name).strip().lower() for name in working["person_name"]], index=working.index)
            )
            scoped_category_filtered = working[(~trainer_rows) | trainer_name_match].drop(columns=["person_name"], errors="ignore").copy()
    return {
        "dashboard_role_labels": dashboard_role_labels,
        "scoped_category_filtered": scoped_category_filtered,
        "scoped_category_summary": summarize_category_df(scoped_category_filtered),
        "individual_performance_df": summarize_individual_performance_from_df(scoped_category_filtered, dashboard_role_labels),
        "role_summaries": summarize_role_groups_from_df(scoped_category_filtered, dashboard_role_labels)
        if not scoped_category_filtered.empty
        else {label: {"total_submissions": 0, "avg_rating_score": 0.0, "recent_submissions": []} for label in dashboard_role_labels},
    }


def build_sentiment_summary(detail_df: pd.DataFrame) -> dict[str, Any]:
    if detail_df.empty or "Sentiment" not in detail_df.columns:
        return {
            "total": 0,
            "positive_count": 0,
            "neutral_count": 0,
            "negative_count": 0,
            "positive_pct": 0.0,
            "neutral_pct": 0.0,
            "negative_pct": 0.0,
            "trend_df": pd.DataFrame(columns=["day", "positive", "neutral", "negative"]),
        }
    sentiment_df = detail_df.copy()
    sentiment_df["sentiment_label"] = sentiment_df["Sentiment"].apply(normalize_sentiment_label)
    sentiment_df = sentiment_df[sentiment_df["sentiment_label"].isin(["positive", "neutral", "negative"])].copy()
    if sentiment_df.empty:
        return {
            "total": 0,
            "positive_count": 0,
            "neutral_count": 0,
            "negative_count": 0,
            "positive_pct": 0.0,
            "neutral_pct": 0.0,
            "negative_pct": 0.0,
            "trend_df": pd.DataFrame(columns=["day", "positive", "neutral", "negative"]),
        }
    total = int(len(sentiment_df))
    positive_count = int((sentiment_df["sentiment_label"] == "positive").sum())
    neutral_count = int((sentiment_df["sentiment_label"] == "neutral").sum())
    negative_count = int((sentiment_df["sentiment_label"] == "negative").sum())
    sentiment_df["day"] = pd.to_datetime(sentiment_df.get("_ts"), errors="coerce").dt.strftime("%Y-%m-%d")
    sentiment_df["day"] = sentiment_df["day"].fillna("Unknown")
    trend_grouped = (
        sentiment_df.groupby(["day", "sentiment_label"], dropna=False)
        .size()
        .reset_index(name="count")
    )
    trend_df = trend_grouped.pivot(index="day", columns="sentiment_label", values="count").fillna(0).reset_index()
    for col in ["positive", "neutral", "negative"]:
        if col not in trend_df.columns:
            trend_df[col] = 0
    if "day" in trend_df.columns:
        trend_df["day_sort"] = pd.to_datetime(trend_df["day"], format="%Y-%m-%d", errors="coerce")
        trend_df = trend_df.sort_values(by="day_sort").drop(columns=["day_sort"])
    return {
        "total": total,
        "positive_count": positive_count,
        "neutral_count": neutral_count,
        "negative_count": negative_count,
        "positive_pct": round(pct(positive_count, max(total, 1)), 1) if total else 0.0,
        "neutral_pct": round(pct(neutral_count, max(total, 1)), 1) if total else 0.0,
        "negative_pct": round(pct(negative_count, max(total, 1)), 1) if total else 0.0,
        "trend_df": trend_df,
    }


def build_category_summary(category_filtered: pd.DataFrame) -> dict[str, Any]:
    return summarize_category_df(category_filtered)


def _monthly_record_score_pct(monthly_filtered: pd.DataFrame) -> pd.Series:
    if monthly_filtered.empty:
        return pd.Series(dtype=float)
    cols = [col for col in monthly_score_columns() if col in monthly_filtered.columns]
    if not cols:
        return pd.Series(dtype=float)
    scores = monthly_filtered[cols].apply(pd.to_numeric, errors="coerce")
    return (scores.mean(axis=1).fillna(0.0) / 5.0 * 100.0).round(2)


def build_alert_metrics_from_filtered_df(
    monthly_filtered: pd.DataFrame,
    scoped_category_filtered: pd.DataFrame,
    technical_filtered: pd.DataFrame,
    role_summaries: dict[str, dict[str, Any]],
    dashboard_role_labels: list[str],
    sentiment_summary: dict[str, Any],
    connected: bool,
    errors: list[Any] | None = None,
) -> dict[str, Any]:
    alerts: list[dict[str, str]] = []
    total_records = safe_count(monthly_filtered) + safe_count(scoped_category_filtered) + safe_count(technical_filtered)
    if total_records == 0:
        alerts.append(
            {
                "severity": "low",
                "title": "No alerts available because no records match the selected filters",
                "message": "Try widening month, district, trade, year, semester, or role group.",
            }
        )
        return {"alerts": alerts, "critical_count": 0}
    if not connected:
        alerts.append({"severity": "high", "title": "API disconnected", "message": "Live API is unreachable. Dashboard is using available mirrored records."})
    supervisor_submissions = int(role_summaries.get("Supervisor", {}).get("total_submissions", 0))
    if "Supervisor" in dashboard_role_labels and supervisor_submissions == 0:
        alerts.append({"severity": "high", "title": "Supervisor submissions missing", "message": "No supervisor feedback is visible for the selected scope."})
    avg_attendance = safe_avg(monthly_filtered.get("attendance_pct", pd.Series(dtype=float))) if not monthly_filtered.empty else 0.0
    if safe_count(monthly_filtered) > 0 and avg_attendance < 75:
        alerts.append({"severity": "medium", "title": "Attendance below threshold", "message": f"Average attendance is {avg_attendance:.1f}% for the selected scope."})
    avg_rating = safe_avg(scoped_category_filtered.get("avg_rating_score", pd.Series(dtype=float))) if not scoped_category_filtered.empty else 0.0
    if safe_count(scoped_category_filtered) > 0 and avg_rating < 2.5:
        alerts.append({"severity": "high", "title": "Low feedback rating", "message": f"Average category rating is {avg_rating:.2f}/4 for the selected scope."})
    if sentiment_summary.get("total", 0) > 0 and safe_float(sentiment_summary.get("negative_pct", 0.0)) > 30:
        alerts.append({"severity": "high", "title": "Negative sentiment spike", "message": f"Negative sentiment is {safe_float(sentiment_summary.get('negative_pct', 0.0)):.1f}% of sentiment records."})
    for err in (errors or [])[:2]:
        if value_or_empty(err).strip():
            alerts.append({"severity": "medium", "title": "Recent API issue", "message": value_or_empty(err)[:180]})
    critical_count = len([a for a in alerts if a.get("severity") == "high"])
    return {"alerts": alerts[:5], "critical_count": critical_count}


def build_kpi_metrics_from_filtered_df(
    monthly_filtered: pd.DataFrame,
    scoped_category_filtered: pd.DataFrame,
    technical_filtered: pd.DataFrame,
    role_summaries: dict[str, dict[str, Any]],
    sentiment_summary: dict[str, Any],
    alert_metrics: dict[str, Any],
    dashboard_role_labels: list[str],
    technical_summary: dict[str, Any] | None = None,
) -> dict[str, Any]:
    technical_summary = technical_summary or {}
    technical_submission_count = safe_count(technical_filtered)
    if technical_submission_count == 0:
        technical_submission_count = int(technical_summary.get("technical_submissions", 0) or 0)
    total_records = safe_count(monthly_filtered) + safe_count(scoped_category_filtered) + technical_submission_count
    monthly_pct_series = _monthly_record_score_pct(monthly_filtered)
    category_pct_series = (
        pd.to_numeric(scoped_category_filtered.get("avg_rating_score", pd.Series(dtype=float)), errors="coerce").fillna(0.0) / 4.0 * 100.0
        if not scoped_category_filtered.empty
        else pd.Series(dtype=float)
    )
    technical_pct_series = (
        pd.to_numeric(technical_filtered.get("technical_score_pct", pd.Series(dtype=float)), errors="coerce").fillna(0.0)
        if not technical_filtered.empty
        else pd.Series(([safe_float(technical_summary.get("technical_accuracy_pct", 0.0))] if technical_submission_count > 0 else []), dtype=float)
    )
    overall_avg_pct = safe_avg(pd.concat([monthly_pct_series, category_pct_series, technical_pct_series], ignore_index=True), default=0.0)

    total_series: list[float] = []
    overall_series: list[float] = []
    if total_records > 0:
        month_sources: list[pd.DataFrame] = []
        if not monthly_filtered.empty and "month" in monthly_filtered.columns:
            month_sources.append(monthly_filtered[["month"]].assign(kind_count=1))
        if not scoped_category_filtered.empty and "month" in scoped_category_filtered.columns:
            month_sources.append(scoped_category_filtered[["month"]].assign(kind_count=1))
        if not technical_filtered.empty and "month" in technical_filtered.columns:
            month_sources.append(technical_filtered[["month"]].assign(kind_count=1))
        if month_sources:
            total_series = (
                pd.concat(month_sources, ignore_index=True)
                .groupby("month", dropna=False)["kind_count"]
                .sum()
                .sort_index()
                .tolist()
            )
        if not monthly_filtered.empty:
            monthly_series = (
                monthly_filtered.assign(score_pct=_monthly_record_score_pct(monthly_filtered))
                .groupby("month", dropna=False)["score_pct"]
                .mean()
                .reset_index()
            )
        else:
            monthly_series = pd.DataFrame(columns=["month", "score_pct"])
        if not scoped_category_filtered.empty:
            category_series = (
                scoped_category_filtered.assign(score_pct=(pd.to_numeric(scoped_category_filtered["avg_rating_score"], errors="coerce").fillna(0.0) / 4.0 * 100.0))
                .groupby("month", dropna=False)["score_pct"]
                .mean()
                .reset_index()
            )
        else:
            category_series = pd.DataFrame(columns=["month", "score_pct"])
        if not technical_filtered.empty:
            technical_series_df = (
                technical_filtered.assign(score_pct=pd.to_numeric(technical_filtered["technical_score_pct"], errors="coerce").fillna(0.0))
                .groupby("month", dropna=False)["score_pct"]
                .mean()
                .reset_index()
            )
        elif technical_submission_count > 0:
            technical_series_df = pd.DataFrame(
                [{"month": "Current", "score_pct": safe_float(technical_summary.get("technical_accuracy_pct", 0.0))}]
            )
        else:
            technical_series_df = pd.DataFrame(columns=["month", "score_pct"])
        if not monthly_series.empty or not category_series.empty or not technical_series_df.empty:
            overall_series = (
                pd.concat([monthly_series, category_series, technical_series_df], ignore_index=True)
                .groupby("month", dropna=False)["score_pct"]
                .mean()
                .sort_index()
                .round(2)
                .tolist()
            )

    role_metrics: list[dict[str, Any]] = []
    for role_label in dashboard_role_labels[:3]:
        role_avg = safe_float(role_summaries.get(role_label, {}).get("avg_rating_score", 0.0))
        role_total = int(role_summaries.get(role_label, {}).get("total_submissions", 0))
        role_series: list[float] = []
        if not scoped_category_filtered.empty:
            role_df = scoped_category_filtered[scoped_category_filtered["role_group"].astype(str).str.lower() == role_label.lower()].copy()
            if not role_df.empty:
                role_series = (
                    role_df.groupby("month", dropna=False)["avg_rating_score"]
                    .mean()
                    .sort_index()
                    .round(2)
                    .tolist()
                )
        if role_label.lower() == "trainee" and (not technical_filtered.empty or technical_submission_count > 0):
            trainee_total = technical_submission_count
            if not technical_filtered.empty:
                trainee_avg = safe_avg(technical_filtered.get("technical_rating_4", pd.Series(dtype=float)), default=0.0)
                trainee_series = (
                    technical_filtered.groupby("month", dropna=False)["technical_rating_4"]
                    .mean()
                    .sort_index()
                    .round(2)
                    .tolist()
                )
            else:
                trainee_avg = round((safe_float(technical_summary.get("technical_accuracy_pct", 0.0)) / 100.0) * 4.0, 2) if trainee_total > 0 else 0.0
                trainee_series = [trainee_avg] if trainee_total > 0 else []
            if role_total > 0 and trainee_total > 0:
                role_avg = round(((role_avg * role_total) + (trainee_avg * trainee_total)) / (role_total + trainee_total), 2)
            elif trainee_total > 0:
                role_avg = round(trainee_avg, 2)
            role_total += trainee_total
            role_series = role_series + trainee_series if role_series else trainee_series
        delta_text, delta_dir = safe_delta(role_series)
        role_metrics.append(
            {
                "title": f"{role_label} Rating",
                "value": f"{role_avg:.2f}/4" if role_total > 0 else "No data",
                "delta_text": delta_text,
                "delta_dir": delta_dir,
                "spark_values": role_series if role_series else [role_avg],
                "status_color": ROLE_ACCENT_COLORS.get(role_label, "#94a3b8") if role_total > 0 else "#64748b",
                "subcopy": f"{role_total} submissions" if role_total > 0 else "No records",
            }
        )

    total_delta_text, total_delta_dir = safe_delta(total_series)
    overall_delta_text, overall_delta_dir = safe_delta(overall_series)
    negative_series = sentiment_summary.get("trend_df", pd.DataFrame()).get("negative", pd.Series(dtype=float)).tolist() if sentiment_summary.get("total", 0) else []
    negative_delta_text, negative_delta_dir = safe_delta(negative_series)
    alerts_list = alert_metrics.get("alerts", []) or []
    actual_negative_count = int(sentiment_summary.get("negative_count", 0))
    critical_count = len([a for a in alerts_list if a.get("severity") == "high"])
    return {
        "total": {
            "title": "Total Feedback Submitted",
            "value": f"{total_records:,}",
            "delta_text": total_delta_text,
            "delta_dir": total_delta_dir,
            "spark_values": total_series if total_series else [float(total_records)],
            "status_color": "#38bdf8",
            "subcopy": "No data" if total_records == 0 else f"{safe_count(scoped_category_filtered)} category / {safe_count(monthly_filtered)} monthly / {technical_submission_count} technical",
        },
        "overall": {
            "title": "Overall Average Score",
            "value": f"{overall_avg_pct:.1f}%" if total_records > 0 else "No data",
            "delta_text": overall_delta_text,
            "delta_dir": overall_delta_dir,
            "spark_values": overall_series if overall_series else [overall_avg_pct],
            "status_color": "#8b5cf6" if total_records > 0 else "#64748b",
            "subcopy": "Actual filtered records only" if total_records > 0 else "No records",
        },
        "roles": role_metrics,
        "negative": {
            "title": "Negative Sentiment Alerts",
            "value": str(actual_negative_count) if sentiment_summary.get("total", 0) > 0 else "No data",
            "delta_text": negative_delta_text,
            "delta_dir": negative_delta_dir,
            "spark_values": negative_series if negative_series else [float(actual_negative_count)],
            "status_color": "#ef4444" if actual_negative_count > 0 else "#22c55e",
            "subcopy": f"{critical_count} critical alerts" if alerts_list else "No alerts",
        },
    }


@st.cache_data(show_spinner=False, ttl=15)
def fetch_live_dashboard_bundle(base_url: str, filters: dict[str, Any], viewer_role: str, refresh_nonce: int = 0) -> dict[str, Any]:
    del refresh_nonce
    dashboard_role_form_ids = get_dashboard_role_form_ids(viewer_role)
    params = {
        "month": None if filters.get("month") in {None, "", "All"} else filters.get("month"),
        "district": None if filters.get("district") in {None, "", "All"} else filters.get("district"),
        "trade_name": None if filters.get("trade_name") in {None, "", "All"} else filters.get("trade_name"),
        "year": None if filters.get("year") in {None, "", "All"} else int(filters.get("year")),
        "semester": None if filters.get("semester") in {None, "", "All"} else int(filters.get("semester")),
    }

    def timed_get(path: str, query: dict[str, Any] | None = None) -> tuple[bool, Any, float]:
        start = datetime.now()
        ok, data = api_get(base_url, path, params=query)
        elapsed = (datetime.now() - start).total_seconds() * 1000.0
        return ok, data, round(elapsed, 1)

    health_ok, health_data, health_ms = timed_get("/health")
    summary_ok, summary_data, summary_ms = timed_get("/dashboard/summary", params)
    trend_ok, trend_data, _ = timed_get("/dashboard/trend", {"group_by": "month", **{k: v for k, v in params.items() if k in {"district", "trade_name"}}})
    category_ok, category_data, _ = timed_get("/category-feedback/summary", {})
    category_trend_ok, category_trend_data, _ = timed_get("/category-feedback/trend", {"group_by": "month"})
    technical_ok, technical_data, technical_ms = timed_get("/authority/technical-summary", params)
    combined_ok, combined_data, combined_ms = timed_get("/authority/combined-summary", params)
    combined_trend_ok, combined_trend_data, _ = timed_get("/authority/combined-trend", {"group_by": "month", **{k: v for k, v in params.items() if v is not None}})

    role_summaries: dict[str, dict[str, Any]] = {}
    for role_label, form_ids in dashboard_role_form_ids.items():
        total_submissions = 0
        weighted_score = 0.0
        recent_rows: list[dict[str, Any]] = []
        for form_id in form_ids:
            ok_role, data_role, _ = timed_get("/category-feedback/summary", {"form_id": form_id})
            if ok_role and isinstance(data_role, dict):
                count = int(data_role.get("total_submissions", 0))
                avg = safe_float(data_role.get("avg_rating_score", 0.0))
                total_submissions += count
                weighted_score += avg * count
                recent_rows.extend(data_role.get("recent_submissions", []) or [])
        role_summaries[role_label] = {
            "total_submissions": total_submissions,
            "avg_rating_score": round(weighted_score / total_submissions, 2) if total_submissions else 0.0,
            "recent_submissions": recent_rows[:10],
        }

    return {
        "connected": bool(health_ok),
        "health": health_data if isinstance(health_data, dict) else {"detail": health_data},
        "health_response_ms": health_ms,
        "summary": summary_data if summary_ok and isinstance(summary_data, dict) else {},
        "summary_response_ms": summary_ms,
        "trend": trend_data if trend_ok and isinstance(trend_data, list) else [],
        "category_summary": category_data if category_ok and isinstance(category_data, dict) else {},
        "category_trend": category_trend_data if category_trend_ok and isinstance(category_trend_data, list) else [],
        "technical_summary": technical_data if technical_ok and isinstance(technical_data, dict) else {},
        "technical_response_ms": technical_ms,
        "combined_summary": combined_data if combined_ok and isinstance(combined_data, dict) else {},
        "combined_response_ms": combined_ms,
        "combined_trend": combined_trend_data if combined_trend_ok and isinstance(combined_trend_data, list) else [],
        "dashboard_roles": list(dashboard_role_form_ids.keys()),
        "role_summaries": role_summaries,
        "errors": [data for ok, data in [(health_ok, health_data), (summary_ok, summary_data), (combined_ok, combined_data)] if not ok],
    }


def filter_category_roles(category_df: pd.DataFrame, role_labels: list[str]) -> pd.DataFrame:
    if category_df.empty or not role_labels or "role_group" not in category_df.columns:
        return category_df.copy()
    allowed = {label.lower() for label in role_labels}
    return category_df[category_df["role_group"].astype(str).str.lower().isin(allowed)].copy()


def summarize_category_df(category_df: pd.DataFrame) -> dict[str, Any]:
    if category_df.empty:
        return {"top_forms": []}
    grouped = (
        category_df.groupby(["form_id", "form_title"], dropna=False)
        .agg(
            submissions=("form_id", "count"),
            avg_rating_score=("avg_rating_score", "mean"),
        )
        .reset_index()
        .sort_values(by=["submissions", "avg_rating_score"], ascending=[False, False])
    )
    top_forms: list[dict[str, Any]] = []
    for _, row in grouped.head(10).iterrows():
        top_forms.append(
            {
                "form_id": value_or_empty(row.get("form_id")),
                "form_title": value_or_empty(row.get("form_title")),
                "submissions": int(row.get("submissions", 0)),
                "avg_rating_score": round(safe_float(row.get("avg_rating_score", 0.0)), 2),
            }
        )
    return {"top_forms": top_forms}


def build_detail_table(mirrored: dict[str, pd.DataFrame], filters: dict[str, Any], role_labels: list[str] | None = None) -> pd.DataFrame:
    category_df = apply_dashboard_filters(mirrored.get("category", pd.DataFrame()), filters)
    technical_df = apply_dashboard_filters(mirrored.get("technical", pd.DataFrame()), filters)
    if role_labels:
        category_df = filter_category_roles(category_df, role_labels)
        if not technical_df.empty:
            technical_df = technical_df[technical_df["role_group"].astype(str).str.lower().isin({label.lower() for label in role_labels})].copy()

    detail_rows: list[dict[str, Any]] = []
    if not category_df.empty:
        sorted_cat = category_df.sort_values(by="submitted_at", ascending=False).reset_index(drop=True)
        for idx, row in sorted_cat.iterrows():
            details = row.get("basic_details", {}) if isinstance(row.get("basic_details", {}), dict) else {}
            rating = safe_float(row.get("avg_rating_score", 0.0))
            submitted_at = row.get("submitted_at")
            detail_rows.append(
                {
                    "Name": infer_feedback_name(details, value_or_empty(row.get("role_group", "Feedback")), idx + 1),
                    "Role": value_or_empty(row.get("role_group", "Trainee")),
                    "Trade": value_or_empty(row.get("trade_name")),
                    "Rating": round(rating, 2),
                    "Sentiment": sentiment_from_rating(rating),
                    "Submitted Time": submitted_at.strftime("%d %b %Y, %I:%M %p") if pd.notna(submitted_at) else "Unavailable",
                    "Status": status_from_rating(rating),
                    "_ts": submitted_at.to_pydatetime() if pd.notna(submitted_at) else None,
                    "_comment": value_or_empty(row.get("comment_text")),
                }
            )

    if not technical_df.empty:
        sorted_tech = technical_df.sort_values(by="submitted_at", ascending=False).reset_index(drop=True)
        for idx, row in sorted_tech.iterrows():
            rating = safe_float(row.get("technical_rating_4", 0.0))
            submitted_at = row.get("submitted_at")
            detail_rows.append(
                {
                    "Name": value_or_empty(row.get("student_name")) or f"Trainee Submission {idx + 1}",
                    "Role": value_or_empty(row.get("role_group", "Trainee")),
                    "Trade": value_or_empty(row.get("trade_name")),
                    "Rating": round(rating, 2),
                    "Sentiment": sentiment_from_rating(rating),
                    "Submitted Time": submitted_at.strftime("%d %b %Y, %I:%M %p") if pd.notna(submitted_at) else "Unavailable",
                    "Status": status_from_rating(rating),
                    "_ts": submitted_at.to_pydatetime() if pd.notna(submitted_at) else None,
                    "_comment": f"Technical test submitted with {int(row.get('correct_answers', 0))}/{int(row.get('mcq_answered', 0))} correct answers.",
                }
            )

    if not detail_rows:
        return pd.DataFrame(columns=["Name", "Role", "Trade", "Rating", "Sentiment", "Submitted Time", "Status", "_ts", "_comment"])
    return pd.DataFrame(detail_rows)


def summarize_role_groups_from_df(category_df: pd.DataFrame, role_labels: list[str]) -> dict[str, dict[str, Any]]:
    summaries: dict[str, dict[str, Any]] = {}
    for role_label in role_labels:
        role_df = category_df[category_df.get("role_group", pd.Series(dtype=str)).astype(str).str.lower() == role_label.lower()].copy()
        if role_df.empty:
            summaries[role_label] = {
                "total_submissions": 0,
                "avg_rating_score": 0.0,
                "recent_submissions": [],
            }
            continue
        recent_rows: list[dict[str, Any]] = []
        recent_df = role_df.sort_values(by="submitted_at", ascending=False).head(10)
        for _, row in recent_df.iterrows():
            details = row.get("basic_details", {}) if isinstance(row.get("basic_details", {}), dict) else {}
            recent_rows.append(
                {
                    "submitted_at": str(row.get("submitted_at", "")),
                    "source": value_or_empty(row.get("source")),
                    "form_id": value_or_empty(row.get("form_id")),
                    "form_title": value_or_empty(row.get("form_title")),
                    "avg_rating_score": round(safe_float(row.get("avg_rating_score", 0.0)), 2),
                    "comment_text": value_or_empty(row.get("comment_text")),
                    "basic_details": details,
                    "trade_name": value_or_empty(row.get("trade_name")),
                }
            )
        summaries[role_label] = {
            "total_submissions": int(len(role_df)),
            "avg_rating_score": round(float(pd.to_numeric(role_df["avg_rating_score"], errors="coerce").fillna(0.0).mean()), 2),
            "recent_submissions": recent_rows,
        }
    return summaries


def summarize_individual_performance_from_df(category_df: pd.DataFrame, role_labels: list[str]) -> pd.DataFrame:
    if category_df.empty:
        return pd.DataFrame(columns=["Department", "Role", "Name", "Avg Rating", "Submissions", "Latest Submission", "Status"])

    working = category_df.copy()
    working["person_name"] = [
        infer_feedback_name(
            row.get("basic_details", {}) if isinstance(row.get("basic_details", {}), dict) else {},
            value_or_empty(row.get("role_group", "Feedback")),
            idx + 1,
        )
        for idx, (_, row) in enumerate(working.iterrows())
    ]
    if role_labels:
        working = filter_category_roles(working, role_labels)
    if working.empty:
        return pd.DataFrame(columns=["Department", "Role", "Name", "Avg Rating", "Submissions", "Latest Submission", "Status"])

    grouped = (
        working.groupby(["trade_name", "role_group", "person_name"], dropna=False)
        .agg(
            avg_rating_score=("avg_rating_score", "mean"),
            submissions=("person_name", "count"),
            latest_submission=("submitted_at", "max"),
        )
        .reset_index()
    )
    grouped["Department"] = grouped["trade_name"].fillna("").astype(str).replace("", "Unassigned")
    grouped["Role"] = grouped["role_group"].fillna("").astype(str)
    grouped["Name"] = grouped["person_name"].fillna("").astype(str)
    grouped["Avg Rating"] = pd.to_numeric(grouped["avg_rating_score"], errors="coerce").fillna(0.0).round(2)
    grouped["Submissions"] = pd.to_numeric(grouped["submissions"], errors="coerce").fillna(0).astype(int)
    grouped["Latest Submission"] = pd.to_datetime(grouped["latest_submission"], errors="coerce").dt.strftime("%d %b %Y, %I:%M %p").fillna("Unavailable")
    grouped["Status"] = grouped["Avg Rating"].apply(status_from_rating)
    grouped = grouped.sort_values(by=["Department", "Role", "Avg Rating", "Submissions", "Name"], ascending=[True, True, True, False, True])
    return grouped[["Department", "Role", "Name", "Avg Rating", "Submissions", "Latest Submission", "Status"]]


def build_sentiment_trend_df(mirrored: dict[str, pd.DataFrame], filters: dict[str, Any]) -> pd.DataFrame:
    monthly_df = apply_dashboard_filters(mirrored.get("monthly", pd.DataFrame()), filters)
    rows: list[dict[str, Any]] = []
    if not monthly_df.empty:
        working = monthly_df.copy()
        submitted_ts = pd.to_datetime(working.get("submitted_at"), errors="coerce")
        if submitted_ts.notna().any():
            working["day"] = submitted_ts.dt.strftime("%Y-%m-%d")
        else:
            working["day"] = working.get("month", pd.Series(dtype=str)).astype(str)
        rows.extend(
            working.assign(sentiment_label=working["sentiment_label"].apply(normalize_sentiment_label))
            .groupby(["day", "sentiment_label"], dropna=False)
            .size()
            .reset_index(name="count")
            .to_dict(orient="records")
        )
    if not rows:
        return pd.DataFrame(columns=["day", "positive", "neutral", "negative"])
    df = pd.DataFrame(rows)
    grouped = df.groupby(["day", "sentiment_label"], dropna=False)["count"].sum().reset_index()
    pivot = grouped.pivot(index="day", columns="sentiment_label", values="count").fillna(0).reset_index()
    for col in ["positive", "neutral", "negative"]:
        if col not in pivot.columns:
            pivot[col] = 0
    pivot["day_sort"] = pd.to_datetime(pivot["day"], format="%Y-%m-%d", errors="coerce")
    pivot = pivot.sort_values(by="day_sort").drop(columns=["day_sort"])
    return pivot


def build_submission_velocity_df(detail_df: pd.DataFrame) -> pd.DataFrame:
    if detail_df.empty or "_ts" not in detail_df.columns:
        return pd.DataFrame(columns=["bucket", "count"])
    live_df = detail_df.dropna(subset=["_ts"]).copy()
    if live_df.empty:
        return pd.DataFrame(columns=["bucket", "count"])
    live_df["bucket"] = pd.to_datetime(live_df["_ts"]).dt.floor("h")
    grouped = live_df.groupby("bucket", dropna=False).size().reset_index(name="count")
    grouped = grouped.sort_values(by="bucket").tail(12)
    grouped["bucket_label"] = grouped["bucket"].dt.strftime("%I %p").str.lstrip("0")
    return grouped


def plot_dark_line(df: pd.DataFrame, x: str, y: str, title: str, color: str, area: bool = False) -> go.Figure:
    fig = go.Figure()
    fig.add_trace(
        go.Scatter(
            x=df[x],
            y=df[y],
            mode="lines+markers",
            name=title,
            line=dict(color=color, width=3, shape="spline"),
            marker=dict(size=8, color=color),
            fill="tozeroy" if area else None,
            fillcolor="rgba(56, 189, 248, 0.14)" if area else None,
        )
    )
    fig.update_layout(
        template="plotly_dark",
        paper_bgcolor="rgba(0,0,0,0)",
        plot_bgcolor="rgba(15,23,42,0.28)",
        margin=dict(l=18, r=18, t=24, b=18),
        height=320,
        hovermode="x unified",
        transition_duration=450,
        legend=dict(orientation="h", y=1.12, x=0),
    )
    fig.update_xaxes(showgrid=False)
    fig.update_yaxes(gridcolor="rgba(148,163,184,0.12)", zeroline=False)
    return fig


def plot_role_distribution(role_summaries: dict[str, dict[str, Any]]) -> go.Figure:
    labels = list(role_summaries.keys())
    values = [int(role_summaries[label].get("total_submissions", 0)) for label in labels]
    colors = [ROLE_ACCENT_COLORS.get(label, "#94a3b8") for label in labels]
    fig = go.Figure(
        data=[
            go.Pie(
                labels=labels,
                values=values,
                hole=0.66,
                marker=dict(colors=colors),
                textinfo="label+percent",
            )
        ]
    )
    fig.update_layout(
        template="plotly_dark",
        paper_bgcolor="rgba(0,0,0,0)",
        plot_bgcolor="rgba(0,0,0,0)",
        margin=dict(l=10, r=10, t=10, b=10),
        height=320,
    )
    return fig


def plot_category_scores(category_summary: dict[str, Any]) -> go.Figure:
    top_forms = category_summary.get("top_forms", []) or []
    df = pd.DataFrame(top_forms).head(8)
    if df.empty:
        return go.Figure()
    df["label"] = df["form_title"].astype(str).str.slice(0, 24)
    fig = go.Figure(
        data=[
            go.Bar(
                x=df["avg_rating_score"],
                y=df["label"],
                orientation="h",
                marker=dict(
                    color=df["avg_rating_score"],
                    colorscale=[[0, "#ef4444"], [0.5, "#f59e0b"], [1, "#22c55e"]],
                ),
                text=df["submissions"].astype(str) + " subs",
                textposition="outside",
                hovertemplate="%{y}<br>Avg score: %{x:.2f}<br>%{text}<extra></extra>",
            )
        ]
    )
    fig.update_layout(
        template="plotly_dark",
        paper_bgcolor="rgba(0,0,0,0)",
        plot_bgcolor="rgba(0,0,0,0)",
        margin=dict(l=12, r=12, t=20, b=12),
        height=320,
    )
    fig.update_xaxes(range=[0, 4], gridcolor="rgba(148,163,184,0.12)")
    return fig


def plot_sentiment_trend(sentiment_df: pd.DataFrame) -> go.Figure:
    fig = go.Figure()
    palette = {"positive": "#22c55e", "neutral": "#f59e0b", "negative": "#ef4444"}
    x_col = "day" if "day" in sentiment_df.columns else "month"
    working = sentiment_df.copy()
    if x_col == "day":
        parsed = pd.to_datetime(working[x_col], errors="coerce")
        working["x_label"] = parsed.dt.strftime("%d %b").fillna(working[x_col].astype(str))
    else:
        working["x_label"] = working[x_col].astype(str)

    compact_view = len(working) <= 7
    for column in ["positive", "neutral", "negative"]:
        if compact_view:
            fig.add_trace(
                go.Bar(
                    x=working["x_label"],
                    y=working[column],
                    name=column.title(),
                    marker=dict(color=palette[column]),
                    hovertemplate="%{x}<br>" + column.title() + ": %{y}<extra></extra>",
                )
            )
        else:
            fig.add_trace(
                go.Scatter(
                    x=working["x_label"],
                    y=working[column],
                    mode="lines+markers",
                    name=column.title(),
                    line=dict(color=palette[column], width=3, shape="spline"),
                    marker=dict(size=7, color=palette[column]),
                )
            )
    fig.update_layout(
        template="plotly_dark",
        paper_bgcolor="rgba(0,0,0,0)",
        plot_bgcolor="rgba(15,23,42,0.28)",
        margin=dict(l=18, r=18, t=24, b=18),
        height=320,
        hovermode="x unified",
        transition_duration=450,
        barmode="group" if compact_view else None,
    )
    fig.update_xaxes(showgrid=False, type="category", tickangle=0)
    fig.update_yaxes(gridcolor="rgba(148,163,184,0.12)", zeroline=False, rangemode="tozero", dtick=1 if compact_view else None)
    return fig


def plot_role_comparison(role_summaries: dict[str, dict[str, Any]]) -> go.Figure:
    labels = list(role_summaries.keys())
    ratings = [safe_float(role_summaries[label].get("avg_rating_score", 0.0)) for label in labels]
    colors = [ROLE_ACCENT_COLORS.get(label, "#94a3b8") for label in labels]
    fig = go.Figure(
        data=[
            go.Bar(
                x=labels,
                y=ratings,
                marker=dict(color=colors),
                text=[f"{v:.2f}/4" for v in ratings],
                textposition="outside",
                hovertemplate="%{x}<br>Avg rating: %{y:.2f}/4<extra></extra>",
            )
        ]
    )
    fig.update_layout(
        template="plotly_dark",
        paper_bgcolor="rgba(0,0,0,0)",
        plot_bgcolor="rgba(0,0,0,0)",
        margin=dict(l=12, r=12, t=20, b=12),
        height=320,
    )
    fig.update_yaxes(range=[0, 4], gridcolor="rgba(148,163,184,0.12)")
    return fig


def render_kpi_card(title: str, value: str, delta_text: str, delta_dir: str, spark_values: list[float], color: str, status_color: str, subcopy: str) -> None:
    spark_svg = build_sparkline_svg(spark_values, color)
    st.markdown(
        f"""
        <div class="kpi-card">
            <div style="display:flex;align-items:center;justify-content:space-between;gap:10px;">
                <div class="kpi-label">{title}</div>
                <span class="status-indicator" style="background:{status_color};"></span>
            </div>
            <div class="kpi-value">{value}</div>
            <div style="display:flex;justify-content:space-between;align-items:center;gap:10px;">
                <div class="kpi-delta {delta_dir}">{delta_text}</div>
                <div class="kpi-sub">{subcopy}</div>
            </div>
            <div class="sparkline-wrap">{spark_svg}</div>
        </div>
        """,
        unsafe_allow_html=True,
    )


def render_empty_state(title: str, subtitle: str, icon: str = "INFO") -> None:
    st.markdown(
        f"""
        <div class="empty-state">
            <div class="empty-icon">{icon}</div>
            <div class="empty-title">{title}</div>
            <div class="empty-copy">{subtitle}</div>
        </div>
        """,
        unsafe_allow_html=True,
    )


def render_activity_panel(detail_df: pd.DataFrame) -> None:
    live_df = detail_df.sort_values(by=["_ts", "Submitted Time"], ascending=False).head(8)
    body_html = ""
    if live_df.empty:
        body_html = (
            '<div class="empty-state">'
            '<div class="empty-icon">FEED</div>'
            '<div class="empty-title">No live activity available</div>'
            '<div class="empty-copy">Try changing filters or wait for new submissions.</div>'
            "</div>"
        )
    else:
        item_html: list[str] = []
        for _, row in live_df.iterrows():
            dt = row.get("_ts")
            relative = format_relative_time(dt)
            rating = safe_float(row.get("Rating", 0.0))
            role_text = value_or_empty(row.get("Role", "Feedback"))
            trade_text = value_or_empty(row.get("Trade", "General"))
            sentiment_text = value_or_empty(row.get("Sentiment", ""))
            status_text = value_or_empty(row.get("Status", "Healthy"))
            message = f"{role_text} submitted {trade_text} feedback."
            sentiment_class = sentiment_text.lower() if sentiment_text.lower() in {"positive", "neutral", "negative"} else "neutral"
            status_class = "status-good" if status_text in {"Excellent", "Healthy"} else ("status-warn" if status_text in {"Good", "Watch"} else "status-bad")
            item_html.append(
                (
                    f'<div class="activity-item">'
                    f'<div class="activity-time">{relative}</div>'
                    f'<div class="activity-text">'
                    f"<div>{message}</div>"
                    f'<div class="activity-meta">'
                    f'<span class="activity-badge role">{role_text}</span>'
                    f'<span class="activity-badge rating">{rating:.1f}/4</span>'
                    f'<span class="activity-badge {sentiment_class}">{sentiment_text or "No sentiment"}</span>'
                    f'<span class="activity-badge {status_class}">{status_text}</span>'
                    f"</div>"
                    f"</div>"
                    f"</div>"
                )
            )
        body_html = "".join(item_html)
    st.markdown(
        f'<div class="feed-card"><div class="panel-title">Live Activity Feed</div><div class="panel-subtitle">Recent submissions and operational shifts across active forms.</div>{body_html}</div>',
        unsafe_allow_html=True,
    )


def render_alert_card(alert: dict[str, str]) -> None:
    severity = value_or_empty(alert.get("severity", "low")).lower()
    st.markdown(
        f"""
        <div class="alert-item {severity}">
            <div class="alert-dot"></div>
            <div class="alert-text"><strong>{alert.get('title')}</strong><br/>{alert.get('message')}</div>
        </div>
        """,
        unsafe_allow_html=True,
    )


def render_alert_panel(alerts: list[dict[str, str]]) -> None:
    body_html = ""
    if not alerts:
        body_html = (
            '<div class="empty-state">'
            '<div class="empty-icon">ALERT</div>'
            '<div class="empty-title">No alerts available</div>'
            '<div class="empty-copy">No rules are currently triggered for the selected scope.</div>'
            "</div>"
        )
    else:
        alert_html: list[str] = []
        for alert in alerts:
            severity = value_or_empty(alert.get("severity", "low")).lower()
            alert_html.append(
                (
                    f'<div class="alert-item {severity}">'
                    f'<div class="alert-dot"></div>'
                    f"<div class=\"alert-text\"><strong>{alert.get('title')}</strong><br/>{alert.get('message')}</div>"
                    f"</div>"
                )
            )
        body_html = "".join(alert_html)
    st.markdown(
        f'<div class="status-card"><div class="panel-title">Smart Alerts</div><div class="panel-subtitle">Threshold-based signals surfaced from live submissions and health telemetry.</div>{body_html}</div>',
        unsafe_allow_html=True,
    )


def render_report_actions(actions: list[dict[str, Any]]) -> None:
    st.markdown('<div class="dashboard-card"><div class="panel-title">Report Actions</div><div class="panel-subtitle">Export exactly what the dashboard is currently showing.</div></div>', unsafe_allow_html=True)
    for action in actions:
        st.download_button(
            value_or_empty(action.get("label")),
            data=action.get("data", b"") or b"",
            file_name=value_or_empty(action.get("file_name", "report.csv")),
            mime=value_or_empty(action.get("mime", "text/csv")),
            use_container_width=True,
            disabled=bool(action.get("disabled", False)),
        )


def render_auto_refresh(interval_seconds: int = 15) -> None:
    st.caption(
        f"Auto refresh is armed for every {interval_seconds} seconds, but full-page reload is disabled to preserve login state. Use Manual Refresh for a safe refresh."
    )


def build_dashboard_alerts(bundle: dict[str, Any], role_summaries: dict[str, dict[str, Any]], sentiment_df: pd.DataFrame) -> list[dict[str, str]]:
    alerts: list[dict[str, str]] = []
    summary = bundle.get("summary", {})
    combined = bundle.get("combined_summary", {})
    total_role_submissions = sum(int(v.get("total_submissions", 0)) for v in role_summaries.values())
    trainer_submissions = int(role_summaries.get("Trainer", {}).get("total_submissions", 0))
    supervisor_submissions = int(role_summaries.get("Supervisor", {}).get("total_submissions", 0))
    negative_count = int(summary.get("negative_count", 0))
    total_count = int(summary.get("total_submissions", 0))
    negative_rate = pct(negative_count, total_count)

    if not bundle.get("connected"):
        alerts.append({"severity": "high", "title": "API disconnected", "message": "Dashboard is running from local mirrored data because FastAPI is unreachable."})
    if total_role_submissions > 0 and pct(trainer_submissions, total_role_submissions) < 20:
        alerts.append({"severity": "medium", "title": "Low trainer participation", "message": "Trainer-originated submissions are under the expected contribution baseline."})
    if supervisor_submissions == 0:
        alerts.append({"severity": "high", "title": "Supervisor submissions missing", "message": "No supervisor feedback is visible for the selected scope."})
    if negative_rate >= 25:
        alerts.append({"severity": "high", "title": "Negative sentiment spike", "message": f"Negative sentiment has reached {negative_rate:.1f}% of current submissions."})
    if safe_float(summary.get("avg_attendance_pct", 0.0)) < 75:
        alerts.append({"severity": "medium", "title": "Attendance below threshold", "message": "Average attendance is below the 75% watch threshold."})
    if safe_float(combined.get("risk_index", 0.0)) >= 35:
        alerts.append({"severity": "high", "title": "Elevated risk index", "message": f"Combined risk index is {safe_float(combined.get('risk_index', 0.0)):.1f}, indicating weak technical or category performance."})
    if alerts:
        return alerts[:4]
    if not sentiment_df.empty:
        alerts.append({"severity": "low", "title": "Stable sentiment", "message": "Positive and neutral sentiment remain dominant for the current filter set."})
    return alerts

def category_basic_prefill(field_name: str, role_key: str) -> str:
    if role_key not in {"iti_trainee", "student"}:
        return ""
    ctx = st.session_state.get("category_generate_payload") or {}
    if not isinstance(ctx, dict):
        return ""

    lower = field_name.strip().lower()
    if "institute" in lower:
        return str(ctx.get("institute_name", "") or "")
    if lower in {"trade", "trade / course", "iti trade"} or lower.startswith("trade"):
        return str(ctx.get("trade_name", "") or "")
    if "district" in lower:
        return str(ctx.get("district", "") or "")
    if lower == "month":
        return str(ctx.get("month", "") or "")
    if lower == "year":
        v = ctx.get("year")
        return "" if v is None else str(v)
    if "semester" in lower:
        v = ctx.get("semester")
        return "" if v is None else str(v)
    if "batch" in lower or "session" in lower:
        month = str(ctx.get("month", "") or "")
        return month if month else ""
    return ""


if "generated_questions" not in st.session_state:
    st.session_state.generated_questions = None

if "last_feedback_response" not in st.session_state:
    st.session_state.last_feedback_response = None

if "last_generate_payload" not in st.session_state:
    st.session_state.last_generate_payload = None

if "trend_state" not in st.session_state:
    st.session_state.trend_state = None
if "category_generated_questions" not in st.session_state:
    st.session_state.category_generated_questions = []
if "category_generate_payload" not in st.session_state:
    st.session_state.category_generate_payload = {}
if "auth_user" not in st.session_state:
    st.session_state.auth_user = None
if "auth_password" not in st.session_state:
    st.session_state.auth_password = ""
if "auth_session" not in st.session_state:
    st.session_state.auth_session = {
        "is_authenticated": False,
        "logged_in_at": None,
        "last_seen_at": None,
        "username": "",
        "role": "",
        "persistent_session_id": "",
        "expires_at": None,
    }
if "student_submit_thank_you" not in st.session_state:
    st.session_state.student_submit_thank_you = ""
if "category_submit_thank_you" not in st.session_state:
    st.session_state.category_submit_thank_you = ""
if "dashboard_auto_refresh" not in st.session_state:
    st.session_state.dashboard_auto_refresh = False
if "dashboard_refresh_nonce" not in st.session_state:
    st.session_state.dashboard_refresh_nonce = 0
if "dashboard_last_synced" not in st.session_state:
    st.session_state.dashboard_last_synced = None
if "dashboard_export_data" not in st.session_state:
    st.session_state.dashboard_export_data = b""
if "dashboard_demo_mode" not in st.session_state:
    st.session_state.dashboard_demo_mode = False

inject_dashboard_styles()
mirrored_data = load_mirrored_feedback_data(DEFAULT_API_BASE)
filter_options = build_filter_options(mirrored_data)
auth_user = st.session_state.get("auth_user")
if not auth_user:
    restored_session_id = get_auth_query_session_id()
    if restore_persistent_auth_session(restored_session_id):
        auth_user = st.session_state.get("auth_user")
if auth_user:
    st.session_state.auth_session["is_authenticated"] = True
    st.session_state.auth_session["last_seen_at"] = datetime.now().isoformat()
    st.session_state.auth_session["username"] = value_or_empty(auth_user.get("username"))
    st.session_state.auth_session["role"] = value_or_empty(auth_user.get("role")).lower()

st.sidebar.markdown('<div class="sidebar-section"><div class="sidebar-title">Feedback Intelligence</div><div class="sidebar-copy">Production-style control surface for live ITI feedback, category, and technical analytics.</div></div>', unsafe_allow_html=True)

api_base = st.sidebar.text_input(
    "FastAPI Base URL",
    value=DEFAULT_API_BASE,
    help="Example: http://127.0.0.1:8000",
)

connected, status_msg = check_api(api_base)
if connected:
    st.sidebar.markdown(
        '<div class="sidebar-section"><div class="sidebar-title">API Status</div><div class="sidebar-copy"><span class="badge-pill"><span class="pulse-dot"></span> Connected</span></div></div>',
        unsafe_allow_html=True,
    )
else:
    st.sidebar.markdown(
        f'<div class="sidebar-section"><div class="sidebar-title">API Status</div><div class="sidebar-copy"><span class="badge-pill" style="color:#fecaca;">Offline</span><div style="margin-top:8px;">{status_msg}</div></div></div>',
        unsafe_allow_html=True,
    )

st.sidebar.markdown("---")
st.sidebar.markdown("### Account Access")

if auth_user:
    st.sidebar.markdown(
        f'<div class="sidebar-section"><div class="sidebar-title">Logged In User</div><div class="sidebar-copy"><strong>{value_or_empty(auth_user.get("username"))}</strong><br/>{value_or_empty(auth_user.get("role"))}</div></div>',
        unsafe_allow_html=True,
    )
    assigned_trade_label = value_or_empty(auth_user.get("assigned_trade")).strip()
    assigned_year_value = auth_user.get("assigned_year")
    if assigned_trade_label or assigned_year_value:
        details = []
        if assigned_trade_label:
            details.append(f"Trade: `{assigned_trade_label}`")
        if assigned_year_value:
            details.append(f"Year: `{assigned_year_value}`")
        st.sidebar.caption(" | ".join(details))
    auth_meta = st.session_state.get("auth_session", {})
    if auth_meta.get("logged_in_at"):
        st.sidebar.caption(f"Session started: {value_or_empty(auth_meta.get('logged_in_at')).replace('T', ' ')[:19]}")
    if auth_meta.get("expires_at"):
        st.sidebar.caption(f"Session expires: {value_or_empty(auth_meta.get('expires_at')).replace('T', ' ')[:19]}")
    if st.sidebar.button("Logout", use_container_width=True):
        remove_persistent_auth_session(value_or_empty(auth_meta.get("persistent_session_id")))
        set_auth_query_session_id("")
        clear_auth_session()
        st.rerun()
else:
    with st.sidebar.expander("Login", expanded=True):
        with st.form("sidebar_login_form"):
            s_login_username = st.text_input("Username", key="sidebar_login_username")
            s_login_password = st.text_input("Password", type="password", key="sidebar_login_password")
            s_login_submit = st.form_submit_button("Login", use_container_width=True)
        if s_login_submit:
            ok_login, login_data = api_post(
                api_base,
                "/auth/login",
                {"username": s_login_username.strip(), "password": s_login_password},
            )
            if ok_login and isinstance(login_data, dict):
                set_auth_session(login_data, s_login_password)
                persistent_session_id = create_persistent_auth_session(login_data, s_login_password)
                st.session_state.auth_session["persistent_session_id"] = persistent_session_id
                st.session_state.auth_session["expires_at"] = (
                    datetime.now() + timedelta(hours=AUTH_SESSION_TTL_HOURS)
                ).isoformat()
                set_auth_query_session_id(persistent_session_id)
                st.sidebar.success("Login successful.")
                st.rerun()
            else:
                st.sidebar.error("Login failed.")

selected_month = "All"
selected_district = "All"
selected_trade = "All"
selected_year = "All"
selected_semester = "All"
selected_role_group = "All"

analytics_sidebar_allowed = bool(auth_user) and value_or_empty((auth_user or {}).get("role")).lower() in ANALYTICS_ALLOWED_ROLES
if analytics_sidebar_allowed:
    analytics_viewer_role = value_or_empty((auth_user or {}).get("role")).lower()
    trade_filter_label = "Department" if analytics_viewer_role == "principal" else "Trade Name"
    analytics_assigned_trade = value_or_empty((auth_user or {}).get("assigned_trade")).strip()
    st.sidebar.markdown("---")
    st.sidebar.markdown('<div class="sidebar-section"><div class="sidebar-title">Dashboard Filters</div><div class="sidebar-copy">These controls apply only to the analytics dashboard.</div></div>', unsafe_allow_html=True)
    selected_month = st.sidebar.selectbox("Month", options=filter_options["months"], index=0, key="dashboard_filter_month")
    selected_district = st.sidebar.selectbox("District", options=filter_options["districts"], index=0, key="dashboard_filter_district")
    if analytics_viewer_role == "trainer" and analytics_assigned_trade:
        selected_trade = st.sidebar.selectbox(trade_filter_label, options=[analytics_assigned_trade], index=0, key="dashboard_filter_trade", disabled=True)
        st.sidebar.caption(f"Locked to assigned department: `{analytics_assigned_trade}`")
    else:
        selected_trade = st.sidebar.selectbox(trade_filter_label, options=filter_options["trades"], index=0, key="dashboard_filter_trade")
    selected_year = st.sidebar.selectbox("Year", options=filter_options["years"], index=0, key="dashboard_filter_year")
    selected_semester = st.sidebar.selectbox("Semester", options=filter_options["semesters"], index=0, key="dashboard_filter_semester")
    selected_role_group = st.sidebar.selectbox("Role Group", options=filter_options["roles"], index=0, key="dashboard_filter_role")
    st.session_state.dashboard_auto_refresh = st.sidebar.toggle("Auto refresh every 15s", value=st.session_state.dashboard_auto_refresh)
    st.session_state.dashboard_demo_mode = st.sidebar.toggle("Use demo preview data when real data is empty", value=st.session_state.dashboard_demo_mode)

    sidebar_filters = {
        "month": selected_month,
        "district": selected_district,
        "trade_name": selected_trade,
        "year": selected_year,
        "semester": selected_semester,
        "role_group": selected_role_group,
    }
    sidebar_source = build_filtered_dashboard_data(
        api_base,
        mirrored_data,
        sidebar_filters,
        analytics_viewer_role,
        auth_user=auth_user if isinstance(auth_user, dict) else None,
        use_demo_preview=st.session_state.dashboard_demo_mode,
    )
    sidebar_role_scope = build_principal_scope_data(
        sidebar_source["category_filtered"],
        analytics_viewer_role,
        auth_user=auth_user if isinstance(auth_user, dict) else None,
    )
    sidebar_export_df = build_detail_table(
        {"category": sidebar_role_scope["scoped_category_filtered"], "technical": sidebar_source.get("technical_filtered", pd.DataFrame())},
        {"month": "All", "district": "All", "trade_name": "All", "year": "All", "semester": "All", "role_group": "All"},
    ).drop(columns=["_ts", "_comment"], errors="ignore")
    st.session_state.dashboard_export_data = sidebar_export_df.to_csv(index=False).encode("utf-8") if not sidebar_export_df.empty else b""

    export_payload = st.session_state.get("dashboard_export_data", b"") or b""
    st.sidebar.download_button(
        "Export report",
        data=export_payload,
        file_name=f"feedback_intelligence_report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv",
        mime="text/csv",
        use_container_width=True,
        disabled=not bool(export_payload),
    )
    st.sidebar.markdown('<div class="sidebar-section"><div class="sidebar-title">Settings</div><div class="sidebar-copy">Use filters to scope live analytics. Auto-refresh no longer forces a page reload, so login stays intact.</div></div>', unsafe_allow_html=True)
else:
    st.session_state.dashboard_export_data = b""

with st.sidebar.expander("Register User", expanded=False):
    registerer_role = value_or_empty((st.session_state.get("auth_user") or {}).get("role")).lower()
    role_labels = {
        "admin": "Admin",
        "principal": "Principal",
        "trainer": "Trainer",
        "ojt_trainer": "OJT Trainer",
        "supervisor": "Supervisor",
        "trainee": "Trainee",
    }
    permission_map = {
        "admin": ["admin", "principal", "trainer", "ojt_trainer", "supervisor", "trainee"],
    }
    allowed_registration_roles = permission_map.get(registerer_role, [])
    auth_headers = auth_api_headers()
    trade_options = get_available_trade_names() + ["Other"]

    if not st.session_state.get("auth_user"):
        s_reg_role = "trainee"
        st.caption("Student self-registration is available here. Other roles must be created by admin.")
        with st.form("sidebar_public_register_form"):
            s_reg_full_name = st.text_input("Full Name", key="sidebar_public_reg_full_name")
            s_reg_username = st.text_input("Username", key="sidebar_public_reg_username")
            s_reg_email = st.text_input("Email", key="sidebar_public_reg_email")
            s_reg_password = st.text_input("Password", type="password", key="sidebar_public_reg_password")
            s_reg_confirm_password = st.text_input("Confirm Password", type="password", key="sidebar_public_reg_confirm_password")
            s_reg_district = st.text_input("District", key="sidebar_public_reg_district")
            s_reg_department = st.text_input("Department", key="sidebar_public_reg_department")
            s_reg_trade_choice = st.selectbox(
                "Department / Trade",
                trade_options,
                index=0,
                key="sidebar_public_reg_trade_choice",
            )
            s_reg_assigned_trade_other = ""
            if s_reg_trade_choice == "Other":
                s_reg_assigned_trade_other = st.text_input("Custom Department / Trade", key="sidebar_public_reg_assigned_trade_other")
            s_reg_assigned_year = st.selectbox(
                "Year",
                options=["", "1", "2"],
                index=0,
                key="sidebar_public_reg_assigned_year",
            )
            s_reg_semester = st.selectbox(
                "Semester",
                options=["", "1", "2", "3", "4"],
                index=0,
                key="sidebar_public_reg_semester",
            )
            s_reg_status = st.selectbox(
                "Status",
                options=["active", "inactive"],
                index=0,
                key="sidebar_public_reg_status",
            )
            s_reg_submit = st.form_submit_button("Create Student Account", use_container_width=True)
        if s_reg_submit:
            selected_trade = s_reg_assigned_trade_other.strip() if s_reg_trade_choice == "Other" else s_reg_trade_choice.strip()
            if not s_reg_username.strip():
                st.sidebar.error("Username is required.")
            elif not s_reg_email.strip():
                st.sidebar.error("Email is required.")
            elif len(s_reg_password) < 8:
                st.sidebar.error("Password must be at least 8 characters long.")
            elif s_reg_password != s_reg_confirm_password:
                st.sidebar.error("Confirm Password must match.")
            elif not selected_trade:
                st.sidebar.error("Assigned department/trade is required for trainee.")
            elif not value_or_empty(s_reg_assigned_year).strip() or not value_or_empty(s_reg_semester).strip():
                st.sidebar.error("Year and semester are required for trainee.")
            else:
                registration_payload = {
                    "full_name": s_reg_full_name.strip(),
                    "username": s_reg_username.strip(),
                    "email": s_reg_email.strip(),
                    "password": s_reg_password,
                    "role": s_reg_role,
                    "district": s_reg_district.strip(),
                    "department": s_reg_department.strip(),
                    "assigned_trade": selected_trade,
                    "assigned_year": value_or_empty(s_reg_assigned_year).strip() or None,
                    "semester": value_or_empty(s_reg_semester).strip() or None,
                    "status": s_reg_status,
                }
                with st.spinner("Creating student account..."):
                    ok_reg, reg_data = api_post(
                        api_base,
                        "/register-user",
                        registration_payload,
                    )
                if ok_reg:
                    created_user = (reg_data or {}).get("user", {}) if isinstance(reg_data, dict) else {}
                    st.sidebar.success(
                        f"Student account created for {value_or_empty(created_user.get('username'))}. You can log in now."
                    )
                    for key in [
                        "sidebar_public_reg_full_name",
                        "sidebar_public_reg_username",
                        "sidebar_public_reg_email",
                        "sidebar_public_reg_password",
                        "sidebar_public_reg_confirm_password",
                        "sidebar_public_reg_district",
                        "sidebar_public_reg_department",
                        "sidebar_public_reg_trade_choice",
                        "sidebar_public_reg_assigned_trade_other",
                        "sidebar_public_reg_assigned_year",
                        "sidebar_public_reg_semester",
                        "sidebar_public_reg_status",
                    ]:
                        if key in st.session_state:
                            del st.session_state[key]
                    st.rerun()
                else:
                    reg_error = value_or_empty((reg_data or {}).get("detail")) if isinstance(reg_data, dict) else value_or_empty(reg_data)
                    st.sidebar.error(reg_error or "Student account creation failed.")
    elif not allowed_registration_roles:
        st.info("This role is not allowed to create new user accounts.")
    elif not auth_headers:
        st.warning("This session does not have a valid access token. Log in again before creating users.")
    else:
        selected_role_index = 0
        existing_selected_role = st.session_state.get("sidebar_reg_role")
        if existing_selected_role in allowed_registration_roles:
            selected_role_index = allowed_registration_roles.index(existing_selected_role)
        s_reg_role = st.selectbox(
            "Role",
            allowed_registration_roles,
            index=selected_role_index,
            key="sidebar_reg_role",
            format_func=lambda value: role_labels.get(value, value.replace("_", " ").title()),
        )
        requires_trade_assignment = s_reg_role in {"trainer", "ojt_trainer", "trainee", "supervisor"}
        requires_year_assignment = s_reg_role == "trainee"
        requires_semester_assignment = s_reg_role == "trainee"
        requires_department_scope = s_reg_role == "supervisor"

        with st.form("sidebar_register_form"):
            s_reg_full_name = st.text_input("Full Name", key="sidebar_reg_full_name")
            s_reg_username = st.text_input("Username", key="sidebar_reg_username")
            s_reg_email = st.text_input("Email", key="sidebar_reg_email")
            s_reg_password = st.text_input("Password", type="password", key="sidebar_reg_password")
            s_reg_confirm_password = st.text_input("Confirm Password", type="password", key="sidebar_reg_confirm_password")
            s_reg_district = st.text_input("District", key="sidebar_reg_district")
            s_reg_department = st.text_input("Department", key="sidebar_reg_department")
            s_reg_trade_choice = trade_options[0]
            if requires_trade_assignment:
                trade_label = "Department / Trade"
                s_reg_trade_choice = st.selectbox(
                    trade_label,
                    trade_options,
                    index=0,
                    key="sidebar_reg_trade_choice",
                    help="Use the assigned department/trade scope for this account.",
                )
            s_reg_assigned_trade_other = ""
            if requires_trade_assignment and s_reg_trade_choice == "Other":
                s_reg_assigned_trade_other = st.text_input("Custom Department / Trade", key="sidebar_reg_assigned_trade_other")
            s_reg_assigned_year = None
            if requires_year_assignment:
                s_reg_assigned_year = st.selectbox(
                    "Year",
                    options=["", "1", "2"],
                    index=0,
                    key="sidebar_reg_assigned_year",
                )
            s_reg_semester = ""
            if requires_semester_assignment:
                s_reg_semester = st.selectbox(
                    "Semester",
                    options=["", "1", "2", "3", "4"],
                    index=0,
                    key="sidebar_reg_semester",
                )
            s_reg_status = st.selectbox(
                "Status",
                options=["active", "inactive"],
                index=0,
                key="sidebar_reg_status",
            )
            s_reg_submit = st.form_submit_button("Create Account", use_container_width=True)

        reset_col, hint_col = st.columns([1, 1])
        with reset_col:
            if st.button("Reset Form", use_container_width=True, key="sidebar_reg_reset_button"):
                clear_registration_form_state()
                st.rerun()
        with hint_col:
            if s_reg_role == "trainee":
                st.caption("Trainee requires department/trade, year, and semester.")
            elif s_reg_role in {"trainer", "ojt_trainer"}:
                st.caption("Trainer roles require assigned department/trade.")
            elif requires_department_scope:
                st.caption("Supervisor requires department/trade scope.")

        if s_reg_submit:
            selected_trade = ""
            if requires_trade_assignment:
                selected_trade = s_reg_assigned_trade_other.strip() if s_reg_trade_choice == "Other" else s_reg_trade_choice.strip()
            if not s_reg_username.strip():
                st.sidebar.error("Username is required.")
            elif not s_reg_email.strip():
                st.sidebar.error("Email is required.")
            elif len(s_reg_password) < 8:
                st.sidebar.error("Password must be at least 8 characters long.")
            elif s_reg_password != s_reg_confirm_password:
                st.sidebar.error("Confirm Password must match.")
            elif s_reg_role in {"trainer", "ojt_trainer", "trainee"} and not selected_trade:
                st.sidebar.error("Assigned department/trade is required for this role.")
            elif s_reg_role == "trainee" and (not value_or_empty(s_reg_assigned_year).strip() or not value_or_empty(s_reg_semester).strip()):
                st.sidebar.error("Year and semester are required for trainee.")
            elif requires_department_scope and not (s_reg_department.strip() or selected_trade):
                st.sidebar.error("Department or trade scope is required for supervisor.")
            else:
                registration_payload = {
                    "full_name": s_reg_full_name.strip(),
                    "username": s_reg_username.strip(),
                    "email": s_reg_email.strip(),
                    "password": s_reg_password,
                    "role": s_reg_role,
                    "district": s_reg_district.strip(),
                    "department": s_reg_department.strip(),
                    "assigned_trade": selected_trade,
                    "assigned_year": value_or_empty(s_reg_assigned_year).strip() or None,
                    "semester": value_or_empty(s_reg_semester).strip() or None,
                    "status": s_reg_status,
                }
                with st.spinner("Creating user account..."):
                    ok_reg, reg_data = api_post(
                        api_base,
                        "/register-user",
                        registration_payload,
                        headers=auth_headers,
                    )
                if ok_reg:
                    created_user = (reg_data or {}).get("user", {}) if isinstance(reg_data, dict) else {}
                    st.sidebar.success(
                        f"Created {role_labels.get(value_or_empty(created_user.get('role')).lower(), value_or_empty(created_user.get('role')))} account for {value_or_empty(created_user.get('username'))}."
                    )
                    clear_registration_form_state()
                    st.rerun()
                else:
                    reg_error = value_or_empty((reg_data or {}).get("detail")) if isinstance(reg_data, dict) else value_or_empty(reg_data)
                    st.sidebar.error(reg_error or "Account creation failed.")

        with st.spinner("Loading recently created users..."):
            ok_recent, recent_data = api_get(
                api_base,
                "/register-user/recent",
                params={"limit": 10},
                headers=auth_headers,
            )
        st.markdown("**Recently Created Users**")
        if ok_recent and isinstance(recent_data, list) and recent_data:
            recent_rows = []
            for row in recent_data:
                recent_rows.append(
                    {
                        "Full Name": value_or_empty((row or {}).get("full_name")),
                        "Username": value_or_empty((row or {}).get("username")),
                        "Email": value_or_empty((row or {}).get("email")),
                        "Role": role_labels.get(value_or_empty((row or {}).get("role")).lower(), value_or_empty((row or {}).get("role")).replace("_", " ").title()),
                        "Trade": value_or_empty((row or {}).get("assigned_trade")),
                        "Year": value_or_empty((row or {}).get("assigned_year")),
                        "Semester": value_or_empty((row or {}).get("semester")),
                        "District": value_or_empty((row or {}).get("district")),
                        "Status": value_or_empty((row or {}).get("status")).title(),
                        "Created By": value_or_empty((row or {}).get("created_by")),
                        "Created At": value_or_empty((row or {}).get("created_at")).replace("T", " ")[:19],
                    }
                )
            st.dataframe(pd.DataFrame(recent_rows), use_container_width=True, hide_index=True)
        elif ok_recent:
            st.caption("No users have been created in your visible scope yet.")
        else:
            st.caption("Recently created users could not be loaded for this session.")

current_login_role = value_or_empty((st.session_state.get("auth_user") or {}).get("role")).lower()
if current_login_role in {"iti_trainee", "trainee", "ojt_institute_officer", "ojt_trainer", "ojt_supervisor", "supervisor"}:
    st.markdown(
        """
        <style>
        div[data-testid="stTabs"] div[role="tablist"] button:nth-child(3) {
            display: none !important;
        }
        div[data-testid="stTabs"] div[data-testid="stTabPanel"]:nth-of-type(3) {
            display: none !important;
        }
        </style>
        """,
        unsafe_allow_html=True,
    )
elif current_login_role == "principal":
    st.markdown(
        """
        <style>
        div[data-testid="stTabs"] div[role="tablist"] button:nth-child(1) {
            display: none !important;
        }
        div[data-testid="stTabs"] div[data-testid="stTabPanel"]:nth-of-type(1) {
            display: none !important;
        }
        </style>
        """,
        unsafe_allow_html=True,
    )

tab1, tab2, tab3 = st.tabs(
    ["1. Technical Flow (DB)", "2. Category Feedback", "3. Category Analytics"]
)

with tab1:
    st.subheader("Trainer and Student Technical Feedback")
    st.caption("Trainer enters topics and generates questions. Students only answer generated questions.")
    if st.session_state.get("student_submit_thank_you"):
        st.success(st.session_state.student_submit_thank_you)
        st.session_state.student_submit_thank_you = ""
    auth_ctx = st.session_state.get("auth_user")
    if not auth_ctx:
        st.warning("Login from the sidebar to access role-based feedback flow.")
    else:
        current_role = value_or_empty(auth_ctx.get("role")).lower()
        assigned_trade = value_or_empty(auth_ctx.get("assigned_trade")).strip()
        assigned_year = auth_ctx.get("assigned_year")
        if current_role == "admin":
            st.info("Admin access is limited to registering trainer, principal, institute officer, and supervisor accounts from the sidebar.")

        if current_role in TRAINER_TECH_ROLES:
            st.markdown("### Trainer: Trade Question Bank")
            st.caption("Add or update a trade first, then generate a question set from the saved theory and practical bank.")
            if current_role == "trainer" and not assigned_trade:
                st.error("This trainer account has no assigned trade. Register the trainer with an assigned trade before using the trainer flow.")
            else:
                trainer_trade_locked = current_role == "trainer" and bool(assigned_trade)
                if st.session_state.get("trade_bank_save_message"):
                    st.success(st.session_state.trade_bank_save_message)
                    st.session_state.trade_bank_save_message = ""

                with st.expander("Step 1 - Add New Trade and Upload Questions", expanded=True):
                    s1, s2, s3 = st.columns(3)
                    with s1:
                        setup_trade = st.text_input(
                            "New / Existing Trade Name",
                            value=assigned_trade or st.session_state.get("trainer_active_trade", "Electrician"),
                            key="setup_trade_name",
                            disabled=trainer_trade_locked,
                        )
                    with s2:
                        setup_year_level = st.selectbox("Year Level", ["I", "II"], index=0, key="setup_year_level")
                    with s3:
                        setup_month = st.text_input("Month", value="2026-09", key="setup_month")

                    combined_files = st.file_uploader(
                        "Upload Combined Theory + Practical File",
                        type=QUESTION_UPLOAD_TYPES,
                        accept_multiple_files=True,
                        key="setup_combined_uploads",
                        help="Use this when one CSV/JSON/XLSX/PDF/DOCX contains both theory and practical questions. Structured files should include a source column with theory/practical.",
                    )
                    theory_files = st.file_uploader(
                        "Upload Trade Theory Questions Only",
                        type=QUESTION_UPLOAD_TYPES,
                        accept_multiple_files=True,
                        key="setup_theory_uploads",
                        help="Use CSV/JSON/XLSX for best results. PDF/DOCX/images use text extraction or OCR when installed.",
                    )
                    practical_files = st.file_uploader(
                        "Upload Trade Practical Questions Only",
                        type=QUESTION_UPLOAD_TYPES,
                        accept_multiple_files=True,
                        key="setup_practical_uploads",
                        help="Images inside documents are saved and shown with student questions when detected.",
                    )

                    save_trade_bank = st.button("Save Trade Question Bank", use_container_width=True, key="save_trade_bank_btn")
                    if save_trade_bank:
                        if not setup_trade.strip():
                            st.error("Enter a trade name before saving.")
                        elif not combined_files and not theory_files and not practical_files:
                            st.error("Upload a combined file or at least one theory/practical question file.")
                        else:
                            combined_df = parse_uploaded_question_files(
                                combined_files,
                                default_source="",
                                default_trade=setup_trade.strip(),
                                default_year_level=setup_year_level,
                                default_month=setup_month.strip(),
                            )
                            theory_df = parse_uploaded_question_files(
                                theory_files,
                                default_source="theory",
                                default_trade=setup_trade.strip(),
                                default_year_level=setup_year_level,
                                default_month=setup_month.strip(),
                            )
                            practical_df = parse_uploaded_question_files(
                                practical_files,
                                default_source="practical",
                                default_trade=setup_trade.strip(),
                                default_year_level=setup_year_level,
                                default_month=setup_month.strip(),
                            )
                            upload_df = pd.concat([combined_df, theory_df, practical_df], ignore_index=True)
                            upload_df = upload_df[upload_df["question_text"].fillna("").astype(str).str.strip() != ""]
                            upload_df["source"] = upload_df["source"].fillna("").astype(str).str.strip().str.lower()
                            missing_source = upload_df["source"].eq("")
                            if missing_source.any() and combined_files:
                                upload_df.loc[missing_source, "source"] = "theory"
                                st.warning(
                                    "Some combined-file rows did not specify theory/practical, so they were saved as theory. "
                                    "Use a source column or Theory/Practical sheet names for exact separation."
                                )
                            upload_df = upload_df[upload_df["source"].isin(["theory", "practical"])]
                            if upload_df.empty:
                                st.error("No valid questions could be parsed. For combined structured files, include a source column with theory or practical.")
                            else:
                                source_counts = upload_df["source"].value_counts().to_dict()
                                topic_preview = (
                                    upload_df.groupby(["source", "topic"], dropna=False)
                                    .size()
                                    .reset_index(name="questions")
                                    .sort_values(["source", "questions"], ascending=[True, False])
                                )
                                st.markdown("**Parsed Topic Summary**")
                                st.dataframe(topic_preview.head(30), use_container_width=True, hide_index=True)
                                ingest_payload = {
                                    "username": value_or_empty(auth_ctx.get("username")),
                                    "password": st.session_state.get("auth_password", ""),
                                    "rows": upload_df.fillna("").astype(str).to_dict(orient="records"),
                                    "default_trade": setup_trade.strip(),
                                    "default_year_level": setup_year_level,
                                    "default_month": setup_month.strip(),
                                }
                                ok_ingest, ingest_data = api_post(api_base, "/trainer/ingest-question-bank", ingest_payload)
                                if ok_ingest and isinstance(ingest_data, dict):
                                    st.session_state.trainer_active_trade = setup_trade.strip()
                                    load_topic_bank.cache_clear()
                                    st.session_state.trade_bank_save_message = (
                                        f"Saved {setup_trade.strip()} question bank. "
                                        f"Theory: {int(source_counts.get('theory', 0))}, "
                                        f"Practical: {int(source_counts.get('practical', 0))}. "
                                        f"New: {value_or_empty(ingest_data.get('inserted'))}, "
                                        f"Skipped existing: {value_or_empty(ingest_data.get('skipped'))}."
                                    )
                                    st.rerun()
                                else:
                                    st.error("Question bank save failed.")
                                    show_json_block("Question Bank Ingest Error", ingest_data)

                st.markdown("### Trainer: Generate Question Set")
                base_bank_df = load_topic_bank()
                topic_source_df = base_bank_df.copy()
                uploaded_df = empty_question_bank_df()

            subject_options = [
                "Trade Theory",
                "Trade Practical",
                "Workshop Calculation",
                "Engineering Drawing",
                "Employability Skills",
            ]
            subject_mode = st.radio(
                "Subject Selection",
                options=["Single Subject", "Multiple Subjects"],
                index=1,
                horizontal=True,
                key="trainer_subject_mode",
            )
            if subject_mode == "Single Subject":
                subject_count = 1
            else:
                subject_count_selection = st.selectbox(
                    "Number of Chapter / Topic Entries",
                    options=["Select Number", 2, 3],
                    index=0,
                    key="trainer_subject_count",
                    help="Choose how many individual chapter/topic selections you want to include in this generated question set.",
                )
                subject_count = 0 if subject_count_selection == "Select Number" else int(subject_count_selection)
            trainer_generate_submit = False
            selected_subjects: list[str] = []
            selected_topics: list[str] = []
            if subject_mode != "Single Subject" and subject_count == 0:
                st.info("Select the number of chapter/topic entries first.")
            else:
                with st.form("trainer_generate_qset_form"):
                    st.markdown("#### Chapter / Topic Selection")
                    entry_help = (
                        "Choose the chapter/topic entries directly. "
                        "The category is already controlled by Question Bank Mode above."
                    )
                    st.caption(entry_help)
                    support_col_1, support_col_2, support_col_3 = st.columns(3)
                    with support_col_1:
                        t_trade = st.text_input(
                            "Trade Name",
                            value=assigned_trade or st.session_state.get("trainer_active_trade", "Electrician"),
                            disabled=trainer_trade_locked,
                        )
                    with support_col_2:
                        if assigned_year in {1, 2}:
                            t_year = int(assigned_year)
                            st.text_input("Year", value=str(t_year), disabled=True)
                            semester_options = [1, 2] if t_year == 1 else [3, 4]
                            t_semester = st.selectbox(
                                "Semester",
                                semester_options,
                                index=0,
                                key="db_trainer_sem",
                                help="Semester is used only to save the generated set in the correct academic slot.",
                            )
                        else:
                            trainer_year_index = 0
                            t_year = st.selectbox("Year", [1, 2], index=trainer_year_index, key="db_trainer_year")
                            semester_options = [1, 2] if int(t_year) == 1 else [3, 4]
                            t_semester = st.selectbox(
                                "Semester",
                                semester_options,
                                index=0,
                                key="db_trainer_sem",
                            )
                    with support_col_3:
                        t_q_mode = st.selectbox(
                            "Selected Category",
                            options=["theory", "practical", "both"],
                            index=2,
                            help="This dropdown is the active category selector. theory = only uploaded theory-bank questions, practical = only uploaded practical-bank questions, both = mix both uploaded banks.",
                        )
                    if t_q_mode == "theory":
                        normalized_subject = "Trade Theory"
                    elif t_q_mode == "practical":
                        normalized_subject = "Trade Practical"
                    else:
                        normalized_subject = "Trade Theory"
                    selector_col_1, selector_col_2 = st.columns(2)
                    for idx in range(subject_count):
                        target_col = selector_col_1 if idx % 2 == 0 else selector_col_2
                        with target_col:
                            st.markdown(f"**Entry {idx + 1}**")
                            topic_opts = collect_topic_options(topic_source_df, t_trade, int(t_year), normalized_subject) if normalized_subject else []
                            topic_select_options = ["Select Topic"] + topic_opts if topic_opts else ["Select Topic"]
                            topic_val = st.selectbox(
                                f"Topic {idx + 1} for selected category",
                                topic_select_options,
                                index=0,
                                key=f"trainer_topic_{idx + 1}",
                                help="Choose the exact chapter/topic from the uploaded bank for the category selected above.",
                                )
                            if t_q_mode == "both":
                                entry_subject = "Trade Theory" if idx % 2 == 0 else "Trade Practical"
                            else:
                                entry_subject = normalized_subject
                            selected_subjects.append(entry_subject)
                            selected_topics.append("" if topic_val == "Select Topic" else topic_val.strip())
                    st.markdown("#### Save Details")
                    save_col_1, save_col_2, save_col_3 = st.columns(3)
                    available_bank_months = collect_bank_month_options(topic_source_df, t_trade, int(t_year))
                    with save_col_1:
                        if available_bank_months:
                            default_month = available_bank_months[-1]
                            t_month = st.selectbox(
                                "Month",
                                options=available_bank_months,
                                index=len(available_bank_months) - 1,
                                key="trainer_question_month",
                                help="Only uploaded bank months available for this trade and year are shown here.",
                            )
                            if len(available_bank_months) == 1:
                                st.caption(f"Available uploaded bank month: {default_month}")
                        else:
                            t_month = st.text_input("Month (YYYY-MM)", value="2026-09")
                    with save_col_2:
                        t_district = st.text_input("District", value="Vijayawada")
                    with save_col_3:
                        t_institute = st.text_input("Institute Name", value="Govt ITI")
                        t_question_count = st.number_input(
                            "Number of Questions to Generate",
                            min_value=1,
                            max_value=50,
                            value=3,
                            step=1,
                        )
                    trainer_generate_submit = st.form_submit_button("Generate and Save Question Set", use_container_width=True)

            if trainer_generate_submit:
                valid_pairs = [(s, t) for s, t in zip(selected_subjects, selected_topics) if s and t]
                if len(valid_pairs) != subject_count:
                    st.error("Select a topic for each chapter/topic entry before generating the question set.")
                    st.stop()
                while len(selected_subjects) < 3:
                    selected_subjects.append("")
                while len(selected_topics) < 3:
                    selected_topics.append("")
                trainer_payload = {
                    "username": value_or_empty(auth_ctx.get("username")),
                    "password": st.session_state.get("auth_password", ""),
                    "month": t_month.strip(),
                    "district": t_district.strip(),
                    "institute_name": t_institute.strip(),
                    "trade_name": t_trade.strip(),
                    "year": int(t_year),
                    "semester": int(t_semester),
                    "question_mode": t_q_mode,
                    "question_count": int(t_question_count),
                    "subject_1": selected_subjects[0],
                    "topic_1": selected_topics[0],
                    "subject_2": selected_subjects[1],
                    "topic_2": selected_topics[1],
                    "subject_3": selected_subjects[2],
                    "topic_3": selected_topics[2],
                    "question_pool": topic_source_df.fillna("").astype(str).to_dict(orient="records"),
                }
                ok_tq, tq_data = api_post(api_base, "/trainer/generate-question-set", trainer_payload)
                if ok_tq and isinstance(tq_data, dict):
                    st.success(f"Question set saved (ID: {tq_data.get('question_set_id')}).")
                    st.caption(
                        f"Mode: {value_or_empty(tq_data.get('question_mode'))} | "
                        f"Target per student: {value_or_empty(tq_data.get('target_question_count'))} | "
                        f"Pool size: {value_or_empty(tq_data.get('question_pool_size'))}"
                    )
                    tq_questions = tq_data.get("questions", [])
                    if isinstance(tq_questions, list) and tq_questions:
                        readable_rows = []
                        for q in tq_questions:
                            readable_rows.append(
                                {
                                    "Subject": value_or_empty(q.get("subject")),
                                    "Topic": value_or_empty(q.get("topic")),
                                    "Question": value_or_empty(q.get("question")),
                                    "Has Image": "Yes" if value_or_empty(q.get("question_image")) else "No",
                                    "A": value_or_empty(q.get("option_a")),
                                    "B": value_or_empty(q.get("option_b")),
                                    "C": value_or_empty(q.get("option_c")),
                                    "D": value_or_empty(q.get("option_d")),
                                }
                            )
                        st.markdown("#### Generated Questions")
                        st.dataframe(pd.DataFrame(readable_rows), use_container_width=True, hide_index=True)
                else:
                    st.error("Question set generation failed.")
                    show_json_block("Generation Error", tq_data)
        else:
            # Non-trainer roles should simply not see the trainer section.
            pass

        st.markdown("---")
        if current_role in STUDENT_TECH_ROLES:
            st.markdown("### Student: Load Questions and Submit Technical Responses")
            student_trade_locked = bool(assigned_trade)
            student_year_locked = assigned_year in {1, 2}
            student_trade_value = assigned_trade or "Electrician"
            with st.form("student_load_qset_form"):
                s1, s2, s3 = st.columns(3)
                with s1:
                    s_trade = st.text_input(
                        "Trade Name",
                        value=student_trade_value,
                        disabled=student_trade_locked,
                    )
                with s2:
                    student_year_index = 0 if assigned_year != 2 else 1
                    if student_year_locked:
                        s_year = int(assigned_year)
                        st.text_input("Year", value=str(s_year), disabled=True)
                        semester_options = [1, 2] if int(s_year) == 1 else [3, 4]
                        s_semester = st.selectbox("Semester", semester_options, index=0, key="db_student_sem")
                    else:
                        s_year = st.selectbox("Year", [1, 2], index=student_year_index, key="db_student_year")
                        semester_options = [1, 2] if int(s_year) == 1 else [3, 4]
                        s_semester = st.selectbox("Semester", semester_options, index=0, key="db_student_sem")
                with s3:
                    load_qset = st.form_submit_button("Load Latest Question Set", use_container_width=True)
            if student_trade_locked:
                st.caption(f"Locked to assigned trade: `{assigned_trade}`")
            if student_year_locked:
                st.caption(f"Locked to assigned year: `{assigned_year}`")

            if load_qset:
                load_payload = {
                    "username": value_or_empty(auth_ctx.get("username")),
                    "password": st.session_state.get("auth_password", ""),
                    "trade_name": s_trade.strip(),
                    "year": int(s_year),
                    "semester": int(s_semester),
                }
                ok_sq, sq_data = api_post(api_base, "/student/latest-question-set", load_payload)
                if ok_sq and isinstance(sq_data, dict):
                    st.success(f"Loaded question set ID: {sq_data.get('id')}")
                    st.session_state["db_student_qset"] = sq_data
                else:
                    st.session_state["db_student_qset"] = None
                    st.error("Could not load question set.")
                    show_json_block("Student Question Set Error", sq_data)

            qset = st.session_state.get("db_student_qset")
            if isinstance(qset, dict) and qset.get("questions"):
                st.markdown("#### Questions for Student")
                questions = qset.get("questions", [])
                with st.form("student_submit_tech_feedback_form"):
                    responses_payload = []
                    for idx, q in enumerate(questions, start=1):
                        subject = value_or_empty(q.get("subject"))
                        topic = value_or_empty(q.get("topic"))
                        question = value_or_empty(q.get("question"))
                        question_image = value_or_empty(q.get("question_image"))
                        st.markdown(f"**Q{idx}. {question}**")
                        st.caption(f"Subject: {subject} | Topic: {topic}")
                        if question_image:
                            st.image(resolve_question_image_source(question_image), caption=f"Question image {idx}", use_container_width=True)

                        option_map = {
                            "A": value_or_empty(q.get("option_a")),
                            "B": value_or_empty(q.get("option_b")),
                            "C": value_or_empty(q.get("option_c")),
                            "D": value_or_empty(q.get("option_d")),
                        }
                        mcq_available = all(option_map[k] for k in ["A", "B", "C", "D"])
                        selected_option = None
                        selected_option_text = None
                        response_text = ""
                        if mcq_available:
                            display_opts = [f"{k}. {option_map[k]}" for k in ["A", "B", "C", "D"]]
                            picked = st.radio(
                                f"Select Answer {idx}",
                                options=display_opts,
                                key=f"student_response_option_{idx}",
                            )
                            selected_option = picked.split(".", 1)[0].strip()
                            selected_option_text = option_map.get(selected_option, "")
                        else:
                            response_text = st.text_area(
                                f"Technical Answer {idx}",
                                key=f"student_response_text_{idx}",
                                height=90,
                            )
                        responses_payload.append(
                            {
                                "subject": subject,
                                "topic": topic,
                                "question": question,
                                "response_text": response_text.strip(),
                                "selected_option": selected_option,
                                "selected_option_text": selected_option_text,
                                "confidence_score": 3.0,
                            }
                        )

                    st.markdown("**Context for this submission**")
                    c1, c2, c3 = st.columns(3)
                    with c1:
                        ctx_month = st.text_input("Month", value=value_or_empty(qset.get("month", "")), disabled=True)
                        ctx_district = st.text_input("District", value=value_or_empty(qset.get("district", "")), disabled=True)
                    with c2:
                        ctx_institute = st.text_input("Institute Name", value=value_or_empty(qset.get("institute_name", "")), disabled=True)
                        ctx_trade = st.text_input("Trade Name", value=value_or_empty(qset.get("trade_name", "")), disabled=True)
                    with c3:
                        ctx_year = st.number_input("Year", min_value=1, max_value=4, value=int(qset.get("year", 1)), step=1, disabled=True)
                        ctx_sem = st.number_input("Semester", min_value=1, max_value=8, value=int(qset.get("semester", 1)), step=1, disabled=True)

                    submit_tech = st.form_submit_button("Submit Student Technical Feedback", use_container_width=True)

                if submit_tech:
                    submit_payload = {
                        "username": value_or_empty(auth_ctx.get("username")),
                        "password": st.session_state.get("auth_password", ""),
                        "question_set_id": int(qset.get("id")),
                        "month": ctx_month.strip(),
                        "district": ctx_district.strip(),
                        "institute_name": ctx_institute.strip(),
                        "trade_name": ctx_trade.strip(),
                        "year": int(ctx_year),
                        "semester": int(ctx_sem),
                        "responses": responses_payload,
                    }
                    ok_sf, sf_data = api_post(api_base, "/student/submit-technical-feedback", submit_payload)
                    if ok_sf:
                        st.session_state.student_submit_thank_you = "Feedback submitted successfully. Duplicate submissions are not allowed."
                        st.session_state["db_student_qset"] = None
                        st.rerun()
                    else:
                        error_text = value_or_empty((sf_data or {}).get("detail")) if isinstance(sf_data, dict) else value_or_empty(sf_data)
                        normalized_error = error_text.lower()
                        if "duplicate technical feedback submission is not allowed" in normalized_error:
                            st.error("You have already submitted this test. Duplicate submissions are not allowed.")
                        else:
                            st.error(error_text or "Student technical feedback submission failed.")
        else:
            pass

with tab2:
    st.subheader("Category-wise Feedback Collection")
    st.caption("Uses dynamic form templates from your uploaded ITI/OJT category formats.")
    if st.session_state.get("category_submit_thank_you"):
        st.success(st.session_state.category_submit_thank_you)
        st.session_state.category_submit_thank_you = ""
    auth_ctx = st.session_state.get("auth_user")
    if not auth_ctx:
        st.warning("Login from the sidebar to submit category feedback.")
        selected_role_key = ""
    else:
        login_role = value_or_empty(auth_ctx.get("role")).lower()
        if login_role == "admin":
            st.info("Admin has no category feedback access.")
            selected_role_key = ""
        else:
            selected_role_key = LOGIN_TO_CATEGORY_ROLE.get(login_role, "")
            st.info(f"Logged role: `{login_role}`. Submitting feedback as `{selected_role_key}`.")

    if selected_role_key:
        if selected_role_key in {"student"}:
            st.info("Student question generation has moved to `Technical Flow (DB)` tab. Students now answer trainer-generated questions only.")

        ok_forms, forms_data = api_get(api_base, "/feedback/forms")
        if not ok_forms:
            st.error("Could not load feedback form templates.")
            show_json_block("Error", forms_data)
        elif not isinstance(forms_data, list) or len(forms_data) == 0:
            st.info("No category templates found. Please ensure `data/feedback_form_templates.json` is present.")
        else:
            allowed_ids = ROLE_ALLOWED_FORM_IDS.get(selected_role_key, set())
            role_filtered_forms = [
                f for f in forms_data
                if (not allowed_ids) or (str(f.get("form_id", "")).strip() in allowed_ids)
            ]

            sources = sorted({str(f.get("source", "")).strip() for f in role_filtered_forms if str(f.get("source", "")).strip()})
            source_filter = st.selectbox("Category Source", options=["All"] + sources, index=0)

            filtered_forms = role_filtered_forms
            if source_filter != "All":
                filtered_forms = [f for f in role_filtered_forms if str(f.get("source", "")).strip() == source_filter]

            if not filtered_forms:
                st.warning(f"No forms available for selected role: {selected_role_key}")
            else:
                form_options = {
                    f"{f.get('form_title', f.get('form_id', 'Unknown'))} ({f.get('source', '')})": f
                    for f in filtered_forms
                }
                selected_label = st.selectbox("Feedback Form", options=list(form_options.keys()), index=0)
                selected_form = form_options[selected_label]

                basic_fields = selected_form.get("basic_fields", [])
                parameters = selected_form.get("parameters", [])
                rating_scale = selected_form.get("rating_scale", ["Excellent", "Good", "Average", "Poor"])

                with st.form("category_feedback_form"):
                    st.markdown("### Basic Details")
                    basic_details_payload: dict[str, str] = {}
                    if basic_fields:
                        cols = st.columns(2)
                        for idx, field_name in enumerate(basic_fields):
                            col = cols[idx % 2]
                            field_text = str(field_name).strip()
                            field_key_base = "".join(ch if ch.isalnum() else "_" for ch in field_text.lower())
                            lower_name = field_text.lower()
                            with col:
                                if "duration" in lower_name or "period" in lower_name:
                                    st.markdown(f"**{field_text}**")
                                    default_from = date.today()
                                    default_to = date.today()
                                    ctx_month = str((st.session_state.get("category_generate_payload") or {}).get("month", "") or "")
                                    if ctx_month:
                                        try:
                                            parsed = datetime.strptime(ctx_month, "%Y-%m")
                                            default_from = date(parsed.year, parsed.month, 1)
                                            default_to = default_from
                                        except Exception:
                                            pass
                                    from_date = st.date_input(
                                        f"{field_text} From",
                                        value=default_from,
                                        key=f"cat_basic_{idx}_{field_key_base}_from",
                                    )
                                    to_date = st.date_input(
                                        f"{field_text} To",
                                        value=default_to,
                                        key=f"cat_basic_{idx}_{field_key_base}_to",
                                    )
                                    from_val = from_date.strftime("%d-%m-%Y")
                                    to_val = to_date.strftime("%d-%m-%Y")
                                    basic_details_payload[field_text] = f"From {from_val} To {to_val}"
                                elif lower_name == "date" or lower_name.endswith(" date"):
                                    picked_date = st.date_input(
                                        field_text,
                                        value=date.today(),
                                        key=f"cat_basic_{idx}_{field_key_base}_date",
                                    )
                                    basic_details_payload[field_text] = picked_date.strftime("%d-%m-%Y")
                                else:
                                    default_val = category_basic_prefill(field_text, selected_role_key)
                                    val = st.text_input(
                                        field_text,
                                        value=default_val,
                                        key=f"cat_basic_{idx}_{field_key_base}",
                                    )
                                    basic_details_payload[field_text] = val.strip()
                    else:
                        st.info("No basic detail fields configured in template.")

                    st.markdown("### Parameter Ratings")
                    parameter_scores_payload = []
                    if parameters:
                        for param in parameters:
                            p_no = param.get("s_no", "")
                            p_name = str(param.get("parameter", "")).strip()
                            if not p_name:
                                continue
                            rating = st.radio(
                                f"{p_no}. {p_name}" if p_no != "" else p_name,
                                options=rating_scale,
                                horizontal=True,
                                key=f"cat_param_{selected_form.get('form_id','')}_{p_no}_{p_name}",
                            )
                            parameter_scores_payload.append({"parameter": p_name, "rating": rating})
                    else:
                        st.info("No parameter rows configured for this form.")

                    category_comment = st.text_area("Comment (optional)", value="", height=100)
                    submitted_category = st.form_submit_button("Submit Category Feedback", use_container_width=True)

                if submitted_category:
                    payload = {
                        "submitted_by_role": selected_role_key,
                        "source": str(selected_form.get("source", "")).strip(),
                        "form_id": str(selected_form.get("form_id", "")).strip(),
                        "form_title": str(selected_form.get("form_title", "")).strip(),
                        "basic_details": basic_details_payload,
                        "parameter_scores": parameter_scores_payload,
                        "comment_text": category_comment,
                    }
                    ok_submit, submit_data = api_post(api_base, "/submit-category-feedback", payload)
                    if ok_submit:
                        st.session_state.category_submit_thank_you = "Feedback submitted successfully. Duplicate submissions are not allowed."
                        st.rerun()
                    else:
                        error_text = value_or_empty((submit_data or {}).get("detail")) if isinstance(submit_data, dict) else value_or_empty(submit_data)
                        normalized_error = error_text.lower()
                        if "duplicate category feedback submission detected" in normalized_error:
                            st.error("This feedback has already been submitted. Duplicate submissions are not allowed.")
                        else:
                            st.error(error_text or "Category feedback submission failed.")

with tab3:
    auth_ctx = st.session_state.get("auth_user")
    role = value_or_empty((auth_ctx or {}).get("role", "")).lower()
    if not auth_ctx:
        st.warning("Login from the sidebar to access dashboard analytics.")
    elif role not in ANALYTICS_ALLOWED_ROLES:
        st.info("Your role has no analytics dashboard access.")
    else:
        dashboard_filters = {
            "month": selected_month,
            "district": selected_district,
            "trade_name": selected_trade,
            "year": selected_year,
            "semester": selected_semester,
            "role_group": selected_role_group,
        }

        hero_left, hero_right = st.columns([4.6, 1.6])
        last_synced = st.session_state.get("dashboard_last_synced")
        if not last_synced:
            last_synced = datetime.now()
        with hero_right:
            st.markdown("")
            if st.button("Manual Refresh", use_container_width=True, type="primary"):
                st.session_state.dashboard_refresh_nonce += 1
                fetch_live_dashboard_bundle.clear()
                load_mirrored_feedback_data.clear()
                st.rerun()
            st.toggle("Auto refresh", key="dashboard_auto_refresh_header", value=st.session_state.dashboard_auto_refresh, disabled=True)
            st.caption("Sidebar toggle controls refresh cadence.")

        if st.session_state.dashboard_auto_refresh:
            render_auto_refresh(15)

        loading_container = st.container()
        with loading_container:
            s1, s2, s3 = st.columns(3)
            for col in (s1, s2, s3):
                with col:
                    st.markdown('<div class="skeleton"></div>', unsafe_allow_html=True)

        mirrored_data = load_mirrored_feedback_data(api_base)
        live_bundle = fetch_live_dashboard_bundle(api_base, dashboard_filters, current_login_role, st.session_state.dashboard_refresh_nonce)
        st.session_state.dashboard_last_synced = datetime.now()
        loading_container.empty()

        dashboard_data = build_filtered_dashboard_data(
            api_base,
            mirrored_data,
            dashboard_filters,
            role,
            auth_user=auth_ctx if isinstance(auth_ctx, dict) else None,
            use_demo_preview=st.session_state.dashboard_demo_mode,
        )
        effective_trade_label = value_or_empty(dashboard_data.get("effective_filters", {}).get("trade_name", selected_trade)) or selected_trade
        hero_pills = [
            '<span class="meta-pill"><span class="pulse-dot"></span> LIVE</span>',
            f'<span class="meta-pill">Last synced {last_synced.strftime("%d %b %Y, %I:%M:%S %p")}</span>',
            f'<span class="meta-pill">Filters: {selected_month} | {effective_trade_label} | {selected_role_group}</span>',
        ]
        if st.session_state.dashboard_demo_mode:
            hero_pills.append('<span class="meta-pill">Demo Preview Mode Armed</span>')
        hero_meta_html = "".join(hero_pills)
        with hero_left:
            st.markdown(
                f"""
                <div class="dashboard-hero">
                    <h1>Feedback Intelligence Dashboard</h1>
                    <p>Enterprise-style visibility for Principal, Trainer, OJT Trainer, Supervisor, and Trainee feedback performance.</p>
                    <div class="hero-meta">
                        {hero_meta_html}
                    </div>
                </div>
                """,
                unsafe_allow_html=True,
            )
        scope_data = build_principal_scope_data(
            dashboard_data["category_filtered"],
            role,
            auth_user=auth_ctx if isinstance(auth_ctx, dict) else None,
        )
        dashboard_role_labels = scope_data["dashboard_role_labels"]
        scoped_category_filtered = scope_data["scoped_category_filtered"]
        scoped_category_summary = scope_data["scoped_category_summary"]
        individual_performance_df = scope_data["individual_performance_df"]
        category_role_summaries = scope_data["role_summaries"]
        role_summaries = {
            label: {
                "total_submissions": int(data.get("total_submissions", 0)),
                "avg_rating_score": safe_float(data.get("avg_rating_score", 0.0)),
                "recent_submissions": list(data.get("recent_submissions", [])),
            }
            for label, data in category_role_summaries.items()
        }
        monthly_filtered = dashboard_data["monthly_filtered"]
        technical_filtered = dashboard_data["technical_filtered"]
        category_filtered = scoped_category_filtered
        detail_df = build_detail_table(
            {"category": scoped_category_filtered, "technical": technical_filtered},
            {"month": "All", "district": "All", "trade_name": "All", "year": "All", "semester": "All", "role_group": "All"},
        )
        if "Trainee" in role_summaries and not technical_filtered.empty:
            technical_total = int(len(technical_filtered))
            technical_avg = safe_avg(technical_filtered.get("technical_rating_4", pd.Series(dtype=float)), default=0.0)
            trainee_existing_total = int(role_summaries["Trainee"].get("total_submissions", 0))
            trainee_existing_avg = safe_float(role_summaries["Trainee"].get("avg_rating_score", 0.0))
            if trainee_existing_total > 0:
                merged_avg = round(
                    ((trainee_existing_avg * trainee_existing_total) + (technical_avg * technical_total))
                    / max(trainee_existing_total + technical_total, 1),
                    2,
                )
            else:
                merged_avg = round(technical_avg, 2)
            tech_recent_rows = [
                {
                    "submitted_at": str(row.get("submitted_at", "")),
                    "source": "technical_flow",
                    "form_id": "technical_question_test",
                    "form_title": "Trainer Generated Technical Test",
                    "avg_rating_score": round(safe_float(row.get("technical_rating_4", 0.0)), 2),
                    "comment_text": f"Technical score {safe_float(row.get('technical_score_pct', 0.0)):.1f}%",
                    "basic_details": {"name": value_or_empty(row.get("student_name"))},
                    "trade_name": value_or_empty(row.get("trade_name")),
                }
                for _, row in technical_filtered.sort_values(by="submitted_at", ascending=False).head(10).iterrows()
            ]
            role_summaries["Trainee"] = {
                "total_submissions": trainee_existing_total + technical_total,
                "avg_rating_score": merged_avg,
                "recent_submissions": (tech_recent_rows + list(role_summaries["Trainee"].get("recent_submissions", [])))[:10],
            }
        elif "Trainee" in role_summaries and int(live_bundle.get("technical_summary", {}).get("technical_submissions", 0) or 0) > 0:
            technical_total = int(live_bundle.get("technical_summary", {}).get("technical_submissions", 0) or 0)
            technical_avg = round((safe_float(live_bundle.get("technical_summary", {}).get("technical_accuracy_pct", 0.0)) / 100.0) * 4.0, 2)
            trainee_existing_total = int(role_summaries["Trainee"].get("total_submissions", 0))
            trainee_existing_avg = safe_float(role_summaries["Trainee"].get("avg_rating_score", 0.0))
            if trainee_existing_total > 0:
                merged_avg = round(
                    ((trainee_existing_avg * trainee_existing_total) + (technical_avg * technical_total))
                    / max(trainee_existing_total + technical_total, 1),
                    2,
                )
            else:
                merged_avg = technical_avg
            role_summaries["Trainee"] = {
                "total_submissions": trainee_existing_total + technical_total,
                "avg_rating_score": merged_avg,
                "recent_submissions": list(role_summaries["Trainee"].get("recent_submissions", [])),
            }
        submission_velocity_df = build_submission_velocity_df(detail_df)
        sentiment_summary = build_sentiment_summary(detail_df)
        sentiment_trend_df = sentiment_summary["trend_df"]
        alert_metrics = build_alert_metrics_from_filtered_df(
            monthly_filtered,
            scoped_category_filtered,
            technical_filtered,
            role_summaries,
            dashboard_role_labels,
            sentiment_summary,
            bool(live_bundle.get("connected")),
            live_bundle.get("errors", []),
        )
        alerts = alert_metrics["alerts"]
        kpi_metrics = build_kpi_metrics_from_filtered_df(
            monthly_filtered,
            scoped_category_filtered,
            technical_filtered,
            role_summaries,
            sentiment_summary,
            alert_metrics,
            dashboard_role_labels,
            live_bundle.get("technical_summary", {}),
        )
        report_export_df = detail_df.drop(columns=[c for c in ["_ts", "_comment"] if c in detail_df.columns], errors="ignore")
        st.session_state.dashboard_export_data = report_export_df.to_csv(index=False).encode("utf-8") if not report_export_df.empty else b""

        report_scope = {
            "Month": selected_month,
            "District": selected_district,
            "Trade" if role != "principal" else "Department": effective_trade_label,
            "Year": selected_year,
            "Semester": selected_semester,
            "Role Group": selected_role_group,
            "Record Count": str(len(report_export_df)),
        }
        report_actions = [
            {
                "label": "Export Filtered CSV",
                "data": report_export_df.to_csv(index=False).encode("utf-8") if not report_export_df.empty else b"",
                "file_name": f"filtered_feedback_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv",
                "mime": "text/csv",
                "disabled": report_export_df.empty,
            },
            {
                "label": "Export Principal Performance Report",
                "data": individual_performance_df.to_csv(index=False).encode("utf-8") if not individual_performance_df.empty else b"",
                "file_name": f"principal_performance_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv",
                "mime": "text/csv",
                "disabled": individual_performance_df.empty,
            },
            {
                "label": "Export Category Analytics Report",
                "data": pd.DataFrame(scoped_category_summary.get('top_forms', [])).to_csv(index=False).encode("utf-8")
                if scoped_category_summary.get("top_forms")
                else b"",
                "file_name": f"category_analytics_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv",
                "mime": "text/csv",
                "disabled": not bool(scoped_category_summary.get("top_forms")),
            },
            {
                "label": "Export Sentiment Report",
                "data": monthly_filtered[["month", "district", "trade_name", "sentiment_label"]].to_csv(index=False).encode("utf-8")
                if not monthly_filtered.empty and {"month", "district", "trade_name", "sentiment_label"}.issubset(set(monthly_filtered.columns))
                else b"",
                "file_name": f"sentiment_report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv",
                "mime": "text/csv",
                "disabled": monthly_filtered.empty,
            },
            {
                "label": "Generate Monthly Summary PDF",
                "data": b"",
                "file_name": "monthly_summary.pdf",
                "mime": "application/pdf",
                "disabled": True,
            },
        ]
        role_label_phrase = ", ".join(dashboard_role_labels[:-1]) + f" and {dashboard_role_labels[-1]}" if len(dashboard_role_labels) > 1 else dashboard_role_labels[0]
        role_comparison_subtitle = f"{role_label_phrase} quality rating side by side."
        role_distribution_subtitle = f"Role-wise split of {role_label_phrase} participation."
        using_demo_preview = bool(dashboard_data["using_demo_preview"])

        st.markdown('<div class="section-label">Top KPI Signal Layer</div>', unsafe_allow_html=True)
        k_cols = st.columns(6)
        render_order = [kpi_metrics["total"], kpi_metrics["overall"], *kpi_metrics["roles"], kpi_metrics["negative"]]
        for col, metric in zip(k_cols, render_order):
            with col:
                render_kpi_card(
                    value_or_empty(metric.get("title")),
                    value_or_empty(metric.get("value")),
                    value_or_empty(metric.get("delta_text")),
                    value_or_empty(metric.get("delta_dir")),
                    metric.get("spark_values", [0.0]) or [0.0],
                    ROLE_ACCENT_COLORS.get(value_or_empty(metric.get("title")).replace(" Rating", ""), "#38bdf8"),
                    value_or_empty(metric.get("status_color", "#38bdf8")),
                    value_or_empty(metric.get("subcopy")),
                )

        overview_tab, live_tab, category_tab, sentiment_tab, health_tab, reports_tab = st.tabs(
            ["Overview", "Live Feed", "Category Analytics", "Sentiment Monitor", "Technical Health", "Reports"]
        )

        with overview_tab:
            row1_left, row1_right = st.columns([1.8, 1.1])
            with row1_left:
                st.markdown('<div class="panel-title">Submission Trend</div><div class="panel-subtitle">Recent submission velocity across active feedback channels.</div>', unsafe_allow_html=True)
                if submission_velocity_df.empty:
                    render_empty_state("No live activity available", "Try changing filters or wait for new submissions.", "FLOW")
                else:
                    st.plotly_chart(
                        plot_dark_line(submission_velocity_df, "bucket_label", "count", "Submissions", "#38bdf8", area=True),
                        use_container_width=True,
                        key="overview_submission_trend_live",
                    )
            with row1_right:
                st.markdown(f'<div class="panel-title">Feedback Distribution</div><div class="panel-subtitle">{role_distribution_subtitle}</div>', unsafe_allow_html=True)
                if sum(int(v.get("total_submissions", 0)) for v in role_summaries.values()) <= 0:
                    render_empty_state("No role distribution available", "Role comparison appears when scoped category records are available.", "ROLE")
                else:
                    st.plotly_chart(plot_role_distribution(role_summaries), use_container_width=True, key="overview_role_distribution")

            row2_left, row2_right = st.columns([1.25, 1.35])
            with row2_left:
                st.markdown('<div class="panel-title">Category Analytics</div><div class="panel-subtitle">Top category forms ranked by score and contribution.</div>', unsafe_allow_html=True)
                top_forms = scoped_category_summary.get("top_forms", []) or []
                if top_forms:
                    st.plotly_chart(plot_category_scores(scoped_category_summary), use_container_width=True, key="overview_category_scores")
                else:
                    render_empty_state(
                        "No category feedback found for the selected filters",
                        "Try changing month, district, trade, year, semester, or role group.",
                        "FORM",
                    )
            with row2_right:
                st.markdown('<div class="panel-title">Sentiment Trend</div><div class="panel-subtitle">Positive, neutral and negative movement over time.</div>', unsafe_allow_html=True)
                if sentiment_summary.get("total", 0) <= 0 or sentiment_trend_df.empty:
                    render_empty_state("No sentiment data available for selected filters", "Sentiment widgets appear only when filtered sentiment records exist.", "SENT")
                else:
                    st.plotly_chart(plot_sentiment_trend(sentiment_trend_df), use_container_width=True, key="overview_sentiment_trend")

            row3_left, row3_right = st.columns([1.35, 1])
            with row3_left:
                st.markdown(f'<div class="panel-title">Role Comparison</div><div class="panel-subtitle">{role_comparison_subtitle}</div>', unsafe_allow_html=True)
                if sum(int(v.get("total_submissions", 0)) for v in role_summaries.values()) <= 0:
                    render_empty_state("No role ratings available", "Role ratings will appear when the selected scope has category records.", "RATE")
                else:
                    st.plotly_chart(plot_role_comparison(role_summaries), use_container_width=True, key="overview_role_comparison")
            with row3_right:
                st.markdown('<div class="panel-title">Data Readiness</div><div class="panel-subtitle">Shows whether live or demo data is currently driving the dashboard.</div>', unsafe_allow_html=True)
                st.markdown(
                    f"""
                    <div class="dashboard-card">
                        <div class="panel-title">Current Data Source</div>
                        <div class="panel-subtitle">{'Demo Preview Data' if using_demo_preview else 'Real mirrored/API-backed data'}</div>
                        <div class="kpi-sub">API Connected: {'Yes' if live_bundle.get('connected') else 'No'}</div>
                        <div class="kpi-sub">Filtered category records: {len(scoped_category_filtered)}</div>
                        <div class="kpi-sub">Filtered sentiment-tech records: {len(monthly_filtered)}</div>
                    </div>
                    """,
                    unsafe_allow_html=True,
                )

            row4_left, row4_right = st.columns(2)
            with row4_left:
                render_activity_panel(detail_df)
            with row4_right:
                render_alert_panel(alerts)

            st.markdown('<div class="section-label">Detailed Feedback Data Table</div>', unsafe_allow_html=True)
            search_term = st.text_input("Search feedback records", value="", placeholder="Search by name, role, trade or sentiment")
            sort_by = st.selectbox("Sort by", ["Submitted Time", "Rating", "Role", "Trade"], index=0)
            page_size = st.selectbox("Rows per page", [5, 10, 20, 30], index=1)
            working_df = detail_df.copy()
            if search_term.strip() and not working_df.empty:
                mask = working_df[["Name", "Role", "Trade", "Sentiment", "Status"]].astype(str).apply(
                    lambda col: col.str.contains(search_term.strip(), case=False, na=False)
                )
                working_df = working_df[mask.any(axis=1)]
            if not working_df.empty:
                sort_col = "Submitted Time" if sort_by == "Submitted Time" else sort_by
                ascending = sort_by not in {"Submitted Time", "Rating"}
                if sort_col in working_df.columns:
                    working_df = working_df.sort_values(by=sort_col, ascending=ascending)
                total_pages = max(1, (len(working_df) + page_size - 1) // page_size)
                page = st.number_input("Page", min_value=1, max_value=total_pages, value=1, step=1)
                start_idx = (int(page) - 1) * page_size
                paged_df = working_df.iloc[start_idx : start_idx + page_size].copy()
                st.dataframe(
                    paged_df.drop(columns=[c for c in ["_ts", "_comment"] if c in paged_df.columns], errors="ignore"),
                    use_container_width=True,
                    hide_index=True,
                )
            else:
                render_empty_state("No filtered records found", "Try widening the filters or enable demo preview mode.", "DATA")

        with live_tab:
            st.markdown("### Live Operations Feed")
            live_a, live_b = st.columns([1.4, 1])
            with live_a:
                render_activity_panel(detail_df)
            with live_b:
                render_alert_panel(alerts)
                st.markdown(
                    f"""
                        <div class="dashboard-card">
                            <div class="panel-title">Live Sync Status</div>
                            <div class="panel-subtitle">Operational heartbeat and freshness indicators.</div>
                            <div class="kpi-value" style="font-size:1.5rem;">{st.session_state.dashboard_last_synced.strftime("%I:%M:%S %p")}</div>
                            <div class="kpi-sub">Health latency {safe_float(live_bundle.get('health_response_ms', 0.0)):.1f} ms</div>
                            <div class="kpi-sub">{'Demo Preview Data' if using_demo_preview else 'Live filtered view'}</div>
                        </div>
                        """,
                        unsafe_allow_html=True,
                )

        with category_tab:
            st.markdown("### Category Analytics")
            cat_a, cat_b = st.columns(2)
            with cat_a:
                top_forms = scoped_category_summary.get("top_forms", []) or []
                if top_forms:
                    st.plotly_chart(plot_category_scores(scoped_category_summary), use_container_width=True, key="category_tab_scores")
                else:
                    render_empty_state(
                        "No category feedback found for the selected filters",
                        "Try changing month, district, trade, year, semester, or role group.",
                        "FORM",
                    )
            with cat_b:
                if sum(int(v.get("total_submissions", 0)) for v in category_role_summaries.values()) <= 0:
                    render_empty_state("No role-wise category comparison available", "Role comparison appears when category records exist for the current scope.", "ROLE")
                else:
                    st.plotly_chart(plot_role_comparison(category_role_summaries), use_container_width=True, key="category_tab_role_comparison")
            if not category_filtered.empty:
                top_comments = category_filtered[["form_title", "comment_text", "avg_rating_score"]].copy()
                top_comments = top_comments.sort_values(by="avg_rating_score").head(10)
                st.dataframe(top_comments, use_container_width=True, hide_index=True)
            else:
                render_empty_state("No category detail rows available", "Detailed category submissions appear when records match the current scope.", "DETAIL")
            if current_login_role == "principal":
                st.markdown("### Individual Performance by Department")
                st.caption("Each table shows person-level average rating and submission count for the selected department scope.")
                trainer_tab, ojt_tab, supervisor_tab = st.tabs(["Trainer", "OJT Trainer", "Supervisor"])
                for role_label, role_tab in [
                    ("Trainer", trainer_tab),
                    ("OJT Trainer", ojt_tab),
                    ("Supervisor", supervisor_tab),
                ]:
                    with role_tab:
                        role_df = individual_performance_df[individual_performance_df["Role"].astype(str).str.lower() == role_label.lower()].copy()
                        if role_df.empty:
                            render_empty_state(
                                f"No {role_label.lower()} analytics found",
                                "Try changing the department, month, district, year, semester, or role group filters.",
                                "DEPT",
                            )
                        else:
                            st.dataframe(role_df, use_container_width=True, hide_index=True)

        with sentiment_tab:
            st.markdown("### Sentiment Monitor")
            sent_a, sent_b = st.columns([1.4, 1])
            with sent_a:
                if sentiment_summary.get("total", 0) <= 0 or sentiment_trend_df.empty:
                    render_empty_state("No sentiment data available for selected filters", "Do not infer sentiment trends when the filtered sentiment dataset is empty.", "SENT")
                else:
                    st.plotly_chart(plot_sentiment_trend(sentiment_trend_df), use_container_width=True, key="sentiment_tab_trend")
            with sent_b:
                if sentiment_summary.get("total", 0) <= 0:
                    render_empty_state("No sentiment snapshot available", "Sentiment summary appears when sentiment records match the current filters.", "SENT")
                else:
                    st.markdown(
                        f"""
                        <div class="dashboard-card">
                            <div class="panel-title">Sentiment Snapshot</div>
                            <div class="panel-subtitle">All sentiment widgets use the same filtered sentiment dataset.</div>
                            <div class="kpi-sub">Positive: {safe_float(sentiment_summary.get('positive_pct', 0.0)):.1f}%</div>
                            <div class="kpi-sub">Neutral: {safe_float(sentiment_summary.get('neutral_pct', 0.0)):.1f}%</div>
                            <div class="kpi-sub">Negative: {safe_float(sentiment_summary.get('negative_pct', 0.0)):.1f}%</div>
                            <div class="kpi-sub">Total sentiment records: {int(sentiment_summary.get('total', 0))}</div>
                            <div class="kpi-sub">Negative sentiment alerts: {int(sentiment_summary.get('negative_count', 0))}</div>
                        </div>
                        """,
                        unsafe_allow_html=True,
                    )
                render_alert_panel(alerts)

        with health_tab:
            st.markdown("### Technical Health")
            h1, h2, h3, h4, h5, h6 = st.columns(6)
            with h1:
                st.metric("API Status", "Connected" if live_bundle.get("connected") else "Disconnected")
            with h2:
                st.metric("DB/Mirror Status", "Mirrored Snapshot Ready" if not mirrored_data.get("monthly", pd.DataFrame()).empty or not mirrored_data.get("category", pd.DataFrame()).empty else "No local mirror")
            with h3:
                st.metric("Last Response Time", f"{safe_float(live_bundle.get('summary_response_ms', 0.0)):.1f} ms")
            with h4:
                st.metric("Total Records Loaded", safe_count(monthly_filtered) + safe_count(scoped_category_filtered))
            with h5:
                st.metric("Last Refresh Time", st.session_state.dashboard_last_synced.strftime("%I:%M:%S %p"))
            with h6:
                st.metric("Refresh Mode", "Auto" if st.session_state.dashboard_auto_refresh else "Manual")
            health_a, health_b = st.columns([1.2, 1])
            with health_a:
                st.markdown(
                    f"""
                    <div class="dashboard-card">
                        <div class="panel-title">Data Sync Status</div>
                        <div class="panel-subtitle">System-health view only. No topic analytics are shown here.</div>
                        <div class="kpi-sub">Category records loaded: {len(scoped_category_filtered)}</div>
                        <div class="kpi-sub">Sentiment-tech records loaded: {len(monthly_filtered)}</div>
                        <div class="kpi-sub">Current data source: {'Demo Preview Data' if using_demo_preview else 'Real mirrored/API-backed data'}</div>
                        <div class="kpi-sub">API latency: {safe_float(live_bundle.get('health_response_ms', 0.0)):.1f} ms</div>
                    </div>
                    """,
                    unsafe_allow_html=True,
                )
            with health_b:
                errors = live_bundle.get("errors", [])
                if errors:
                    st.markdown("#### API Error Logs")
                    for err in errors:
                        st.code(value_or_empty(err))
                else:
                    render_empty_state("No API error logs", "The current refresh window did not surface backend or connectivity issues.", "HEALTH")

        with reports_tab:
            st.markdown("### Reports")
            report_a, report_b = st.columns([1.25, 1])
            with report_a:
                if report_export_df.empty:
                    render_empty_state("No records available for export under current filters.", "Adjust the filters or enable demo preview mode to generate report outputs.", "REPORT")
                else:
                    st.dataframe(report_export_df.head(50), use_container_width=True, hide_index=True)
            with report_b:
                render_report_actions(report_actions)
                st.markdown(
                    f"""
                    <div class="dashboard-card">
                        <div class="panel-title">Report Scope</div>
                        <div class="panel-subtitle">Current exports reflect the exact dashboard filters and do not silently mix real and demo data.</div>
                        <div class="kpi-sub">Month: {report_scope.get('Month')}</div>
                        <div class="kpi-sub">District: {report_scope.get('District')}</div>
                        <div class="kpi-sub">{'Department' if role == 'principal' else 'Trade'}: {selected_trade}</div>
                        <div class="kpi-sub">Year: {report_scope.get('Year')}</div>
                        <div class="kpi-sub">Semester: {report_scope.get('Semester')}</div>
                        <div class="kpi-sub">Role Group: {report_scope.get('Role Group')}</div>
                        <div class="kpi-sub">Record Count: {report_scope.get('Record Count')}</div>
                    </div>
                    """,
                    unsafe_allow_html=True,
                )
